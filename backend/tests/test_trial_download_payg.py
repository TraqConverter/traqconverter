"""Trial downloads carry a watermark; credit packs without a subscription give pay-as-you-go access."""
import io
import re
import uuid
import zipfile

import fitz
import pytest
from docx import Document

from app.core.plan_features import PLAN_FEATURES, SEAT_LIMITS
from app.dependencies import feature_guard
from app.models.credit import CreditWallet
from app.services import trial_watermark
from app.services.credit_service import CreditService
from app.services.trial_watermark import WATERMARK_TEXT
from tests.test_delivery import stub_export  # noqa: F401
from tests.test_stripe_portal import _post_event, _sub_event, _subscriber, prices  # noqa: F401

MARK = 'string="Trial – OnlineDocTranslator.ai"'


def _header_xml(docx: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(docx)) as z:
        return [z.read(n).decode() for n in z.namelist() if re.fullmatch(r"word/header\d+\.xml", n)]


def _docx_marked(docx: bytes) -> bool:
    return any(MARK in x for x in _header_xml(docx))


def _pdf_pages(pdf: bytes) -> list[str]:
    with fitz.open(stream=pdf, filetype="pdf") as doc:
        return [p.get_text() for p in doc]


def _wallet(db, who, **values):
    db.query(CreditWallet).filter(CreditWallet.team_id == who["team"].id).update(values)
    db.commit()


def _plan(db, who):
    db.expire_all()
    return feature_guard.effective_plan(db, who["user"])


def _get(client, who, project, path=""):
    return client.get(f"/projects/{project.id}/export{path}", headers=who["headers"])


def _ended_trial(db, make_user, purchased=0):
    who = make_user(plan="TRIAL", credits=0, expires_in_days=-1)
    _wallet(db, who, subscription_status="EXPIRED", purchased_credits=purchased)
    return who


# The watermark itself

def test_watermark_pdf_stamps_every_page_including_rotated_ones():
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Portrait")
    doc.new_page(width=842, height=595).insert_text((72, 72), "Landscape")
    rotated = doc.new_page()
    rotated.insert_text((72, 72), "Rotated")
    rotated.set_rotation(90)
    pages = _pdf_pages(trial_watermark.watermark_pdf(doc.tobytes()))
    assert len(pages) == 3
    for text in pages:
        # One large diagonal line and two smaller ones, written into the page content.
        assert text.count(WATERMARK_TEXT) == 3
    with fitz.open(stream=trial_watermark.watermark_pdf(doc.tobytes()), filetype="pdf") as out:
        assert out[2].rotation == 90
        # Page content, not an annotation a viewer can hide or delete.
        assert not list(out[0].annots()) and not list(out[0].widgets())


def test_watermark_docx_reaches_every_header_a_page_can_show():
    doc = Document()
    doc.add_paragraph("One")
    second = doc.add_section()
    doc.add_paragraph("Two")
    second.different_first_page_header_footer = True
    buf = io.BytesIO()
    doc.save(buf)
    out = trial_watermark.watermark_docx(buf.getvalue())
    marked = Document(io.BytesIO(out))
    first, other = marked.sections
    assert not first.header.is_linked_to_previous
    # The second section's default header shows the first's; its first-page header needs its own.
    assert other.header.is_linked_to_previous
    assert not other.first_page_header.is_linked_to_previous
    headers = _header_xml(out)
    assert len(headers) == 2 and all(MARK in x for x in headers)
    for x in headers:
        assert "rotation:315" in x and "z-index:-" in x
        assert "PowerPlusWaterMarkObject" not in x
    assert [p.text for p in marked.paragraphs if p.text] == ["One", "Two"]


def test_watermark_docx_keeps_an_existing_header():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Studio Rossi"
    doc.add_paragraph("Body")
    buf = io.BytesIO()
    doc.save(buf)
    out = Document(io.BytesIO(trial_watermark.watermark_docx(buf.getvalue())))
    assert "Studio Rossi" in [p.text for p in out.sections[0].header.paragraphs]


# Trial exports

def test_trial_docx_export_is_watermarked(client, make_user, make_project):
    trial = make_user(plan="TRIAL")
    r = _get(client, trial, make_project(trial))
    assert r.status_code == 200
    assert _docx_marked(r.content)


def test_trial_pdf_export_is_watermarked_on_every_page(client, make_user, make_project):
    trial = make_user(plan="TRIAL")
    r = _get(client, trial, make_project(trial), "/pdf")
    assert r.status_code == 200
    pages = _pdf_pages(r.content)
    assert pages and all(WATERMARK_TEXT in t for t in pages)


def test_trial_delivery_pdf_is_watermarked_on_every_page(client, storage, make_user, make_project, stub_export):  # noqa: F811
    from tests.test_delivery import _original, _pdf

    trial = make_user(plan="TRIAL")
    project = make_project(trial)
    _original(storage, project, _pdf(2, "Original"), "o.pdf")
    r = _get(client, trial, project, "/delivery.pdf")
    assert r.status_code == 200, r.text
    pages = _pdf_pages(r.content)
    assert len(pages) > 2 and all(WATERMARK_TEXT in t for t in pages)


@pytest.mark.parametrize("plan", ["BASIC", "PRO"])
def test_paid_exports_are_clean(client, make_user, make_project, plan):
    owner = make_user(plan=plan)
    project = make_project(owner)
    docx = _get(client, owner, project)
    pdf = _get(client, owner, project, "/pdf")
    assert docx.status_code == pdf.status_code == 200
    assert not _docx_marked(docx.content)
    assert not any(WATERMARK_TEXT in t for t in _pdf_pages(pdf.content))


def test_trial_cannot_share_with_a_client(client, make_user, make_project):
    trial = make_user(plan="TRIAL")
    project = make_project(trial)
    for body in ({"kind": "pdf"}, {"protected": True, "amount": "40"}):
        r = client.post(f"/projects/{project.id}/delivery-links", json=body, headers=trial["headers"])
        assert r.status_code == 403


def test_ended_trial_without_packs_cannot_download(client, db, make_user, make_project):
    who = _ended_trial(db, make_user)
    assert _plan(db, who) == "EXPIRED"
    assert _get(client, who, make_project(who)).status_code == 403


# Pay as you go

def test_payg_is_basic_plus_certifications_on_one_seat():
    payg = PLAN_FEATURES["PAYG"]
    assert payg == {**PLAN_FEATURES["BASIC"], "certifications": True}
    assert payg["clean_download"] and payg["client_sharing"]
    assert not payg["terminology_memory"] and not payg["glossaries"]
    assert SEAT_LIMITS["PAYG"] == 1


@pytest.mark.parametrize("setup", ["ended_trial", "running_trial", "cancelled"])
def test_purchased_credits_without_a_subscription_give_payg(client, db, make_user, make_project, setup):
    if setup == "ended_trial":
        who = _ended_trial(db, make_user, purchased=10)
    elif setup == "running_trial":
        who = make_user(plan="TRIAL")
        _wallet(db, who, purchased_credits=10)
    else:
        who = make_user(plan="PRO")
        _wallet(db, who, plan_type="EXPIRED", subscription_status="INACTIVE", subscription_credits=0, purchased_credits=10)
    assert _plan(db, who) == "PAYG"
    assert feature_guard.user_has_feature(db, who["user"], "certifications")
    assert not feature_guard.user_has_feature(db, who["user"], "glossaries")
    r = _get(client, who, make_project(who))
    assert r.status_code == 200 and not _docx_marked(r.content)
    wallet = client.get("/billing/wallet", headers=who["headers"]).json()
    assert wallet["tier"] == "PAYG" and wallet["pack_pricing"] == "payg"


def test_payg_seat_limit_is_one(client, db, make_user):
    who = _ended_trial(db, make_user, purchased=10)
    r = client.post("/members/invite", json={"email": "x@traqtest.io", "role": "MEMBER"}, headers=who["headers"])
    assert r.status_code == 403, r.text
    assert r.json()["detail"].startswith("Your pay-as-you-go plan includes up to 1 team members")


def test_payg_ends_when_purchased_credits_hit_zero(client, db, make_user, make_project):
    who = _ended_trial(db, make_user, purchased=2)
    unpaid = make_project(who)
    paid = make_project(who)
    assert _plan(db, who) == "PAYG"
    CreditService.deduct_credits(db, who["team"].id, 2, reference_id=str(paid.id))
    db.commit()
    assert _plan(db, who) == "EXPIRED"
    assert not feature_guard.user_has_feature(db, who["user"], "certifications")
    assert _get(client, who, unpaid).status_code == 403
    # The translation the last purchased credits paid for stays downloadable, clean.
    r = _get(client, who, paid)
    assert r.status_code == 200 and not _docx_marked(r.content)


def test_running_trial_falls_back_to_watermarks_when_packs_run_out(client, db, make_user, make_project):
    who = make_user(plan="TRIAL", credits=0)
    _wallet(db, who, purchased_credits=1)
    assert _plan(db, who) == "PAYG"
    _wallet(db, who, purchased_credits=0)
    assert _plan(db, who) == "TRIAL"
    r = _get(client, who, make_project(who))
    assert r.status_code == 200 and _docx_marked(r.content)


def test_refunded_purchase_does_not_keep_a_clean_download(db, make_user, make_project):
    who = _ended_trial(db, make_user, purchased=2)
    project = make_project(who)
    CreditService.deduct_credits(db, who["team"].id, 2, reference_id=str(project.id))
    db.flush()
    assert trial_watermark.paid_with_purchased_credits(db, project)
    CreditService.refund_usage(db, str(project.id))
    db.commit()
    assert not trial_watermark.paid_with_purchased_credits(db, project)


def test_subscription_credits_do_not_count_as_a_purchase(db, make_user, make_project):
    owner = make_user(plan="PRO")
    project = make_project(owner)
    CreditService.deduct_credits(db, owner["team"].id, 2, reference_id=str(project.id))
    db.commit()
    assert not trial_watermark.paid_with_purchased_credits(db, project)


# Cancellation falls back to pay as you go

def _deleted(owner):
    return {
        "id": f"evt_{uuid.uuid4().hex}",
        "type": "customer.subscription.deleted",
        "data": {"object": {
            "id": "sub_live", "customer": "cus_1", "status": "canceled", "ended_at": 1790000000,
            "metadata": {"user_id": str(owner["user"].id), "team_id": str(owner["team"].id)},
        }},
    }


def test_cancelled_subscription_with_purchased_credits_falls_back_to_payg(client, db, make_user, make_project, monkeypatch):
    owner = _subscriber(db, make_user, plan="STUDIO")
    _wallet(db, owner, purchased_credits=12)
    assert _post_event(client, monkeypatch, _deleted(owner)).json() == {"status": "success"}
    assert _plan(db, owner) == "PAYG"
    db.expire_all()
    wallet = db.query(CreditWallet).filter(CreditWallet.team_id == owner["team"].id).one()
    assert (wallet.subscription_credits, wallet.purchased_credits) == (0, 12)
    r = _get(client, owner, make_project(owner))
    assert r.status_code == 200 and not _docx_marked(r.content)


def test_cancelled_subscription_without_purchased_credits_expires(client, db, make_user, monkeypatch):
    owner = _subscriber(db, make_user, plan="PRO")
    _post_event(client, monkeypatch, _deleted(owner))
    assert _plan(db, owner) == "EXPIRED"


def test_unpaid_subscription_with_purchased_credits_falls_back_to_payg(client, db, make_user, monkeypatch, prices):  # noqa: F811
    owner = _subscriber(db, make_user, plan="PRO")
    _wallet(db, owner, purchased_credits=3)
    _post_event(client, monkeypatch, _sub_event(owner, "price_pro", status="unpaid"))
    assert _plan(db, owner) == "PAYG"


def test_purchased_credits_survive_trial_expiry_and_stay_spendable(db, make_user):
    who = make_user(plan="TRIAL", credits=3, expires_in_days=-1)
    _wallet(db, who, purchased_credits=5)
    assert CreditService.deduct_credits(db, who["team"].id, 5) == 0
    db.commit()
    db.expire_all()
    wallet = db.query(CreditWallet).filter(CreditWallet.team_id == who["team"].id).one()
    assert (wallet.subscription_credits, wallet.purchased_credits) == (0, 0)
