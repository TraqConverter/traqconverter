import io
import json
import re
from datetime import datetime, timedelta
from types import SimpleNamespace

import pytest
from docx import Document

from app.models.learning import DocumentTemplate, TemplateUpload
from app.models.translation_memory import TranslationMemory
from app.services import claude_params, docx_blocks, learning, template_upload

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PROFILE = {
    "document_type": "residence certificate",
    "country": "IT",
    "issuing_authority": "Comune di Esempio",
    "source_language": "Italian",
    "format_variant": "",
    "title": "CERTIFICATO DI RESIDENZA",
}
SOURCE_LINES = [
    "COMUNE DI ESEMPIO",
    "CERTIFICATO DI RESIDENZA",
    "Si certifica che Mario Rossi è residente",
    "in Via Roma 12, Esempio.",
    "Rilasciato il 01/02/2024",
    "L'Ufficiale d'Anagrafe",
]
TRANSLATION = [
    "MUNICIPALITY OF ESEMPIO",
    "RESIDENCE CERTIFICATE",
    "It is hereby certified that Mario Rossi resides at Via Roma 12, Esempio.",
    "Issued on 01/02/2023",
    "[stamp]",
    "The Registry Officer",
]


def _docx(*paras: str) -> bytes:
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf() -> bytes:
    import fitz

    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "COMUNE DI ESEMPIO")
    data = doc.tobytes()
    doc.close()
    return data


class _FakeClaude:
    """Answers in order; an answer can be a dict, an exception, or a function of the request."""

    def __init__(self, answers):
        self.answers = list(answers)
        self.requests = []

    def __call__(self, client, **kwargs):
        self.requests.append(kwargs)
        body = self.answers.pop(0)
        if isinstance(body, Exception):
            raise body
        if callable(body):
            body = body(kwargs)
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=json.dumps(body))],
            stop_reason="end_turn",
            usage=SimpleNamespace(input_tokens=100, output_tokens=20, cache_read_input_tokens=0, cache_creation_input_tokens=0),
            model=kwargs["model"],
        )


@pytest.fixture()
def fake_claude(monkeypatch):
    def install(*answers):
        fake = _FakeClaude(answers)
        monkeypatch.setattr(claude_params, "create_message", fake)
        monkeypatch.setattr(claude_params, "api_key", lambda: "test-key")
        return fake

    return install


@pytest.fixture(autouse=True)
def source_lines(monkeypatch):
    # The real extraction may OCR with a model call; tests never reach the API.
    monkeypatch.setattr(template_upload, "extract_source_lines", lambda data, name: list(SOURCE_LINES))


def _prompt(kwargs) -> str:
    return kwargs["messages"][0]["content"]


def _ids_by_text(kwargs) -> dict[str, str]:
    block = _prompt(kwargs).rsplit("TRANSLATION PARAGRAPHS", 1)[1]
    return {m.group(2): m.group(1) for m in re.finditer(r"^(_b[0-9a-f]{8}): (.*)$", block, re.M)}


def _align(*groups):
    """Alignment answer: (translation text, source line indexes) pairs."""

    def answer(kwargs):
        ids = _ids_by_text(kwargs)
        return {"pairs": [{"paragraph_id": ids.get(text, "_b00000000"), "lines": lines} for text, lines in groups]}

    return answer


def _analyze(client, owner, original=None, translation=None, target="en-GB", names=("scan.pdf", "translation.docx")):
    return client.post(
        "/templates/analyze",
        files={
            "original": (names[0], original if original is not None else _pdf(), "application/octet-stream"),
            "translation": (names[1], translation if translation is not None else _docx(*TRANSLATION), DOCX),
        },
        data={"target_language": target},
        headers=owner["headers"],
    )


def _save(client, owner, upload_id, **fields):
    body = {"upload_id": upload_id, "target_language": "en-GB", **{k: v for k, v in PROFILE.items() if k != "title"}}
    body.update(fields)
    return client.post("/templates/from-upload", json=body, headers=owner["headers"])


def test_analyze_validates_files(client, make_user, storage, fake_claude):
    fake = fake_claude()
    owner = make_user()
    cases = [
        ({"names": ("scan.docx", "translation.docx")}, "PDF, JPG or PNG"),
        ({"names": ("scan.pdf", "translation.pdf")}, "Word .docx"),
        ({"original": b"%PDF" + b"0" * (template_upload.MAX_BYTES + 1)}, "20 MB"),
        ({"translation": b"0" * (template_upload.MAX_BYTES + 1)}, "20 MB"),
        ({"translation": b"not a zip file"}, "couldn't be opened"),
        ({"translation": _docx("", "   ")}, "no text"),
        ({"target": "auto"}, "language"),
    ]
    for kwargs, message in cases:
        resp = _analyze(client, owner, **kwargs)
        assert resp.status_code == 422, kwargs
        assert message in resp.json()["detail"]
    assert fake.requests == []
    assert storage["objects"] == {}


def test_analyze_needs_a_paid_plan(client, make_user, fake_claude):
    fake_claude()
    assert _analyze(client, make_user(plan="TRIAL")).status_code == 403


def test_analyze_returns_profile(client, make_user, storage, fake_claude, db):
    fake = fake_claude(PROFILE)
    owner = make_user()
    resp = _analyze(client, owner)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["profile"] == {
        "document_type": "residence certificate",
        "country": "IT",
        "issuing_authority": "Comune di Esempio",
        "format_variant": "",
        "source_language": "Italian",
    }
    assert body["doc_key"] == "residence-certificate.it.comune-esempio"
    assert body["title"] == "Residence certificate · Comune di Esempio"
    assert body["existing_template"] is None
    assert body["source_lines"] == len(SOURCE_LINES)
    assert len(fake.requests) == 1 and fake.requests[0]["model"] == claude_params.CLASSIFIER_MODEL
    assert sorted(k.split("/")[0] for k in storage["objects"]) == ["pending-templates", "pending-templates"]
    row = db.query(TemplateUpload).one()
    assert row.team_id == owner["team"].id
    assert row.expires_at - row.created_at == timedelta(hours=24)


def test_from_upload_builds_template_with_block_ids(client, make_user, storage, fake_claude, db):
    fake_claude(PROFILE, {"pairs": []})
    owner = make_user()
    upload_id = _analyze(client, owner).json()["upload_id"]
    pending_keys = set(storage["objects"])

    check = client.get(
        "/templates/key-check",
        params={"upload_id": upload_id, "target_language": "en-GB", "document_type": "Residence Certificate",
                "country": "it", "issuing_authority": "Comune di Altrove"},
        headers=owner["headers"],
    )
    assert check.status_code == 200
    assert check.json()["doc_key"] == "residence-certificate.it.comune-altrove"

    resp = _save(client, owner, upload_id, document_type="Residence Certificate", country="it",
                 issuing_authority="Comune di Altrove", format_variant="ANPR model")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["replaced"] is False
    assert body["template"]["doc_key"] == "residence-certificate.it.comune-altrove.anpr-model"
    assert body["template"]["title"] == "Residence certificate · Comune di Altrove"
    assert body["template"]["source_project_id"] is None

    template = db.query(DocumentTemplate).one()
    assert template.source_text == "\n".join(SOURCE_LINES)
    assert template.doc_profile["issuing_authority"] == "Comune di Altrove"
    data = storage["objects"][template.s3_key]
    assert len(docx_blocks.block_ids(data)) == len(TRANSLATION)
    assert docx_blocks.outline(data)
    # The pending upload and its files are gone.
    assert db.query(TemplateUpload).count() == 0
    assert pending_keys <= set(storage["deleted"])

    # Found for the next document of the same kind, from another town.
    found = learning.find_template(db, owner["team"].id, "residence-certificate.it.comune-esempio", "English (UK)")
    assert found is not None and found.id == template.id


def test_from_upload_replaces_existing_template(client, make_user, storage, fake_claude, db):
    owner = make_user()
    old = DocumentTemplate(
        team_id=owner["team"].id,
        doc_key="residence-certificate.it.comune-esempio",
        target_language=learning.lang_key("en-GB"),
        title="Old",
        doc_profile=PROFILE,
        s3_key="uploads/old_template.docx",
        source_text="old",
        use_count=3,
    )
    db.add(old)
    db.commit()
    fake_claude(PROFILE, {"pairs": []})
    analysed = _analyze(client, owner).json()
    assert analysed["existing_template"]["id"] == str(old.id)

    resp = _save(client, owner, analysed["upload_id"])
    assert resp.status_code == 200, resp.text
    assert resp.json()["replaced"] is True
    db.expire_all()
    rows = db.query(DocumentTemplate).all()
    assert len(rows) == 1 and rows[0].id == old.id
    assert rows[0].source_project_id is None and rows[0].use_count == 3
    assert rows[0].s3_key != "uploads/old_template.docx"
    assert "uploads/old_template.docx" in storage["deleted"]


def test_alignment_stores_approved_pairs_and_drops_failures(client, make_user, fake_claude, db):
    fake = fake_claude(
        PROFILE,
        _align(
            (TRANSLATION[0], [0]),
            (TRANSLATION[1], [1]),
            (TRANSLATION[2], [2, 3]),
            (TRANSLATION[3], [4]),  # the year differs
            (TRANSLATION[5], [5, 5]),  # not consecutive
            ("unknown paragraph", [5]),
        ),
    )
    owner = make_user()
    upload_id = _analyze(client, owner).json()["upload_id"]
    resp = _save(client, owner, upload_id)
    assert resp.status_code == 200, resp.text
    assert resp.json()["memory_lines_added"] == 3
    assert resp.json()["memory_note"] is None

    request = fake.requests[1]
    assert request["model"] == claude_params.CLASSIFIER_MODEL
    assert request["output_config"]["format"]["type"] == "json_schema"
    prompt = _prompt(request)
    assert "[3] in Via Roma 12, Esempio." in prompt
    assert "[stamp]" not in prompt.rsplit("TRANSLATION PARAGRAPHS", 1)[1]

    rows = {r.source_text: r for r in db.query(TranslationMemory).all()}
    assert set(rows) == {
        "COMUNE DI ESEMPIO",
        "CERTIFICATO DI RESIDENZA",
        "Si certifica che Mario Rossi è residente in Via Roma 12, Esempio.",
    }
    for r in rows.values():
        assert r.origin == "approved"
        assert r.project_id is None
        assert r.team_id == owner["team"].id
        assert (r.source_language, r.target_language) == ("it", "en-GB")
    assert rows["CERTIFICATO DI RESIDENZA"].translated_text == "RESIDENCE CERTIFICATE"


def test_checked_pairs_rules():
    paragraphs = [{"id": f"_b0000000{i}", "text": t} for i, t in enumerate(TRANSLATION)]
    raw = [
        {"paragraph_id": "_b00000000", "lines": [0]},
        {"paragraph_id": "_b00000001", "lines": [0]},  # line already used
        {"paragraph_id": "_b00000002", "lines": [0, 1, 2, 3, 4]},  # too many lines
        {"paragraph_id": "_b00000002", "lines": [3, 2]},  # out of order
        {"paragraph_id": "_b00000005", "lines": [99]},  # out of range
        {"paragraph_id": "_b00000001", "lines": []},
    ]
    assert template_upload.checked_pairs(raw, SOURCE_LINES, paragraphs) == [("COMUNE DI ESEMPIO", "MUNICIPALITY OF ESEMPIO")]


def test_certification_page_is_not_aligned():
    doc = docx_blocks._Doc.load(docx_blocks.tag_blocks(_docx("Body text here", "I certify this translation", "After")))
    ps = list(doc.paragraphs())
    start, end = docx_blocks._marker_pair(900, "_cert_start")
    ps[1].insert(0, start)
    ps[1].append(end)
    cert_end_start, cert_end_end = docx_blocks._marker_pair(901, "_cert_end")
    ps[1].append(cert_end_start)
    ps[1].append(cert_end_end)
    texts = [p["text"] for p in template_upload.alignable_paragraphs(doc.dump())]
    assert "Body text here" in texts
    assert "I certify this translation" not in texts


def test_alignment_failure_still_saves_template(client, make_user, fake_claude, db):
    fake_claude(PROFILE, RuntimeError("model unavailable"))
    owner = make_user()
    upload_id = _analyze(client, owner).json()["upload_id"]
    resp = _save(client, owner, upload_id)
    assert resp.status_code == 200, resp.text
    assert resp.json()["memory_lines_added"] == 0
    assert db.query(DocumentTemplate).count() == 1
    assert db.query(TranslationMemory).count() == 0


def test_unknown_source_language_skips_memory(client, make_user, fake_claude, db):
    fake = fake_claude(PROFILE)
    owner = make_user()
    upload_id = _analyze(client, owner).json()["upload_id"]
    resp = _save(client, owner, upload_id, source_language="")
    assert resp.status_code == 200, resp.text
    assert resp.json()["memory_lines_added"] == 0
    assert "source language" in resp.json()["memory_note"]
    assert len(fake.requests) == 1
    assert db.query(DocumentTemplate).count() == 1


def test_memory_needs_pro_plan(client, make_user, fake_claude, db):
    fake = fake_claude(PROFILE)
    owner = make_user(plan="BASIC")
    upload_id = _analyze(client, owner).json()["upload_id"]
    resp = _save(client, owner, upload_id)
    assert resp.status_code == 200, resp.text
    assert resp.json()["memory_lines_added"] == 0
    assert "Pro" in resp.json()["memory_note"]
    assert len(fake.requests) == 1


def test_other_team_cannot_use_pending_upload(client, make_user, fake_claude, db):
    fake_claude(PROFILE)
    owner, stranger = make_user(), make_user()
    upload_id = _analyze(client, owner).json()["upload_id"]
    assert _save(client, stranger, upload_id).status_code == 404
    check = client.get(
        "/templates/key-check",
        params={"upload_id": upload_id, "target_language": "en-GB", "document_type": "passport"},
        headers=stranger["headers"],
    )
    assert check.status_code == 404
    assert db.query(TemplateUpload).count() == 1
    assert db.query(DocumentTemplate).count() == 0


def test_team_member_can_confirm_owner_upload(client, make_user, fake_claude, db):
    fake_claude(PROFILE, {"pairs": []})
    owner = make_user()
    member = make_user(team=owner["team"])
    upload_id = _analyze(client, owner).json()["upload_id"]
    assert _save(client, member, upload_id).status_code == 200


def test_pending_upload_expires(client, make_user, storage, fake_claude, db):
    fake_claude(PROFILE, PROFILE)
    owner = make_user()
    first = _analyze(client, owner).json()["upload_id"]
    second = _analyze(client, owner).json()["upload_id"]
    rows = {str(r.id): r for r in db.query(TemplateUpload).all()}
    for r in rows.values():
        r.expires_at = datetime.utcnow() - timedelta(minutes=1)
    rows[second].expires_at = datetime.utcnow() + timedelta(hours=1)
    first_key = rows[first].original_key
    db.commit()

    resp = _save(client, owner, first)
    assert resp.status_code == 410
    assert db.query(TemplateUpload).filter(TemplateUpload.id == first).count() == 0
    assert first_key in storage["deleted"]

    later = datetime.utcnow() + timedelta(hours=2)
    assert template_upload.purge_expired(now=later) == 1
    db.expire_all()
    assert db.query(TemplateUpload).count() == 0
    assert db.query(DocumentTemplate).count() == 0
