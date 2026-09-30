import io
import json
import time
import uuid
from datetime import datetime, timedelta
from types import SimpleNamespace

import pytest
from docx import Document

from app.models.glossary import Glossary
from app.models.learning import DocumentTemplate, PendingLearning
from app.services import claude_params, docx_blocks, glossary_service, learning, template_fill

KEY = "residence-certificate.it.comune-esempio"
PROFILE = {
    "document_type": "residence certificate",
    "country": "IT",
    "issuing_authority": "Comune di Esempio",
    "source_language": "Italian",
    "format_variant": "",
    "title": "CERTIFICATO DI RESIDENZA",
    "doc_key": KEY,
}
SOURCE = [
    "L'UFFICIALE D'ANAGRAFE, visti gli atti d'ufficio,",
    "che il Sig. BIANCHI LUCA, nato a Bari il 12/03/1987,",
    "Ufficio Anagrafe e Stato Civile",
]


def _docx(*paras: str) -> bytes:
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _texts(data: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(data)).paragraphs]


class _FakeClaude:
    def __init__(self, answers):
        self.answers = list(answers)
        self.requests = []

    def __call__(self, client, **kwargs):
        self.requests.append(kwargs)
        body = self.answers.pop(0)
        if isinstance(body, Exception):
            raise body
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=json.dumps(body))],
            stop_reason="end_turn",
            usage=SimpleNamespace(input_tokens=100, output_tokens=20, cache_read_input_tokens=0, cache_creation_input_tokens=0),
            model=kwargs["model"],
        )


@pytest.fixture()
def fake_claude(monkeypatch):
    def install(*answers):
        fake = _FakeClaude(answers)
        monkeypatch.setattr(claude_params, "create_message", fake)
        monkeypatch.setattr(claude_params, "api_key", lambda: "test-key")
        return fake

    return install


@pytest.fixture()
def no_claude(monkeypatch):
    def boom(*a, **kw):
        raise AssertionError("no model call expected here")

    monkeypatch.setattr(claude_params, "create_message", boom)


@pytest.fixture()
def doc_project(db, storage, make_user, make_project):
    def _make(owner=None, doc_key=KEY, target="English", paras=None):
        owner = owner or make_user()
        project = make_project(owner, segments=SOURCE)
        key = f"uploads/{uuid.uuid4()}_authored.docx"
        storage["objects"][key] = _docx(*(paras or (
            "CERTIFICATE OF RESIDENCE",
            "THE REGISTRY OFFICER, having examined the records,",
            "that Mr. BIANCHI LUCA, born in Bari on 12/03/1987,",
        )))
        project.authored_docx_s3_key = key
        project.target_language = target
        project.doc_key = doc_key
        project.doc_profile = dict(PROFILE, doc_key=doc_key) if doc_key else None
        db.commit()
        return owner, project

    return _make



def test_doc_key_is_deterministic_across_spelling_variants():
    a = learning.build_doc_key({"document_type": "Driving Licence", "country": "it", "issuing_authority": "Motorizzazione  Civile", "format_variant": "EU card model 2013"})
    b = learning.build_doc_key({"document_type": "driving licence ", "country": "IT", "issuing_authority": "MOTORIZZAZIONE CIVILE", "format_variant": "EU Card Model 2013"})
    assert a == b == "driving-licence.it.motorizzazione-civile.eu-card-model-2013"
    assert learning.build_doc_key({"document_type": "Certificato", "country": "", "issuing_authority": "Comune di Città"}) == "certificato.xx.comune-citta"
    assert learning.build_doc_key(PROFILE) == KEY


def test_classify_document_returns_profile_with_key(fake_claude):
    answer = {k: v for k, v in PROFILE.items() if k != "doc_key"}
    answer["country"] = "it"
    fake = fake_claude(answer)
    profile = learning.classify_document(b"hello", "scan.txt", "CERTIFICATO DI RESIDENZA")
    assert profile["doc_key"] == KEY and profile["country"] == "IT"
    req = fake.requests[0]
    assert req["model"] == claude_params.CLASSIFIER_MODEL
    assert req["output_config"]["format"]["type"] == "json_schema"


def test_profile_failure_never_raises(db, make_user, make_project, fake_claude):
    project = make_project(make_user())
    fake_claude(RuntimeError("api down"))
    assert learning.profile_project(db, project, b"x", "text") is None
    assert project.doc_key is None



def test_language_keys_match_codes_and_names():
    assert glossary_service.lang_key("it-IT") == glossary_service.lang_key("Italian") == glossary_service.lang_key("it") == "it"
    assert glossary_service.lang_key("English (UK)") == glossary_service.lang_key("en-US") == "en"
    assert glossary_service.lang_key("zh-TW") != glossary_service.lang_key("zh-CN")
    assert glossary_service.lang_key("auto") == ""


def test_glossary_matches_across_language_formats_and_skips_rejected(db, make_user):
    owner = make_user()
    team = owner["team"].id
    db.add_all([
        Glossary(team_id=team, source_language="Italian", target_language="English", source_term="Comune", target_term="Municipality"),
        Glossary(team_id=team, source_language="it", target_language="en", source_term="Anagrafe", target_term="Registry", origin="rejected"),
        Glossary(team_id=team, source_language="French", target_language="English", source_term="Mairie", target_term="Town hall"),
    ])
    db.commit()
    terms = glossary_service.get_glossary(db, team, "it-IT", "en-GB")
    assert [t.source_term for t in terms] == ["Comune"]
    assert glossary_service.get_glossary(db, team, "auto", "en-GB") == []



def test_export_captures_template(client, db, storage, doc_project, monkeypatch):
    import app.routers.export as export_router

    owner, project = doc_project()
    monkeypatch.setattr(export_router, "generate_docx", lambda *a, **kw: io.BytesIO(b"docx"))
    r = client.get(f"/projects/{project.id}/export", headers=owner["headers"])
    assert r.status_code == 200
    template = db.query(DocumentTemplate).one()
    assert (template.doc_key, template.target_language, template.team_id) == (KEY, "en", owner["team"].id)
    assert template.source_project_id == project.id and template.title.startswith("Residence certificate")
    assert "BIANCHI" in template.source_text
    assert _texts(storage["objects"][template.s3_key])[1].startswith("THE REGISTRY OFFICER")


def test_save_as_template_endpoint_and_newer_replaces_older(client, db, storage, doc_project):
    owner, first = doc_project()
    r = client.post(f"/projects/{first.id}/template", headers=owner["headers"])
    assert r.status_code == 200, r.text
    again = client.post(f"/projects/{first.id}/template", headers=owner["headers"])
    assert again.json()["id"] == r.json()["id"]
    assert db.query(DocumentTemplate).count() == 1

    _, second = doc_project(owner=owner, paras=("CERTIFICATE OF RESIDENCE", "THE REGISTRAR certifies"))
    client.post(f"/projects/{second.id}/template", headers=owner["headers"])
    db.expire_all()
    template = db.query(DocumentTemplate).one()
    assert template.source_project_id == second.id
    assert _texts(storage["objects"][template.s3_key])[1] == "THE REGISTRAR certifies"
    assert len(storage["deleted"]) == 1 and storage["deleted"][0] != template.s3_key


def test_certify_captures_template(client, db, doc_project):
    owner, project = doc_project()
    r = client.patch(f"/projects/{project.id}/review-status", headers=owner["headers"], json={"status": "CERTIFIED"})
    assert r.status_code == 200
    assert db.query(DocumentTemplate).count() == 1


def test_template_lookup_is_team_scoped_and_language_aware(client, db, doc_project, make_user):
    owner, project = doc_project(target="en-GB")
    client.post(f"/projects/{project.id}/template", headers=owner["headers"])
    other = make_user()

    assert learning.find_template(db, owner["team"].id, KEY, "English").id
    assert learning.find_template(db, owner["team"].id, KEY + ".multilingual-extract", "en-US") is not None
    assert learning.find_template(db, owner["team"].id, KEY, "French") is None
    assert learning.find_template(db, owner["team"].id, "driving-licence.it.motorizzazione-civile", "en") is None
    # Another town's residence certificate reuses it; another country's doesn't.
    assert learning.find_template(db, owner["team"].id, "residence-certificate.it.comune-altrove", "en") is not None
    assert learning.find_template(db, owner["team"].id, "residence-certificate.es.ayuntamiento-madrid", "en") is None
    assert learning.find_template(db, other["team"].id, KEY, "en") is None

    assert client.get("/templates", headers=other["headers"]).json() == []
    listed = client.get("/templates", headers=owner["headers"]).json()
    assert len(listed) == 1 and listed[0]["document_type"] == "residence certificate"
    tid = listed[0]["id"]
    assert client.delete(f"/templates/{tid}", headers=other["headers"]).status_code == 404
    assert client.delete(f"/templates/{tid}", headers=owner["headers"]).status_code == 200
    assert db.query(DocumentTemplate).count() == 0


def test_save_as_template_other_tenant_404(client, doc_project, make_user):
    _, project = doc_project()
    other = make_user()
    assert client.post(f"/projects/{project.id}/template", headers=other["headers"]).status_code == 404
    assert client.get(f"/projects/{project.id}/learning", headers=other["headers"]).status_code == 404



TEMPLATE_SOURCE = "Il Sig. BIANCHI LUCA, nato a Bari il 12/03/1987. Pratica 7781."
NEW_SOURCE = "La Sig.ra ROSSI GIULIA, nata a Lecce il 05/11/1992. Pratica 8120."


def _template_docx() -> bytes:
    return docx_blocks.tag_blocks(_docx("CERTIFICATE OF RESIDENCE", "Mr. BIANCHI LUCA, born in Bari on 12/03/1987.", "Case 7781"))


def _fill(**kw):
    return template_fill.fill_from_template(
        source_data=b"%PDF-1.4",
        file_name="new.pdf",
        source_text=NEW_SOURCE,
        template_docx=kw.pop("template", _template_docx()),
        template_source=TEMPLATE_SOURCE,
        source_lang="Italian",
        target_lang="English",
        terminology=kw.pop("terminology", ""),
    )


def _set_ops(template: bytes, texts: list[str]) -> dict:
    ids = docx_blocks.block_ids(template)
    return {"operations": [{"op": "set_text", "target": ids[i], "content": t} for i, t in texts], "notes": "x"}


def test_template_fill_applies_operations(fake_claude):
    template = _template_docx()
    fake = fake_claude(_set_ops(template, [(1, "Ms. ROSSI GIULIA, born in Lecce on 05/11/1992."), (2, "Case 8120")]))
    out, stats = _fill(template=template, terminology="TEAM TERMINOLOGY — use exactly these translations:\n- Comune → Municipality\n")
    assert _texts(out) == ["CERTIFICATE OF RESIDENCE", "Ms. ROSSI GIULIA, born in Lecce on 05/11/1992.", "Case 8120"]
    assert len(docx_blocks.block_ids(out)) == 3
    assert stats["calls"] == 1 and stats["leftovers"] == [] and stats["missing_numbers"] == 0
    req = fake.requests[0]
    assert req["messages"][0]["content"][0]["type"] == "document"
    assert req["messages"][0]["content"][0]["cache_control"]["type"] == "ephemeral"
    prompt = req["messages"][0]["content"][1]["text"]
    assert TEMPLATE_SOURCE in prompt and "Comune → Municipality" in prompt
    assert req["output_config"]["format"]["type"] == "json_schema"


def test_template_fill_retries_when_old_data_remains_then_gives_up(fake_claude):
    template = _template_docx()
    stale = _set_ops(template, [(2, "Case 8120")])
    fake = fake_claude(stale, stale)
    with pytest.raises(template_fill.TemplateFillError):
        _fill(template=template)
    assert len(fake.requests) == 2
    feedback = fake.requests[1]["messages"][-1]["content"]
    assert "BIANCHI" in feedback and "1987" in feedback


def test_template_fill_rejects_missing_source_numbers(fake_claude):
    template = _template_docx()
    wrong = _set_ops(template, [(1, "Ms. ROSSI GIULIA, born in Lecce."), (2, "Case")])
    fake_claude(wrong, wrong)
    with pytest.raises(template_fill.TemplateFillError):
        _fill(template=template)


def test_worker_uses_template_and_falls_back(db, storage, doc_project, make_user, make_project, monkeypatch, tmp_path):
    from app.services import translation_processor as tp

    owner, done = doc_project()
    learning.capture_template(db, done, owner["user"])
    template = db.query(DocumentTemplate).one()

    new = make_project(owner, segments=["La Sig.ra ROSSI GIULIA"])
    new.doc_key, new.doc_profile = KEY, PROFILE
    db.commit()

    seen = {}

    def fake_fill(**kw):
        seen.update(kw)
        return docx_blocks.tag_blocks(_docx("filled")), {"seconds": 1}

    monkeypatch.setattr(template_fill, "fill_from_template", fake_fill)
    job = tp._start_template_fill(db, new, "PDF", b"%PDF", "La Sig.ra ROSSI GIULIA", "TERMS")
    assert tp._finish_template_fill(db, new, job, tmp_path) is True
    db.refresh(template)
    assert new.template_id == template.id and template.use_count == 1 and template.last_used_at
    assert _texts(storage["objects"][new.authored_docx_s3_key]) == ["filled"]
    assert seen["terminology"] == "TERMS" and seen["template_source"] == template.source_text
    assert _texts(seen["template_docx"])[0] == "CERTIFICATE OF RESIDENCE"

    def failing(**kw):
        raise template_fill.TemplateFillError("coverage")

    monkeypatch.setattr(template_fill, "fill_from_template", failing)
    new.authored_docx_s3_key = None
    new.template_id = None
    job = tp._start_template_fill(db, new, "PDF", b"%PDF", "x", "")
    assert tp._finish_template_fill(db, new, job, tmp_path) is False
    assert new.template_id is None and new.authored_docx_s3_key is None

    assert tp._start_template_fill(db, new, "DOCX", b"x", "x", "") is None
    other = make_user()
    stranger = make_project(other)
    stranger.doc_key = KEY
    db.commit()
    assert tp._start_template_fill(db, stranger, "PDF", b"%PDF", "x", "") is None



def _open(client, owner, project):
    r = client.get(f"/projects/{project.id}/document", headers=owner["headers"])
    return r.content, int(r.headers["X-Document-Version"])


def test_typed_edit_only_queues_rows(client, db, doc_project, no_claude):
    owner, project = doc_project()
    data, v = _open(client, owner, project)
    ids = docx_blocks.block_ids(data)
    started = time.monotonic()
    r = client.post(
        f"/projects/{project.id}/document/edits",
        headers=owner["headers"],
        json={"version": v, "edits": [
            {"block_id": ids[1], "text": "THE REGISTRAR, having examined the records,"},
            {"block_id": ids[2], "text": "that Mr. BIANCHI LUCA, born in Bari on 12/03/1986,"},
        ]},
    )
    assert r.status_code == 200 and time.monotonic() - started < 2
    rows = db.query(PendingLearning).all()
    assert len(rows) == 1
    assert rows[0].before_text.startswith("THE REGISTRY OFFICER") and rows[0].after_text.startswith("THE REGISTRAR")


def test_processor_learns_terms_after_quiet_period(client, db, doc_project, fake_claude):
    owner, project = doc_project()
    data, v = _open(client, owner, project)
    bid = docx_blocks.block_ids(data)[1]
    client.post(f"/projects/{project.id}/document/edits", headers=owner["headers"],
                json={"version": v, "edits": [{"block_id": bid, "text": "THE REGISTRY CLERK, having examined the records,"}]})
    client.post(f"/projects/{project.id}/document/edits", headers=owner["headers"],
                json={"version": v + 1, "edits": [{"block_id": bid, "text": "THE REGISTRAR, having examined the records,"}]})

    fake = fake_claude({"terms": [
        {"source_term": "Ufficiale d'Anagrafe", "target_term": "Registrar", "confidence": 0.95},
        {"source_term": "Prefettura", "target_term": "Prefecture", "confidence": 0.9},
        {"source_term": "Stato Civile", "target_term": "Registrar", "confidence": 0.3},
    ]})
    assert learning.process_pending(now=datetime.utcnow()) == 0
    assert fake.requests == []

    assert learning.process_pending(now=datetime.utcnow() + timedelta(seconds=40)) == 1
    prompt = fake.requests[0]["messages"][0]["content"]
    assert "BEFORE: THE REGISTRY OFFICER" in prompt and "AFTER: THE REGISTRAR" in prompt and "BIANCHI" in prompt
    assert "REGISTRY CLERK" not in prompt
    term = db.query(Glossary).one()
    assert (term.source_term, term.target_term, term.origin, term.usage_count) == ("Ufficiale d'Anagrafe", "Registrar", "learned", 1)
    assert term.learned_from_project_id == project.id and term.source_language == "Italian" and term.target_language == "English"
    assert db.query(PendingLearning).count() == 0

    learning.upsert_terms(db, project, [{"source_term": "ufficiale d'anagrafe", "target_term": "Registrar", "confidence": 0.8}], "the Registrar", "Ufficiale d'Anagrafe")
    db.commit()
    db.refresh(term)
    assert term.usage_count == 2 and db.query(Glossary).count() == 1


def test_manual_terms_are_not_overwritten(db, doc_project):
    owner, project = doc_project()
    db.add(Glossary(team_id=owner["team"].id, source_language="Italian", target_language="English", source_term="Ufficiale d'Anagrafe", target_term="Registry Officer"))
    db.commit()
    learning.upsert_terms(db, project, [{"source_term": "Ufficiale d'Anagrafe", "target_term": "Registrar", "confidence": 0.9}], "Registrar", "Ufficiale d'Anagrafe")
    db.commit()
    assert db.query(Glossary).one().target_term == "Registry Officer"


def test_chat_edits_are_recorded(client, db, doc_project, fake_claude):
    owner, project = doc_project()
    data, v = _open(client, owner, project)
    ids = docx_blocks.block_ids(data)
    fake_claude({"reply": "Done.", "operations": [{"op": "set_text", "target": ids[1], "content": "THE REGISTRAR, having examined the records,"}]})
    r = client.post(f"/projects/{project.id}/document/chat", headers=owner["headers"], json={"version": v, "message": "use Registrar"})
    assert r.status_code == 200, r.text
    row = db.query(PendingLearning).one()
    assert row.origin == "chat" and "REGISTRAR" in row.after_text


def test_reject_learned_term_works_without_glossary_plan(client, db, doc_project, make_user):
    owner, project = doc_project()
    member = make_user(team=owner["team"])
    wallet_plan(db, owner["team"].id, "BASIC")
    term = Glossary(team_id=owner["team"].id, source_language="Italian", target_language="English",
                    source_term="Ufficiale d'Anagrafe", target_term="Registrar", origin="learned", learned_from_project_id=project.id)
    db.add(term)
    db.commit()

    assert client.delete(f"/glossary/{term.id}", headers=member["headers"]).status_code == 403
    assert client.post(f"/learning/terms/{term.id}/reject", headers=make_user()["headers"]).status_code == 404
    assert client.post(f"/learning/terms/{term.id}/reject", headers=member["headers"]).status_code == 200
    db.refresh(term)
    assert term.origin == "rejected"
    assert glossary_service.get_glossary(db, owner["team"].id, "Italian", "English") == []

    learning.upsert_terms(db, project, [{"source_term": "Ufficiale d'Anagrafe", "target_term": "Registrar", "confidence": 0.99}], "Registrar", "Ufficiale d'Anagrafe")
    db.commit()
    db.refresh(term)
    assert term.origin == "rejected"


def wallet_plan(db, team_id, plan):
    from app.models.credit import CreditWallet

    db.query(CreditWallet).filter(CreditWallet.team_id == team_id).update({CreditWallet.plan_type: plan})
    db.commit()


def test_project_learning_and_summary(client, db, doc_project, make_user):
    owner, project = doc_project()
    client.post(f"/projects/{project.id}/template", headers=owner["headers"])
    template = db.query(DocumentTemplate).one()
    project.template_id = template.id
    db.add_all([
        Glossary(team_id=owner["team"].id, source_language="it", target_language="en", source_term="Ufficio Anagrafe", target_term="Registry Office"),
        Glossary(team_id=owner["team"].id, source_language="Italian", target_language="English", source_term="Prefettura", target_term="Prefecture"),
        Glossary(team_id=owner["team"].id, source_language="Italian", target_language="English", source_term="Ufficiale d'Anagrafe",
                 target_term="Registrar", origin="learned", learned_from_project_id=project.id),
    ])
    db.commit()

    body = client.get(f"/projects/{project.id}/learning", headers=owner["headers"]).json()
    assert body["template_used"]["id"] == str(template.id)
    assert body["saved_as_template"]["id"] == str(template.id)
    assert body["doc_profile"]["doc_key"] == KEY
    assert {t["source_term"] for t in body["terms_applied"]} == {"Ufficio Anagrafe", "Ufficiale d'Anagrafe"}
    assert [t["target_term"] for t in body["terms_learned_here"]] == ["Registrar"]

    summary = client.get("/learning/summary", headers=owner["headers"]).json()
    assert summary["templates"] == 1 and summary["learned_terms"] == 1
    assert client.get("/learning/summary", headers=make_user()["headers"]).json()["templates"] == 0



def _team_term(db, owner, **kw):
    db.add(Glossary(team_id=owner["team"].id, source_language="Italian", target_language="English",
                    source_term="Ufficiale d'Anagrafe", target_term="Registrar", origin="learned", **kw))
    db.add(Glossary(team_id=owner["team"].id, source_language="Italian", target_language="English",
                    source_term="Prefettura", target_term="Prefecture"))
    db.commit()


def test_segment_translation_uses_learned_terms_with_auto_source(db, make_user, make_project, monkeypatch):
    from app.services import ai_translation_service as ats

    owner = make_user()
    project = make_project(owner)
    project.source_language = "auto"
    project.target_language = "en-GB"
    project.doc_profile = PROFILE
    db.commit()
    _team_term(db, owner)
    seen = {}
    monkeypatch.setattr(ats, "_call_model", lambda **kw: seen.update(kw) or "x<<<SEG>>>y")
    ats.translate_batch(["L'Ufficiale d'Anagrafe certifica", "altro"], "auto", "en-GB", db=db, project=project)
    assert "Ufficiale d'Anagrafe → Registrar" in seen["cache_prefix"]
    assert "Prefettura" not in seen["cache_prefix"]


def test_chat_edit_prompt_includes_team_terminology(client, db, doc_project, fake_claude):
    owner, project = doc_project()
    _team_term(db, owner)
    data, v = _open(client, owner, project)
    fake = fake_claude({"reply": "No change.", "operations": []})
    client.post(f"/projects/{project.id}/document/chat", headers=owner["headers"], json={"version": v, "message": "check terms"})
    texts = [b.get("text", "") for b in fake.requests[0]["messages"][0]["content"]]
    joined = "\n".join(texts)
    assert "TEAM TERMINOLOGY" in joined and "Ufficiale d'Anagrafe → Registrar" in joined and "Prefettura" not in joined


def test_authored_rebuild_prompt_includes_terminology(monkeypatch):
    from app.services import claude_authored_rebuild as car
    from app.services import claude_multiturn_rebuild as mt

    fake = _FakeClaude([{"x": 1}])
    monkeypatch.setattr(claude_params, "create_message", fake)
    monkeypatch.setattr(claude_params, "api_key", lambda: "k")
    car._call_claude_to_author(b"%PDF", "it", "en", "/tmp/out.docx", [], [], model="claude-opus-4-8",
                               terminology="TEAM TERMINOLOGY — use exactly these translations:\n- Comune → Municipality\n")
    assert "Comune → Municipality" in fake.requests[0]["messages"][0]["content"][1]["text"]

    seen = {}
    monkeypatch.setattr(mt, "author_rebuild_docx_multiturn", lambda *a, **kw: seen.update(kw) or b"docx")
    car.author_rebuild_docx(b"pdf", "it", "en", terminology="T")
    assert seen["terminology"] == "T"


def test_team_terminology_respects_apply_glossary(db, doc_project):
    owner, project = doc_project()
    _team_term(db, owner)
    assert "Registrar" in learning.team_terminology(db, project)
    project.apply_glossary = False
    assert learning.team_terminology(db, project) == ""
