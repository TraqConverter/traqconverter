"""Choosing the model that writes the translation: Claude by default, or GPT-4.1 with Claude still rebuilding the layout."""
import io
import uuid
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import httpx
import openai
import pytest
from docx import Document

from app.models.ai_usage import AiUsage
from app.models.batch import Batch
from app.models.credit import CreditWallet
from app.models.project import ProjectStatus, TranslationProject
from app.models.translation_segment import TranslationSegment
from app.services import ai_translation_service as ats
from app.services import authored_translation, learning
from tests.conftest import make_pdf
from tests.test_ai_metering import hooks  # noqa: F401
from tests.test_plan_gating import (  # noqa: F401
    MEMORY,
    PROFILE,
    SOURCE,
    _add_glossary,
    _add_memory,
    gated_project,
    no_claude,
)

GPT = "gpt-4.1"


def _rebuilt_copy() -> bytes:
    """What Claude's same-language rebuild returns: body text, a table, a header, a tab-aligned line and a date."""
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = SOURCE[2]
    doc.add_paragraph(SOURCE[0])
    p = doc.add_paragraph()
    p.add_run("Nato a ").bold = True
    p.add_run("Bari")
    doc.add_paragraph("Data:\tBari")
    doc.add_paragraph("12/03/1987")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = SOURCE[1]
    table.rows[0].cells[1].text = "[Firma]"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _fake_model(calls):
    def call_model(**kw):
        calls.append(kw)
        user = kw["user"]
        if "INPUT:\n" not in user:
            return f"EN {user}"
        segs = user.split("INPUT:\n", 1)[1].split("\n<<<SEG>>>\n")
        return "<<<SEG>>>".join(f"EN {s}" for s in segs)

    return call_model


def _texts(data: bytes) -> dict:
    doc = Document(io.BytesIO(data))
    return {
        "body": [p.text for p in doc.paragraphs],
        "cells": [c.text for c in doc.tables[0].rows[0].cells],
        "header": doc.sections[0].header.paragraphs[0].text,
    }


def _run(db, project, monkeypatch, model, kind="PDF"):
    from app.services import claude_authored_rebuild as car
    from app.services import translation_processor as tp

    seen = {"author": [], "calls": [], "uploads": {}}
    monkeypatch.setattr(tp, "download_file_from_s3", lambda key, dest: Path(dest).write_bytes(b"%PDF"))
    monkeypatch.setattr(tp, "extract_segments", lambda path: (kind, [SimpleNamespace(text=s, layout={}) for s in SOURCE]))
    monkeypatch.setattr(learning, "classify_document", lambda data, name, hint="": dict(PROFILE))
    monkeypatch.setattr(ats, "_call_model", _fake_model(seen["calls"]))

    def rebuild(**kw):
        Path(kw["output_path"]).write_bytes(b"docx")
        return kw["output_path"]

    def upload(path):
        key = f"uploads/{uuid.uuid4()}_{Path(path).name}"
        seen["uploads"][key] = Path(path).read_bytes()
        return key

    monkeypatch.setattr(tp, "rebuild_output", rebuild)
    monkeypatch.setattr(tp, "upload_file_to_s3", upload)
    monkeypatch.setattr(car, "author_rebuild_docx", lambda **kw: seen["author"].append(kw) or _rebuilt_copy())
    project.status = ProjectStatus.PENDING
    project.model = model
    project.source_language = "Italian"
    project.doc_key, project.doc_profile = None, None
    db.commit()
    tp.process_translation_job(str(project.id))
    db.expire_all()
    return seen


# Pipeline

def test_gpt_pdf_keeps_claudes_layout_and_gpt_writes_every_word(db, gated_project, monkeypatch):
    owner, project = gated_project("PRO")
    _add_glossary(db, owner["team"].id)
    _add_memory(db, owner["team"].id)
    project.ai_instructions = "Use British spelling"
    db.commit()
    seen = _run(db, project, monkeypatch, GPT)

    [author] = seen["author"]
    assert author["reproduce"] is True
    assert author["source_lang"] == author["target_lang"] == "Italian"
    assert not author.get("terminology") and not author.get("instructions")

    assert project.status == ProjectStatus.COMPLETED and project.rebuild_error is None
    texts = _texts(seen["uploads"][project.authored_docx_s3_key])
    # Translation memory first, then the glossary term inside GPT's translation.
    assert texts["body"][0] == MEMORY
    assert texts["header"] == "EN Ufficio Anagrafe e Civil Registry"
    assert texts["body"][1] == "EN Nato a Bari"
    assert texts["body"][2] == "EN Data:\tEN Bari"
    assert texts["body"][3] == "12/03/1987"
    assert texts["cells"] == [f"EN {SOURCE[1]}", "EN [Firma]"]

    assert seen["calls"] and all(c["model_key"] == GPT for c in seen["calls"])
    prompts = [c["cache_prefix"] + c["system"] + c["user"] for c in seen["calls"]]
    assert all("Use British spelling" in p for p in prompts)
    assert any("Civil Registry" in p for p in prompts)


def test_gpt_translates_the_segments_too(db, gated_project, monkeypatch):
    _, project = gated_project("PRO")
    _run(db, project, monkeypatch, GPT)
    segs = db.query(TranslationSegment).filter(TranslationSegment.project_id == project.id).order_by(TranslationSegment.segment_index).all()
    assert [s.translated_text for s in segs] == [f"EN {s}" for s in SOURCE]


def test_claude_pdf_is_unchanged(db, gated_project, monkeypatch):
    owner, project = gated_project("PRO")
    _add_glossary(db, owner["team"].id)
    project.ai_instructions = "Use British spelling"
    db.commit()
    seen = _run(db, project, monkeypatch, "claude-authored")

    [author] = seen["author"]
    assert author["reproduce"] is False and author["target_lang"] == project.target_language
    assert "Civil Registry" in author["terminology"] and author["instructions"] == "Use British spelling"
    # Claude's own translation is the document; nothing translates it again.
    assert _texts(seen["uploads"][project.authored_docx_s3_key])["body"][0] == SOURCE[0]
    assert all(c["model_key"] == "claude-authored" for c in seen["calls"])


@pytest.mark.parametrize("model", [GPT, "claude-authored"])
def test_word_files_have_no_layout_rebuild_and_use_the_chosen_model(db, gated_project, monkeypatch, model):
    _, project = gated_project("PRO")
    project.authored_docx_s3_key = None
    db.commit()
    seen = _run(db, project, monkeypatch, model, kind="DOCX")
    assert seen["author"] == [] and project.authored_docx_s3_key is None
    assert seen["calls"] and all(c["model_key"] == model for c in seen["calls"])


def test_a_failed_gpt_translation_falls_back_to_the_segment_layout(db, gated_project, monkeypatch):
    _, project = gated_project("PRO")
    real = authored_translation.translate_docx

    def no_text(data, translate):
        return real(data, lambda texts: [""] * len(texts))

    monkeypatch.setattr(authored_translation, "translate_docx", no_text)
    _run(db, project, monkeypatch, GPT)
    assert project.status == ProjectStatus.COMPLETED
    assert project.authored_docx_s3_key is None and project.rebuild_error


def test_batch_terms_reach_gpt(db, gated_project, monkeypatch):
    from app.services import batch_terms

    owner, project = gated_project("PRO")
    batch = Batch(team_id=owner["team"].id, name="Rossi family")
    db.add(batch)
    db.flush()
    project.batch_id = batch.id
    db.commit()
    monkeypatch.setattr(batch_terms, "terminology_block", lambda db, project, text=None: "BATCH TERMS: Rossi -> Rossi")
    monkeypatch.setattr(batch_terms, "record_from_segments", lambda *a, **kw: 0)
    seen = _run(db, project, monkeypatch, GPT)
    assert seen["calls"] and all("BATCH TERMS" in c["cache_prefix"] + c["system"] for c in seen["calls"])


# Regenerate

def test_regenerate_on_a_gpt_project_rebuilds_with_claude_and_translates_with_gpt(db, gated_project, monkeypatch):
    from app.services import ai_actions
    from app.services import claude_multiturn_rebuild as mt

    _, project = gated_project("PRO")
    project.model = GPT
    project.source_language = "Italian"
    project.rebuild_status = "running"
    db.commit()
    rebuilds, calls = [], []
    monkeypatch.setattr(mt, "author_rebuild_docx_multiturn", lambda *a, **kw: rebuilds.append((a, kw)) or _rebuilt_copy())
    monkeypatch.setattr(ats, "_call_model", _fake_model(calls))
    ai_actions.run_rebuild(str(project.id), "Call him Signor Bianchi")
    db.expire_all()

    [(args, kw)] = rebuilds
    assert args[1:] == ("Italian", "Italian") and kw["reproduce"] is True
    assert kw["extra_instructions"] == "Call him Signor Bianchi"
    assert project.rebuild_status == "done"
    assert calls and all(c["model_key"] == GPT and "Call him Signor Bianchi" in c["system"] + c["user"] for c in calls)


# Upload and listing

def _upload(client, owner, tmp_path, **data):
    pdf = make_pdf(tmp_path / "doc.pdf", pages=3)
    with open(pdf, "rb") as f:
        return client.post(
            "/projects/upload",
            headers=owner["headers"],
            files={"file": ("doc.pdf", f, "application/pdf")},
            data={"source_language": "Italian", "target_language": "English", **data},
        )


def test_both_choices_cost_the_same(client, db, make_user, tmp_path):
    for model in ("claude-authored", GPT):
        owner = make_user(credits=10)
        r = _upload(client, owner, tmp_path, model=model)
        assert r.status_code == 200, r.text
        assert r.json()["credits_used"] == 3
        assert db.get(TranslationProject, r.json()["project_id"]).model == model
        db.expire_all()
        assert db.query(CreditWallet).filter(CreditWallet.team_id == owner["team"].id).one().subscription_credits == 7
    assert _upload(client, make_user(credits=10), tmp_path, model="gpt-9").status_code == 400


def test_the_listing_offers_claude_then_gpt(client, make_user):
    owner = make_user()
    models = client.get("/projects/translation-models", headers=owner["headers"]).json()["models"]
    assert [m["id"] for m in models] == ["claude-authored", GPT]
    assert models[0]["recommended"] and models[1]["provider"] == "openai"


# Metering

def test_gpt_translation_is_metered_against_the_project(hooks, db, make_user, make_project, monkeypatch):  # noqa: F811
    from app.services import ai_usage

    def handler(request):
        return httpx.Response(200, json={
            "id": "chatcmpl-1",
            "object": "chat.completion",
            "created": 0,
            "model": "gpt-4.1-2025-04-14",
            "choices": [{"index": 0, "message": {"role": "assistant", "content": "Hello"}, "finish_reason": "stop"}],
            "usage": {"prompt_tokens": 1000, "completion_tokens": 100, "total_tokens": 1100},
        })

    monkeypatch.setattr(ats, "_openai_client", openai.OpenAI(api_key="test", http_client=httpx.Client(transport=httpx.MockTransport(handler))))
    owner = make_user()
    project = make_project(owner)
    project.model = GPT
    db.commit()
    with ai_usage.ai_context(action="translation", project_id=project.id):
        assert ats.translate_text("Ciao", "Italian", "English", project=project) == "Hello"
    db.expire_all()
    [row] = db.query(AiUsage).all()
    assert row.model == "gpt-4.1-2025-04-14" and row.project_id == project.id and row.action == "translation"
    assert row.usd == Decimal("0.002800")


# The DOCX text swap

def test_translate_docx_sends_each_passage_once_and_skips_numbers():
    sent = []

    def translate(texts):
        sent.append(list(texts))
        return [t.upper() for t in texts]

    doc = Document()
    for text in ("Bari", "Bari", "  Roma  ", "2024", ""):
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    out = authored_translation.translate_docx(buf.getvalue(), translate)
    assert sent == [["Bari", "Roma"]]
    assert [p.text for p in Document(io.BytesIO(out)).paragraphs] == ["BARI", "BARI", "  ROMA  ", "2024", ""]


def test_translate_docx_refuses_to_leave_source_text():
    with pytest.raises(authored_translation.UntranslatedText):
        authored_translation.translate_docx(_rebuilt_copy(), lambda texts: ["x"] + [""] * (len(texts) - 1))


def test_only_openai_models_need_a_separate_translation():
    assert authored_translation.needs_separate_translation(SimpleNamespace(model=GPT))
    assert authored_translation.needs_separate_translation(SimpleNamespace(model="gpt-4.1-mini"))
    for model in ("claude-authored", "balanced", None, "claude-sonnet-4-6"):
        assert not authored_translation.needs_separate_translation(SimpleNamespace(model=model))

