import io
import json
import uuid
from types import SimpleNamespace

import fitz
import pytest
from docx import Document
from PIL import Image

from app.models.glossary import Glossary
from app.models.project import ProjectStatus, TranslationProject
from app.models.translation_segment import TranslationSegment
from app.services import claude_params, docx_blocks, review_checks, source_map

SRC_LINES = [
    ("CERTIFICATO DI RESIDENZA", "CERTIFICATE OF RESIDENCE"),
    ("che il Sig. BIANCHI LUCA, nato a Bari il 12/03/1987,", "that Mr. BIANCHI LUCA, born in Bari on 12/03/1987,"),
    ("codice fiscale BNCLCU87C12A662X, importo 1.000,00 EUR", "tax code BNCLCU87C12A662X, amount 1,000.00 EUR"),
    ("Esempio, lì 14/09/2026, protocollo n. 5582", "Esempio, on 14/09/2026, protocol no. 5582"),
    ("L'Ufficiale d'Anagrafe", "The Registrar"),
]
GOOD_TRANSLATION = [t for _, t in SRC_LINES]


def _docx(paras) -> bytes:
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _source_pdf(text_layer=True) -> tuple[bytes, list[dict]]:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    layouts = []
    for i, (src, _) in enumerate(SRC_LINES):
        y = 100 + i * 40
        if text_layer:
            page.insert_text((72, y), src, fontsize=11)
        layouts.append({"kind": "pdf_line", "page": 0, "bbox": [72, y - 11, 72 + 5.5 * len(src), y + 3]})
    data = doc.tobytes()
    doc.close()
    return data, layouts


def _png(width=800, height=1100) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (250, 248, 242)).save(buf, "PNG")
    return buf.getvalue()


@pytest.fixture(autouse=True)
def page_cache(tmp_path, monkeypatch):
    monkeypatch.setenv("SOURCE_PAGE_CACHE", str(tmp_path / "pages"))


@pytest.fixture()
def review_project(db, storage, make_user):
    def _make(translation=GOOD_TRANSLATION, plan="PRO", source="pdf", owner=None, boxes=True):
        owner = owner or make_user(plan=plan)
        if source == "pdf":
            data, layouts = _source_pdf()
            name = "certificate.pdf"
        elif source == "scan":
            data, layouts = _source_pdf(text_layer=False)
            name = "certificate.pdf"
        elif source == "image":
            data, layouts = _png(), [{"kind": "image_line", "page": 0, "bbox": [50, 80 + 60 * i, 700, 120 + 60 * i]} for i in range(len(SRC_LINES))]
            name = "certificate.png"
        else:
            data, layouts = _docx([s for s, _ in SRC_LINES]), [{"kind": "docx_paragraph"} for _ in SRC_LINES]
            name = "certificate.docx"
        key = f"uploads/{uuid.uuid4()}_{name}"
        storage["objects"][key] = data
        doc_key = f"uploads/{uuid.uuid4()}_authored.docx"
        storage["objects"][doc_key] = _docx(translation)
        project = TranslationProject(
            user_id=owner["user"].id,
            team_id=owner["team"].id,
            file_name=name,
            file_path=key,
            page_count=1,
            credits_used=1,
            status=ProjectStatus.COMPLETED,
            source_language="Italian",
            target_language="English",
            model="claude-sonnet-4-6",
            authored_docx_s3_key=doc_key,
        )
        db.add(project)
        db.flush()
        for i, ((src, tgt), layout) in enumerate(zip(SRC_LINES, layouts)):
            db.add(TranslationSegment(
                project_id=project.id,
                segment_index=i,
                source_text=src,
                translated_text=tgt,
                layout_meta=layout if boxes == True else (
                    {"kind": "pdf_claude_line", "page": 0, "ocr_source": "claude", "bbox": [10, 10 + 20 * i, 900, 30 + 20 * i], "page_width_pt": 595, "page_height_pt": 842}
                    if boxes == "claude" else {"kind": "pdf_claude_line", "page": 0}
                ),
            ))
        db.commit()
        return owner, project

    return _make


class _Fake:
    """Answers each model call by the structured-output schema it asks for."""

    def __init__(self, names=None, untranslated=None, vision=None):
        self.answers = {"names": names, "untranslated": untranslated, "vision": vision}
        self.calls = []

    def __call__(self, client, **kwargs):
        schema = kwargs["output_config"]["format"]["schema"]
        kind = (
            "names" if schema is review_checks.NAMES_SCHEMA
            else "untranslated" if schema is review_checks.UNTRANSLATED_SCHEMA
            else "vision" if schema is source_map.VISION_SCHEMA
            else "other"
        )
        self.calls.append((kind, kwargs))
        body = self.answers.get(kind)
        if callable(body):
            body = body(kwargs)
        if body is None:
            body = {"names": []} if kind == "names" else {"results": []} if kind == "untranslated" else {"paragraphs": [], "elements": []}
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=json.dumps(body))],
            stop_reason="end_turn",
            usage=SimpleNamespace(input_tokens=100, output_tokens=50),
            model=kwargs["model"],
        )

    def kinds(self):
        return [k for k, _ in self.calls]


@pytest.fixture()
def fake(monkeypatch):
    def install(**answers):
        f = _Fake(**answers)
        monkeypatch.setattr(claude_params, "create_message", f)
        monkeypatch.setattr(claude_params, "api_key", lambda: "test-key")
        return f

    return install


def _doc(client, owner, project):
    r = client.get(f"/projects/{project.id}/document", headers=owner["headers"])
    assert r.status_code == 200, r.text
    return r.content, int(r.headers["X-Document-Version"])


def _ids_by_text(data):
    return {p["text"]: p["id"] for p in source_map.document_paragraphs(data)}


def _checks(client, owner, project, refresh=False):
    r = client.get(f"/projects/{project.id}/checks" + ("?refresh=1" if refresh else ""), headers=owner["headers"])
    assert r.status_code == 200, r.text
    return r.json()


def _kinds(result, severity=None):
    return [(i["kind"], i["severity"]) for i in result["items"] if severity is None or i["severity"] == severity]


def test_source_pages_render_with_auth_and_tenant_isolation(client, review_project, make_user):
    owner, project = review_project()
    url = f"/projects/{project.id}/source/pages"
    info = client.get(url, headers=owner["headers"]).json()
    assert info["available"] and info["count"] == 1
    assert info["pages"][0] == {"width": 595.0, "height": 842.0}

    png = client.get(f"{url}/0.png?scale=1", headers=owner["headers"])
    assert png.status_code == 200 and png.headers["content-type"] == "image/png"
    img = Image.open(io.BytesIO(png.content))
    assert img.size == (595, 842)
    assert client.get(f"{url}/0.png?scale=2", headers=owner["headers"]).status_code == 200
    assert client.get(f"{url}/1.png", headers=owner["headers"]).status_code == 404

    other = make_user()
    assert client.get(url, headers=other["headers"]).status_code == 404
    assert client.get(f"{url}/0.png", headers=other["headers"]).status_code == 404
    assert client.get(f"/projects/{project.id}/source/map", headers=other["headers"]).status_code == 404
    assert client.get(f"/projects/{project.id}/checks", headers=other["headers"]).status_code == 404
    assert client.get(f"{url}/0.png").status_code in (401, 403)


def test_image_source_is_one_page_in_pixels_and_docx_without_converter_is_unavailable(client, review_project, monkeypatch):
    owner, project = review_project(source="image")
    info = client.get(f"/projects/{project.id}/source/pages", headers=owner["headers"]).json()
    assert info["count"] == 1 and info["pages"][0] == {"width": 800.0, "height": 1100.0}

    import app.services.export_service as export_service

    monkeypatch.setattr(export_service, "_convert_docx_to_pdf", lambda data: None)
    owner2, project2 = review_project(source="docx")
    info2 = client.get(f"/projects/{project2.id}/source/pages", headers=owner2["headers"]).json()
    assert info2 == {"available": False, "count": 0, "pages": [], "kind": None}
    view = client.get(f"/projects/{project2.id}/source/map", headers=owner2["headers"]).json()
    assert view["status"] == "unavailable" and view["blocks"] == {}


def test_map_from_segments_on_a_digital_pdf_needs_no_model_call(client, review_project, fake):
    f = fake()
    owner, project = review_project()
    data, _ = _doc(client, owner, project)
    ids = _ids_by_text(data)
    view = client.get(f"/projects/{project.id}/source/map", headers=owner["headers"]).json()
    assert view["status"] == "ready"
    assert "vision" not in f.kinds()
    box = view["blocks"][ids["that Mr. BIANCHI LUCA, born in Bari on 12/03/1987,"]]
    assert box["page"] == 0 and box["confidence"] >= 0.75
    assert box["bbox"][1] == pytest.approx((140 - 11) / 842, abs=0.01)
    assert box["bbox"][0] == pytest.approx(72 / 595, abs=0.01)
    title = view["blocks"][ids["CERTIFICATE OF RESIDENCE"]]
    assert title["bbox"][1] < box["bbox"][1]


def test_claude_ocr_boxes_are_replaced_by_the_text_layer(client, review_project, fake):
    f = fake()
    owner, project = review_project(boxes="claude")
    data, _ = _doc(client, owner, project)
    ids = _ids_by_text(data)
    view = client.get(f"/projects/{project.id}/source/map", headers=owner["headers"]).json()
    assert "vision" not in f.kinds()
    box = view["blocks"][ids["that Mr. BIANCHI LUCA, born in Bari on 12/03/1987,"]]["bbox"]
    line = fitz.open(stream=_source_pdf()[0], filetype="pdf")[0].search_for("BIANCHI LUCA")[0]
    assert box[1] <= line.y0 / 842 <= box[3]
    assert box[0] == pytest.approx(72 / 595, abs=0.01)


def test_scan_runs_one_vision_call_per_page_and_surfaces_uncertain_readings(client, review_project, fake):
    def vision(kwargs):
        listing = kwargs["messages"][0]["content"][1]["text"]
        rows = [line.split(": ", 1) for line in listing.split("\n") if line.startswith("_b")]
        out = []
        for i, (bid, text) in enumerate(rows):
            low = "5582" in text
            out.append({"id": bid, "found": True, "x0": 60, "y0": 100 + 50 * i, "x1": 700, "y1": 140 + 50 * i,
                        "reading": "low" if low else "high", "reason": "digits after 55 smudged" if low else ""})
        return {"paragraphs": out, "elements": [
            {"kind": "stamp", "x0": 300, "y0": 700, "x1": 500, "y1": 900, "text": "COMUNE DI ESEMPIO", "reading": "medium", "block_id": ""},
        ]}

    f = fake(vision=vision)
    owner, project = review_project(source="scan", boxes=False)
    data, _ = _doc(client, owner, project)
    ids = _ids_by_text(data)

    first = client.get(f"/projects/{project.id}/source/map", headers=owner["headers"]).json()
    assert first["status"] == "pending"
    assert f.kinds().count("vision") == 1
    image = f.calls[f.kinds().index("vision")][1]["messages"][0]["content"][0]
    assert image["type"] == "image" and image["source"]["media_type"] == "image/jpeg"

    view = client.get(f"/projects/{project.id}/source/map", headers=owner["headers"]).json()
    assert view["status"] == "ready"
    assert f.kinds().count("vision") == 1
    protocol = ids["Esempio, on 14/09/2026, protocol no. 5582"]
    assert view["blocks"][protocol]["reading"] == "low"
    assert view["blocks"][protocol]["reason"] == "digits after 55 smudged"
    assert len(view["elements"]) == 1 and view["elements"][0]["kind"] == "stamp"

    result = _checks(client, owner, project)
    uncertain = [i for i in result["items"] if i["kind"] == "uncertain"]
    assert len(uncertain) == 1 and uncertain[0]["severity"] == "warning"
    assert uncertain[0]["block_id"] == protocol and uncertain[0]["source"]["page"] == 0
    notation = [i for i in result["items"] if i["kind"] == "notation"]
    assert notation and notation[0]["severity"] == "warning" and notation[0]["source"]["bbox"]
    assert f.kinds().count("vision") == 1


def test_dense_scan_page_lists_every_paragraph_to_the_vision_pass(client, review_project, fake):
    listed = []

    def vision(kwargs):
        listing = kwargs["messages"][0]["content"][1]["text"]
        rows = [line.split(": ", 1)[0] for line in listing.split("\n") if line.startswith("_b")]
        listed.append(rows)
        return {"paragraphs": [{"id": b, "found": True, "x0": 10, "y0": 10, "x1": 400, "y1": 60, "reading": "high", "reason": ""} for b in rows], "elements": []}

    fake(vision=vision)
    translation = [f"Row {i} of the register, entry number {i}" for i in range(170)]
    owner, project = review_project(translation=translation, source="scan", boxes=False)
    data, _ = _doc(client, owner, project)
    client.get(f"/projects/{project.id}/source/map", headers=owner["headers"])
    view = client.get(f"/projects/{project.id}/source/map", headers=owner["headers"]).json()
    assert all(len(rows) <= source_map.MAX_PARAS_PER_CALL for rows in listed)
    assert {b for rows in listed for b in rows} == set(_ids_by_text(data).values())
    assert set(view["blocks"]) == set(_ids_by_text(data).values())


def test_new_paragraph_after_edit_is_mapped_incrementally(client, db, review_project, fake):
    calls = []

    def vision(kwargs):
        listing = kwargs["messages"][0]["content"][1]["text"]
        rows = [line.split(": ", 1)[0] for line in listing.split("\n") if line.startswith("_b")]
        calls.append(rows)
        return {"paragraphs": [{"id": b, "found": True, "x0": 10, "y0": 10, "x1": 400, "y1": 60, "reading": "high", "reason": ""} for b in rows], "elements": []}

    fake(vision=vision)
    owner, project = review_project()
    data, v = _doc(client, owner, project)
    client.get(f"/projects/{project.id}/source/map", headers=owner["headers"])
    assert calls == []

    ops = [{"op": "insert_after", "target": docx_blocks.block_ids(data)[-1], "content": "<w:p><w:r><w:t>Unrelated closing words</w:t></w:r></w:p>"}]
    new_data, _ = docx_blocks.apply_operations(data, ops)
    from app.services import document_editor

    db.refresh(project)
    document_editor.save_version(db, project, new_data, "test", None)
    db.commit()
    view = client.get(f"/projects/{project.id}/source/map", headers=owner["headers"]).json()
    new_id = _ids_by_text(new_data)["Unrelated closing words"]
    assert calls == [[new_id]]
    view = client.get(f"/projects/{project.id}/source/map", headers=owner["headers"]).json()
    assert new_id in view["blocks"] and len(calls) == 1


def test_clean_translation_is_ready_with_formatting_notes_only(client, review_project, fake):
    fake(names={"names": [{"source": "BIANCHI LUCA", "kind": "person", "renderings": ["BIANCHI LUCA"]},
                          {"source": "Bari", "kind": "place", "renderings": ["Bari"]}]})
    owner, project = review_project(plan="BASIC")
    result = _checks(client, owner, project)
    assert result["ready"], result["items"]
    assert result["document_version"] == 1
    messages = [i["message"] for i in result["items"]]
    assert "1.000,00 is formatted as 1,000.00" in messages
    assert all(i["severity"] == "info" for i in result["items"])


def test_missing_number_name_code_and_formatted_date(client, review_project, fake):
    fake(names={"names": [{"source": "BIANCHI LUCA", "kind": "person", "renderings": ["BIANCHI LUCA"]}]})
    translation = [
        "CERTIFICATE OF RESIDENCE",
        "that Mr. ROSSI MARCO, born in Bari on 12 March 1987,",
        "tax code BNCLCU87C12A662X, amount 1,000.00 EUR",
        "Esempio, on 14/09/2026, protocol no.",
        "The Registrar",
    ]
    owner, project = review_project(translation=translation, plan="BASIC")
    data, _ = _doc(client, owner, project)
    ids = _ids_by_text(data)
    result = _checks(client, owner, project)
    assert not result["ready"]
    by_msg = {i["message"]: i for i in result["items"]}
    missing = by_msg["Number 5582 from the source isn't in the translation"]
    assert missing["severity"] == "error"
    assert missing["block_id"] == ids["Esempio, on 14/09/2026, protocol no."]
    assert missing["source"]["page"] == 0 and len(missing["source"]["bbox"]) == 4
    assert by_msg["Name BIANCHI LUCA isn't in the translation"]["severity"] == "error"
    assert by_msg["Date 12/03/1987 is written as 12 march 1987"]["severity"] == "info"


def test_missing_date_and_illegible_downgrade(client, review_project, fake):
    fake()
    translation = list(GOOD_TRANSLATION)
    translation[1] = "that Mr. BIANCHI LUCA, born in Bari,"
    translation[3] = "Esempio, on 14/09/2026, protocol no. 55[illegible]"
    owner, project = review_project(translation=translation, plan="BASIC")
    result = _checks(client, owner, project)
    by_kind = {(i["kind"], i["severity"]): i for i in result["items"]}
    assert by_kind[("date", "error")]["message"] == "Date 12/03/1987 from the source isn't in the translation"
    warn = by_kind[("number", "warning")]
    assert "5582" in warn["message"] and "illegible" in warn["message"]
    assert ("uncertain", "warning") in by_kind


def test_glossary_violation(client, db, review_project, fake):
    fake()
    owner, project = review_project(plan="PRO")
    db.add(Glossary(team_id=owner["team"].id, source_language="Italian", target_language="English",
                    source_term="Ufficiale d'Anagrafe", target_term="Civil Registrar", origin="manual"))
    db.commit()
    result = _checks(client, owner, project)
    item = next(i for i in result["items"] if i["kind"] == "glossary")
    assert item["severity"] == "warning"
    assert "Civil Registrar" in item["message"]
    assert item["source"]["page"] == 0
    assert item["block_id"]


def test_untranslated_paragraph_is_confirmed_once_and_cached(client, review_project, fake):
    def verdict(kwargs):
        text = kwargs["messages"][0]["content"]
        rows = [line.split(": ", 1)[0] for line in text.split("\n") if line.startswith("_b")]
        return {"results": [{"id": b, "untranslated": True} for b in rows]}

    f = fake(untranslated=verdict)
    translation = list(GOOD_TRANSLATION)
    translation[1] = "che il Sig. BIANCHI LUCA, nato a Bari il 12/03/1987,"
    owner, project = review_project(translation=translation, plan="BASIC")
    result = _checks(client, owner, project)
    item = next(i for i in result["items"] if i["kind"] == "untranslated")
    assert item["severity"] == "error" and item["block_id"]
    assert f.kinds().count("untranslated") == 1
    _checks(client, owner, project, refresh=True)
    assert f.kinds().count("untranslated") == 1
    assert f.kinds().count("names") == 1


def test_missing_notation_from_ocr_placeholders(client, db, review_project, fake):
    fake()
    owner, project = review_project(plan="BASIC")
    db.add(TranslationSegment(project_id=project.id, segment_index=10, source_text="[Stamp: COMUNE DI ESEMPIO]", translated_text="[Stamp]",
                              layout_meta={"kind": "pdf_line", "page": 0, "bbox": [300, 600, 450, 700], "claude_kind": "placeholder", "placeholder_kind": "stamp"}))
    db.add(TranslationSegment(project_id=project.id, segment_index=11, source_text="[Signature]", translated_text="[Signature]",
                              layout_meta={"kind": "pdf_line", "page": 0, "bbox": [400, 700, 550, 760], "claude_kind": "placeholder", "placeholder_kind": "signature"}))
    db.commit()
    result = _checks(client, owner, project)
    notes = [i for i in result["items"] if i["kind"] == "notation"]
    assert len(notes) == 2 and all(i["severity"] == "warning" for i in notes)

    owner2, project2 = review_project(translation=GOOD_TRANSLATION + ["[Round stamp: Municipality of Esempio]", "[Signature]"], plan="BASIC")
    db.add(TranslationSegment(project_id=project2.id, segment_index=10, source_text="[Stamp]", translated_text="[Stamp]",
                              layout_meta={"kind": "pdf_line", "page": 0, "bbox": [300, 600, 450, 700], "claude_kind": "placeholder", "placeholder_kind": "stamp"}))
    db.commit()
    assert not [i for i in _checks(client, owner2, project2)["items"] if i["kind"] == "notation"]


def test_certification_page_checked_only_on_plans_with_certifications(client, review_project, fake):
    fake()
    owner, project = review_project(plan="PRO")
    result = _checks(client, owner, project)
    cert = [i for i in result["items"] if i["kind"] == "certification"]
    assert len(cert) == 1 and cert[0]["severity"] == "warning" and not result["ready"]

    data, v = _doc(client, owner, project)
    r = client.post(f"/projects/{project.id}/document/certification", headers=owner["headers"], json={"version": v})
    assert r.status_code == 200, r.text
    result = _checks(client, owner, project)
    assert result["document_version"] == v + 1
    cert = [i for i in result["items"] if i["kind"] == "certification"]
    assert not cert or cert[0]["severity"] == "error"

    owner_b, project_b = review_project(plan="BASIC")
    assert not [i for i in _checks(client, owner_b, project_b)["items"] if i["kind"] == "certification"]


def test_empty_certification_fields_are_an_error(client, review_project, fake):
    fake()
    owner, project = review_project(plan="PRO")
    data, v = _doc(client, owner, project)
    v = client.post(f"/projects/{project.id}/document/certification", headers=owner["headers"], json={"version": v}).json()["version"]
    r = client.put(f"/projects/{project.id}/document/certification", headers=owner["headers"], json={"version": v, "fields": {"translator": ""}})
    assert r.status_code == 200, r.text
    cert = [i for i in _checks(client, owner, project)["items"] if i["kind"] == "certification"]
    assert len(cert) == 1 and cert[0]["severity"] == "error" and "translator" in cert[0]["message"]
    assert cert[0]["block_id"]


def test_checks_are_cached_per_version_and_recomputed_on_edit_or_refresh(client, review_project, fake, monkeypatch):
    fake()
    runs = []
    real = review_checks.compute
    monkeypatch.setattr(review_checks, "compute", lambda *a, **k: runs.append(1) or real(*a, **k))
    owner, project = review_project(plan="BASIC")
    first = _checks(client, owner, project)
    second = _checks(client, owner, project)
    assert len(runs) == 1 and first["checked_at"] == second["checked_at"]

    data, v = _doc(client, owner, project)
    bid = _ids_by_text(data)["The Registrar"]
    r = client.post(f"/projects/{project.id}/document/edits", headers=owner["headers"], json={"version": v, "edits": [{"block_id": bid, "text": "The Registrar."}]})
    assert r.status_code == 200
    third = _checks(client, owner, project)
    assert len(runs) == 2 and third["document_version"] == v + 1
    _checks(client, owner, project, refresh=True)
    assert len(runs) == 3


def test_checks_cache_survives_a_version_number_reused_after_undo(client, review_project, fake, monkeypatch):
    fake()
    runs = []
    real = review_checks.compute
    monkeypatch.setattr(review_checks, "compute", lambda *a, **k: runs.append(1) or real(*a, **k))
    owner, project = review_project(plan="BASIC")
    url = f"/projects/{project.id}/document"
    data, v = _doc(client, owner, project)
    bid = _ids_by_text(data)["The Registrar"]
    client.post(f"{url}/edits", headers=owner["headers"], json={"version": v, "edits": [{"block_id": bid, "text": "The Registrar."}]})
    _checks(client, owner, project)
    client.post(f"{url}/undo", headers=owner["headers"], json={})
    r = client.post(f"{url}/edits", headers=owner["headers"], json={"version": v, "edits": [{"block_id": bid, "text": "Registrar"}]})
    assert r.json()["version"] == v + 1
    _checks(client, owner, project)
    assert len(runs) == 2


def test_dismiss_marks_item_and_can_be_restored(client, review_project, fake, make_user):
    fake()
    translation = list(GOOD_TRANSLATION)
    translation[3] = "Esempio, on 14/09/2026, protocol no."
    owner, project = review_project(translation=translation, plan="BASIC")
    result = _checks(client, owner, project)
    assert not result["ready"]
    errors = [i for i in result["items"] if i["severity"] in ("error", "warning")]
    assert len(errors) == 1
    item_id = errors[0]["id"]

    other = make_user()
    assert client.post(f"/projects/{project.id}/checks/{item_id}/dismiss", headers=other["headers"]).status_code == 404
    r = client.post(f"/projects/{project.id}/checks/{item_id}/dismiss", headers=owner["headers"])
    assert r.status_code == 200 and r.json() == {"id": item_id, "dismissed": True, "document_version": 1}
    after = _checks(client, owner, project)
    assert after["ready"] and next(i for i in after["items"] if i["id"] == item_id)["dismissed"]
    assert client.post(f"/projects/{project.id}/checks/not-an-id/dismiss", headers=owner["headers"]).status_code == 422

    client.delete(f"/projects/{project.id}/checks/{item_id}/dismiss", headers=owner["headers"])
    assert not _checks(client, owner, project)["ready"]


def test_segment_matching_prefers_reading_order_and_skips_weak_matches():
    paras = [{"id": "_b00000001", "text": "Number"}, {"id": "_b00000002", "text": "Something unrelated entirely"}]
    segs = [
        {"index": 0, "page": 0, "bbox": [0.1, 0.1, 0.3, 0.12], "src": "Numero", "tgt": "Number", "toks": {"numero", "number"}, "base": {"number"}},
        {"index": 1, "page": 0, "bbox": [0.1, 0.5, 0.3, 0.52], "src": "Numero", "tgt": "Number", "toks": {"numero", "number"}, "base": {"number"}},
    ]
    out = source_map.match_segments(paras, segs)
    assert out["_b00000001"]["bbox"][1] == 0.1
    assert "_b00000002" not in out


def test_model_boxes_snap_to_stamp_ink_and_text_rows():
    from PIL import ImageDraw

    img = Image.new("RGB", (800, 1000), (250, 248, 242))
    d = ImageDraw.Draw(img)
    d.ellipse((500, 600, 700, 800), outline=(40, 60, 170), width=8)
    d.rectangle((100, 300, 500, 318), fill=(20, 20, 20))
    buf = io.BytesIO()
    img.save(buf, "JPEG")
    jpeg = buf.getvalue()

    comps, w, h = source_map._ink_components(jpeg)
    stamp = source_map.snap_to_ink([0.6, 0.52, 0.85, 0.72], comps, w, h)
    assert stamp[0] == pytest.approx(0.625, abs=0.02) and stamp[1] == pytest.approx(0.6, abs=0.02)
    assert stamp[3] == pytest.approx(0.8, abs=0.02)

    dark, (w, h) = source_map._dark_mask(jpeg)
    line = source_map.snap_to_text([0.15, 0.29, 0.5, 0.305], dark, w, h)
    assert line[1] == pytest.approx(0.3, abs=0.004) and line[3] == pytest.approx(0.318, abs=0.004)
    assert line[0] == pytest.approx(0.125, abs=0.01) and line[2] == pytest.approx(0.625, abs=0.01)


def test_certification_paragraphs_are_identified_exactly():
    import gc
    import io

    from docx import Document

    from app.services import docx_blocks, docx_certification, source_map

    doc = Document()
    for i in range(40):
        doc.add_paragraph(f"Body paragraph {i} MUNICIPALITY OF ESEMPIO 01230045678901")
    buf = io.BytesIO()
    doc.save(buf)
    data = docx_blocks.tag_blocks(buf.getvalue())
    from datetime import date

    content = docx_certification.CertContent(lang="en", day=date(2026, 9, 29), values={
        "translator": "Anna Rossi", "date": "29 September 2026", "source_language": "Italian",
        "target_language": "English", "document": "Certificate", "pages": "1",
    })
    with_cert = docx_certification.add_certification(data, content)
    for _ in range(5):
        gc.collect()
        paras = source_map.document_paragraphs(with_cert)
        body = [p for p in paras if p["text"].startswith("Body paragraph")]
        assert len(body) == 40 and not any(p["cert"] for p in body)
        assert any(p["cert"] for p in paras)


def _pdf_with_table(rows):
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 100
    for row in rows:
        x = 60
        for cell in row:
            page.insert_text((x, y), cell, fontsize=10)
            x += 110
        y += 20
    return doc.tobytes()


def test_word_anchors_place_repeated_table_values_in_their_own_cells():
    from app.services import source_map as sm

    pdf = _pdf_with_table([["1", "126,00", "0,00", "N2.2"], ["2", "4,00", "0,00", "N2.2"], ["Totale", "130,00", "0,00"]])
    paras = [{"id": f"_b{i:08x}", "text": t} for i, t in enumerate(
        ["€126.00", "€0.00", "N2.2", "€4.00", "€0.00", "N2.2", "Total", "€130.00", "€0.00"])]
    res = sm.match_words(paras, sm.text_layer_words(pdf))
    tops = {p["text"] + str(i): res[p["id"]]["bbox"][1] for i, p in enumerate(paras) if p["id"] in res}
    assert tops["€0.001"] < tops["€0.004"] < tops["€0.008"]
    assert tops["N2.22"] < tops["N2.25"]
    for p in paras:
        if p["id"] in res:
            x0, y0, x1, y1 = res[p["id"]]["bbox"]
            assert (x1 - x0) * (y1 - y0) < 0.01


def test_coarse_boxes_rejected_for_short_paragraphs():
    from app.services import source_map as sm

    assert sm._too_coarse({"bbox": [0.05, 0.2, 0.95, 0.6]}, "€0.00")
    assert not sm._too_coarse({"bbox": [0.05, 0.2, 0.95, 0.6]}, "x" * 200)
    assert not sm._too_coarse({"bbox": [0.1, 0.2, 0.2, 0.22]}, "€0.00")
