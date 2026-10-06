"""Regressions from the pipeline and editor audit."""
import tempfile
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.models.project import ProjectStatus
from app.services import translation_processor as tp


@pytest.mark.parametrize(
    "name,expected",
    [
        ("contract.pdf", "contract.pdf"),
        ("Contract 1/2.pdf", "2.pdf"),
        ("/etc/evil.pdf", "evil.pdf"),
        ("../../evil.pdf", "evil.pdf"),
        ("..\\..\\evil.pdf", "evil.pdf"),
        ("..", "source"),
        ("", "source"),
        (None, "source"),
    ],
)
def test_local_file_name_stays_inside_the_folder(name, expected):
    assert tp.local_file_name(name) == expected


def test_job_writes_the_source_inside_its_temp_dir(db, make_user, make_project, monkeypatch):
    owner = make_user()
    project = make_project(owner, status=ProjectStatus.PENDING)
    project.file_name = "/../../../../tmp/escaped_by_name.pdf"
    db.commit()
    written = []

    def download(key, dest):
        written.append(Path(dest))
        Path(dest).write_bytes(b"%PDF")

    def rebuild(**kw):
        written.append(Path(kw["output_path"]))
        Path(kw["output_path"]).write_bytes(b"docx")
        return kw["output_path"]

    monkeypatch.setattr(tp, "download_file_from_s3", download)
    monkeypatch.setattr(tp, "extract_segments", lambda path: ("PDF", [SimpleNamespace(text="Ciao", layout={})]))
    monkeypatch.setattr(tp, "translate_batch", lambda texts, *a, **kw: [f"EN {t}" for t in texts])
    monkeypatch.setattr(tp, "translate_text", lambda text, *a, **kw: f"EN {text}")
    monkeypatch.setattr(tp, "rebuild_output", rebuild)
    monkeypatch.setattr(tp, "upload_file_to_s3", lambda path: "uploads/out.docx")
    tp.process_translation_job(str(project.id))

    temp_root = Path(tempfile.gettempdir()).resolve()
    assert len(written) == 2
    for path in written:
        assert path.resolve().parent.parent == temp_root
    assert [p.name for p in written] == ["escaped_by_name.pdf", "translated_escaped_by_name.pdf"]


def test_a_failed_vision_page_falls_back_instead_of_dropping_the_page(tmp_path, monkeypatch):
    from app.services import claude_vision_ocr, layout_translator
    from tests.conftest import make_pdf

    pdf = make_pdf(tmp_path / "two.pdf", pages=2)
    answers = iter([[{"text": "Pagina 1", "bbox": [10, 10, 200, 40]}], None])
    monkeypatch.setattr(claude_vision_ocr, "is_available", lambda: True)
    monkeypatch.setattr(claude_vision_ocr, "ocr_image", lambda path: next(answers))
    assert layout_translator._extract_pdf_via_claude(str(pdf)) is None

    blank_then_text = iter([[], [{"text": "Pagina 2", "bbox": [10, 10, 200, 40]}]])
    monkeypatch.setattr(claude_vision_ocr, "ocr_image", lambda path: next(blank_then_text))
    segs = layout_translator._extract_pdf_via_claude(str(pdf))
    assert [(s.text, s.layout["page"]) for s in segs] == [("Pagina 2", 1)]


@pytest.mark.parametrize("status", [ProjectStatus.PENDING, ProjectStatus.PROCESSING, ProjectStatus.FAILED])
@pytest.mark.parametrize("path", ["export", "export/pdf", "export/delivery.pdf"])
def test_unfinished_projects_do_not_export(client, make_user, make_project, status, path):
    owner = make_user()
    project = make_project(owner, status=status)
    r = client.get(f"/projects/{project.id}/{path}", headers=owner["headers"])
    assert r.status_code == 409


def _word_file(path):
    from docx import Document
    from docx.oxml import parse_xml

    doc = Document()
    p = doc.add_paragraph("Visit ")
    p._p.append(parse_xml(
        '<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:anchor="x">'
        "<w:r><w:t>our website</w:t></w:r></w:hyperlink>"
    ))
    table = doc.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Merged label"
    table.cell(1, 0).text = "Left"
    table.cell(1, 1).text = "Right"
    doc.save(path)
    return path


def test_word_extraction_and_rebuild_keep_links_and_merged_cells_single(tmp_path):
    from docx import Document

    from app.services import layout_translator as lt

    src = _word_file(tmp_path / "in.docx")
    kind, segs = lt.extract_segments(str(src))
    assert kind == "DOCX"
    assert [s.text for s in segs] == ["Visit our website", "Merged label", "Left", "Right"]

    out = tmp_path / "out.docx"
    lt._rebuild_docx(str(src), str(out), [(f"T:{s.text}", s.layout) for s in segs])
    rebuilt = Document(str(out))
    assert rebuilt.paragraphs[0].text == "T:Visit our website"
    assert rebuilt.tables[0].cell(0, 0).text == "T:Merged label"


def _ocr_reply(monkeypatch, body):
    import json

    from app.services import claude_params, claude_vision_ocr

    monkeypatch.setattr(claude_vision_ocr, "is_available", lambda: True)
    monkeypatch.setattr(
        claude_params, "create_message",
        lambda client, **kw: SimpleNamespace(content=[SimpleNamespace(type="text", text=json.dumps(body))], usage=None),
    )
    return claude_vision_ocr


def test_vision_bboxes_default_to_the_frame_the_model_saw(tmp_path, monkeypatch):
    from PIL import Image

    image = tmp_path / "scan.png"
    # Sent at 1568x392 (long edge cap); the reply gives no width/height.
    Image.new("RGB", (2000, 500), "white").save(image)
    ocr = _ocr_reply(monkeypatch, {"elements": [{"text": "Comune di Bari", "bbox": [100, 100, 500, 150]}]})
    (line,) = ocr.ocr_image(str(image))
    sx, sy = 2000 / 1568, 500 / 392
    assert line["bbox"] == pytest.approx([100 * sx, 100 * sy, 500 * sx, 150 * sy], rel=0.01)


def test_vision_reply_that_is_not_an_object_is_a_failed_read(tmp_path, monkeypatch):
    from PIL import Image

    image = tmp_path / "scan.png"
    Image.new("RGB", (400, 300), "white").save(image)
    assert _ocr_reply(monkeypatch, [{"text": "x"}]).ocr_image(str(image)) is None
    ocr = _ocr_reply(monkeypatch, {"elements": ["stray", {"text": "Kept", "bbox": [1, 1, 100, 40]}]})
    assert [line["text"] for line in ocr.ocr_image(str(image))] == ["Kept"]
