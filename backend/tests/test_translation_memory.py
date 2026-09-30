import io
import uuid
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.models.project import ProjectStatus
from app.models.translation_memory import TranslationMemory
from app.models.translation_segment import TranslationSegment
from app.services import claude_params, docx_blocks, tm_capture, tm_keys, tmx
from app.services import translation_memory_service as tm

SOURCE = [
    "L'UFFICIALE D'ANAGRAFE, visti gli atti d'ufficio,",
    "che il Sig. BIANCHI LUCA, nato a Bari il 12/03/1987,",
    "Ufficio Anagrafe e Stato Civile",
    "CERTIFICATO DI RESIDENZA",
]
FIRST_PASS = [
    "THE REGISTRY OFFICER, having examined the records,",
    "that Mr. BIANCHI LUCA, born in Bari on 12/03/1987,",
    "Registry and Civil Status Office",
    "CERTIFICATE OF RESIDENCE",
]


@pytest.fixture(autouse=True)
def no_claude(monkeypatch):
    def boom(*a, **kw):
        raise AssertionError("no model call expected")

    monkeypatch.setattr(claude_params, "create_message", boom)


def _docx(paras, cert=()) -> bytes:
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    cert_ps = [doc.add_paragraph(text) for text in cert]
    for p, name in ((cert_ps[:1], "_cert_start"), (cert_ps[-1:], "_cert_end")):
        for para in p:
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), "900" if name == "_cert_start" else "901")
            start.set(qn("w:name"), name)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), start.get(qn("w:id")))
            para._p.insert(0, start)
            para._p.append(end)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture()
def tm_project(db, storage, make_user, make_project):
    def _make(owner=None, paras=None, cert=(), source="Italian", target="English"):
        owner = owner or make_user()
        project = make_project(owner, segments=SOURCE)
        for seg in db.query(TranslationSegment).filter(TranslationSegment.project_id == project.id):
            seg.translated_text = FIRST_PASS[seg.segment_index]
        key = f"uploads/{uuid.uuid4()}_authored.docx"
        storage["objects"][key] = _docx(paras or FIRST_PASS[:2], cert)
        project.authored_docx_s3_key = key
        project.source_language, project.target_language = source, target
        db.commit()
        return owner, project

    return _make


def _entries(db, **filters):
    q = db.query(TranslationMemory)
    for k, v in filters.items():
        q = q.filter(getattr(TranslationMemory, k) == v)
    return q.order_by(TranslationMemory.source_text).all()


def _add(db, team_id, source, target, origin="approved", src="it", tgt="en", project_id=None):
    assert tm.upsert_entries(db, team_id, [{
        "source_language": src, "target_language": tgt, "source_text": source,
        "translated_text": target, "origin": origin, "project_id": project_id,
    }]) == 1


# Languages

def test_language_normaliser():
    assert tm_keys.source_lang("Italian") == tm_keys.source_lang("it-IT") == tm_keys.source_lang(" IT ") == "it"
    assert tm_keys.source_lang("auto") == tm_keys.source_lang("") == tm_keys.source_lang(None) == ""
    assert tm_keys.source_lang("Klingon") == ""
    assert tm_keys.target_lang("English (UK)") == tm_keys.target_lang("en_gb") == tm_keys.target_lang("en-GB") == "en-GB"
    assert tm_keys.target_lang("English") == "en" and tm_keys.target_lang("en-US") == "en-US"
    assert tm_keys.target_lang("Portuguese (Brazilian)") == "pt-BR"
    assert tm_keys.target_lang("auto") == ""
    assert tm_keys.source_lang("zh-TW") == tm_keys.source_lang("zh-Hant") == "zh-Hant" and tm_keys.source_lang("zh-CN") == "zh"
    assert tm_keys.family("en-GB") == tm_keys.family("en-US") == "en"
    assert tm_keys.family("zh-TW") != tm_keys.family("zh-CN")
    assert tm_keys.source_hash(" Ufficio \n Anagrafe") == tm_keys.source_hash("Ufficio Anagrafe")
    assert tm_keys.source_hash("Ufficio anagrafe") != tm_keys.source_hash("Ufficio Anagrafe")


# Upsert

def test_upsert_precedence(db, make_user):
    team = make_user()["team"].id
    _add(db, team, "Comune di Esempio", "Town of Esempio", "machine", src="Italian", tgt="English (UK)")
    row = _entries(db)[0]
    assert (row.source_language, row.target_language, row.origin) == ("it", "en-GB", "machine")

    _add(db, team, "Comune  di Esempio ", "Esempio Town", "machine", tgt="en-GB")
    assert [r.translated_text for r in _entries(db)] == ["Esempio Town"]

    _add(db, team, "Comune di Esempio", "Municipality of Esempio", "approved", tgt="en-GB")
    assert tm.store_tm_entry(db, team, "it", "en-GB", "Comune di Esempio", "AI again") == 0
    db.expire_all()
    rows = _entries(db)
    assert len(rows) == 1 and (rows[0].translated_text, rows[0].origin) == ("Municipality of Esempio", "approved")

    _add(db, team, "Comune di Esempio", "Municipality of Esempio (edited)", "manual", tgt="en-GB")
    db.expire_all()
    assert _entries(db)[0].origin == "manual"
    assert tm.store_tm_entry(db, team, "auto", "en-GB", "Altro", "Other") == 0
    assert tm.store_tm_entry(db, team, "it", "en", "x" * 1501, "long") == 0
    assert db.query(TranslationMemory).count() == 1


def test_upsert_failure_is_swallowed_and_keeps_the_session(db, make_user):
    owner = make_user()
    assert tm.store_tm_entry(db, uuid.uuid4(), "it", "en", "Ciao", "Hello") == 0
    owner["user"].full_name = "Still usable"
    db.commit()


def test_lookup_prefers_exact_target_then_same_language(db, make_user):
    team = make_user()["team"].id
    _add(db, team, "Ufficio Anagrafe", "Registry Office (US)", tgt="en-US")
    _add(db, team, "Stato Civile", "Civil Status (US)", tgt="en-US")
    _add(db, team, "Stato Civile", "Civil Status (UK)", tgt="en-GB")
    _add(db, team, "Stato Civile", "Estado civil", tgt="es")
    found = tm.lookup(db, team, "Italian", "English (UK)", ["Ufficio  Anagrafe", "Stato Civile"])
    assert sorted(r.translated_text for r in found.values()) == ["Civil Status (UK)", "Registry Office (US)"]
    assert tm.lookup(db, team, "auto", "en-GB", ["Stato Civile"]) == {}


# Migration data logic

def test_cleanup_plan_resolves_auto_links_projects_and_dedupes():
    t = "team-1"
    rows = [
        {"id": "a", "team_id": t, "source_language": "auto", "target_language": "en-GB", "source_text": "Comune", "translated_text": "Town (old)", "seq": 1},
        {"id": "b", "team_id": t, "source_language": "auto", "target_language": "en-GB", "source_text": "Comune ", "translated_text": "Town", "seq": 2},
        {"id": "c", "team_id": t, "source_language": "auto", "target_language": "en-US", "source_text": "Orfano", "translated_text": "Orphan", "seq": 3},
        {"id": "d", "team_id": t, "source_language": "Italian", "target_language": "English", "source_text": "Anagrafe", "translated_text": "Registry", "seq": 4},
        {"id": "e", "team_id": t, "source_language": "it", "target_language": "en", "source_text": "Anagrafe", "translated_text": "Registry office", "seq": 5},
        {"id": "f", "team_id": t, "source_language": "it", "target_language": "", "source_text": "X", "translated_text": "Y", "seq": 6},
    ]
    old = {"id": "p1", "source_language": "auto", "target_language": "en-GB", "doc_profile": {"source_language": "Italian"}, "created_at": datetime(2026, 1, 1)}
    new = {"id": "p2", "source_language": "Italian", "target_language": "English", "doc_profile": None, "created_at": datetime(2026, 2, 1)}
    plan = tm_keys.plan_cleanup(rows, {"a": [old], "b": [old], "d": [new]})
    kept = {u["id"]: u for u in plan["update"]}
    assert set(kept) == {"b", "d"}
    assert (kept["b"]["source_language"], kept["b"]["target_language"], kept["b"]["project_id"]) == ("it", "en-GB", "p1")
    assert kept["d"]["project_id"] == "p2" and kept["d"]["source_hash"] == tm_keys.source_hash("Anagrafe")
    assert sorted(plan["delete"]) == ["a", "c", "e", "f"]
    again = tm_keys.plan_cleanup([dict(r, **kept[r["id"]]) for r in rows if r["id"] in kept], {"b": [old], "d": [new]})
    assert again["delete"] == [] and {u["id"] for u in again["update"]} == {"b", "d"}


def test_migration_cleanup_runs_on_postgres(db, make_user, make_project):
    import importlib.util

    from sqlalchemy import text

    spec = importlib.util.spec_from_file_location(
        "tm_migration", Path(__file__).resolve().parent.parent / "alembic/versions/d7f3a9c1e5b2_translation_memory_keys.py"
    )
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    owner = make_user()
    project = make_project(owner, segments=["Comune di Esempio"])
    project.source_language, project.target_language, project.doc_profile = "auto", "en-GB", {"source_language": "Italian"}
    db.commit()
    team = owner["team"].id
    for src, tgt, text_, hash_ in (("auto", "en-GB", "Comune di Esempio", "h1"), ("auto", "en-GB", "Comune di Esempio", "h2"), ("auto", "en-US", "Orfano", "h3")):
        db.add(TranslationMemory(team_id=team, source_language=src, target_language=tgt, source_text=text_,
                                 translated_text="t " + hash_, source_hash=hash_))
    db.commit()
    migration.cleanup_rows(db.connection())
    db.commit()
    db.expire_all()
    rows = _entries(db)
    assert len(rows) == 1
    assert (rows[0].source_language, rows[0].target_language, rows[0].project_id) == ("it", "en-GB", project.id)
    assert rows[0].source_hash == tm_keys.source_hash("Comune di Esempio")
    assert db.execute(text("SELECT count(*) FROM translation_memory")).scalar() == 1


# Approved text from the editor

def _open(client, owner, project):
    r = client.get(f"/projects/{project.id}/document", headers=owner["headers"])
    return r.content, int(r.headers["X-Document-Version"])


def test_editor_save_stores_approved_pairs(client, db, tm_project):
    owner, project = tm_project()
    data, v = _open(client, owner, project)
    ids = docx_blocks.block_ids(data)
    r = client.post(f"/projects/{project.id}/document/edits", headers=owner["headers"], json={"version": v, "edits": [
        {"block_id": ids[0], "text": "THE REGISTRAR, having examined the records,"},
        # A changed number no longer matches the source line, so it isn't stored.
        {"block_id": ids[1], "text": "that Mr. BIANCHI LUCA, born in Bari on 12/03/1986,"},
    ]})
    assert r.status_code == 200, r.text
    rows = _entries(db)
    assert [(r.source_text, r.translated_text, r.origin, r.source_language, r.target_language) for r in rows] == [
        (SOURCE[0], "THE REGISTRAR, having examined the records,", "approved", "it", "en"),
    ]
    assert rows[0].project_id == project.id

    # A second edit of the same paragraph maps through the first version and replaces the entry.
    client.post(f"/projects/{project.id}/document/edits", headers=owner["headers"], json={"version": v + 1, "edits": [
        {"block_id": ids[0], "text": "THE REGISTRAR, having examined the office records,"},
    ]})
    db.expire_all()
    assert [r.translated_text for r in _entries(db)] == ["THE REGISTRAR, having examined the office records,"]


def test_editor_save_respects_use_tm(client, db, tm_project):
    owner, project = tm_project()
    project.use_tm = False
    db.commit()
    data, v = _open(client, owner, project)
    ids = docx_blocks.block_ids(data)
    client.post(f"/projects/{project.id}/document/edits", headers=owner["headers"],
                json={"version": v, "edits": [{"block_id": ids[0], "text": "THE REGISTRAR, having examined the records,"}]})
    assert _entries(db) == []


def test_certify_stores_the_document_but_not_the_certification_page(client, db, tm_project):
    owner, project = tm_project(
        paras=FIRST_PASS[:2] + ["[Signature]", "Something the AI invented"],
        cert=["Registry and Civil Status Office", "I certify this translation."],
        source="auto", target="en-GB",
    )
    project.doc_profile = {"source_language": "Italian"}
    db.commit()
    r = client.patch(f"/projects/{project.id}/review-status", headers=owner["headers"], json={"status": "CERTIFIED"})
    assert r.status_code == 200
    rows = _entries(db)
    assert {(r.source_text, r.translated_text) for r in rows} == {(SOURCE[0], FIRST_PASS[0]), (SOURCE[1], FIRST_PASS[1])}
    assert {(r.source_language, r.target_language, r.origin) for r in rows} == {("it", "en-GB", "approved")}


def test_delivery_skips_unknown_source_language(client, db, tm_project):
    owner, project = tm_project(source="auto")
    client.patch(f"/projects/{project.id}/review-status", headers=owner["headers"], json={"status": "CERTIFIED"})
    assert _entries(db) == []


def test_paragraph_mapping_rules():
    index = tm_capture.SegmentIndex([(i, s, t) for i, (s, t) in enumerate(zip(SOURCE, FIRST_PASS))])
    assert index.source_of(["THE REGISTRY OFFICER, having examined the records,"]) == SOURCE[0]
    # Two lines merged into one paragraph map to both source lines.
    merged = FIRST_PASS[0] + " " + FIRST_PASS[1]
    assert index.source_of([merged]) == SOURCE[0] + " " + SOURCE[1]
    assert index.source_of(["Completely unrelated sentence about weather"]) is None
    assert not tm_capture.usable_pair("[Firma]", "[Signature]")
    assert not tm_capture.usable_pair("Nato il 12/03/1987", "Born on 12/03/1986")
    assert not tm_capture.usable_pair("x" * 1501, "y" * 1501)
    assert tm_capture.usable_pair("Nato il 12/03/1987", "Born on 03/12/1987")


# Main translation step

def _approved(db, project, pairs):
    src, tgt = tm.project_pair(project)
    for s, t in pairs:
        _add(db, project.team_id, s, t, "approved", src=src, tgt=tgt)


def test_prompt_block_prefers_approved_and_is_capped(db, make_user, make_project):
    owner = make_user()
    project = make_project(owner)
    _approved(db, project, [(SOURCE[0], "THE REGISTRAR, having examined the records,")])
    _add(db, project.team_id, SOURCE[2], "Registry Office (AI)", "machine")
    _add(db, project.team_id, SOURCE[3], "RESIDENCE CERTIFICATE", "manual", tgt="en-GB")
    block = tm.prompt_block(db, project, "\n".join(SOURCE))
    assert block.startswith(tm.PROMPT_HEADING)
    assert "THE REGISTRAR" in block and "RESIDENCE CERTIFICATE" in block and "Registry Office (AI)" not in block
    assert block.index("RESIDENCE CERTIFICATE") < block.index("THE REGISTRAR")

    short = [f"Riga {i}" for i in range(200)]
    _approved(db, project, [(s, s.replace("Riga", "Line")) for s in short])
    assert tm.prompt_block(db, project, "\n".join(short)).count("- SOURCE:") == tm.PROMPT_MAX_ENTRIES

    long = [f"Riga numero {i} di un paragrafo piuttosto lungo del documento di prova" for i in range(200)]
    _approved(db, project, [(s, s.replace("Riga numero", "Line number")) for s in long])
    block = tm.prompt_block(db, project, "\n".join(long))
    assert 50 < block.count("- SOURCE:") < tm.PROMPT_MAX_ENTRIES and len(block) <= tm.PROMPT_MAX_CHARS

    project.use_tm = False
    assert tm.prompt_block(db, project, "\n".join(SOURCE)) == ""


def _run_job(db, project, monkeypatch):
    from app.services import claude_authored_rebuild as car
    from app.services import translation_processor as tp

    seen = {}
    monkeypatch.setattr(tp, "download_file_from_s3", lambda key, dest: Path(dest).write_bytes(b"%PDF"))
    monkeypatch.setattr(tp, "extract_segments", lambda path: ("PDF", [SimpleNamespace(text=s, layout={}) for s in SOURCE]))
    monkeypatch.setattr(tp.learning, "profile_project", lambda *a, **kw: None)
    monkeypatch.setattr(tp, "translate_batch", lambda texts, *a, **kw: [f"EN {t}" for t in texts])
    monkeypatch.setattr(tp, "translate_text", lambda text, *a, **kw: f"EN {text}")

    def rebuild(**kw):
        Path(kw["output_path"]).write_bytes(b"docx")
        return kw["output_path"]

    monkeypatch.setattr(tp, "rebuild_output", rebuild)
    monkeypatch.setattr(tp, "upload_file_to_s3", lambda path: "uploads/out.docx")
    monkeypatch.setattr(car, "author_rebuild_docx", lambda **kw: seen.update(kw) or b"docx")
    project.status = ProjectStatus.PENDING
    project.model = "claude-authored"
    db.commit()
    tp.process_translation_job(str(project.id))
    return seen


def test_rebuild_receives_the_memory_block(db, make_user, make_project, monkeypatch):
    owner = make_user()
    project = make_project(owner, segments=["x"])
    _approved(db, project, [(SOURCE[0], "THE REGISTRAR, having examined the records,")])
    seen = _run_job(db, project, monkeypatch)
    assert tm.PROMPT_HEADING in seen["terminology"]
    assert "TRANSLATION: THE REGISTRAR, having examined the records," in seen["terminology"]
    db.expire_all()
    seg = db.query(TranslationSegment).filter(TranslationSegment.project_id == project.id, TranslationSegment.segment_index == 0).one()
    assert seg.translated_text == "THE REGISTRAR, having examined the records," and seg.tm_pct == 100
    # First-pass drafts stay out of the memory; only the translator's text goes in.
    assert _entries(db, origin="machine") == []


def test_rebuild_without_memory_when_use_tm_is_off(db, make_user, make_project, monkeypatch):
    owner = make_user()
    project = make_project(owner, segments=["x"])
    project.use_tm = False
    _approved(db, project, [(SOURCE[0], "THE REGISTRAR, having examined the records,")])
    seen = _run_job(db, project, monkeypatch)
    assert tm.PROMPT_HEADING not in (seen["terminology"] or "")
    assert _entries(db, origin="machine") == []


def test_auto_source_without_detection_stores_nothing(db, make_user, make_project, monkeypatch):
    owner = make_user()
    project = make_project(owner, segments=["x"])
    project.source_language = "auto"
    _run_job(db, project, monkeypatch)
    assert db.query(TranslationMemory).count() == 0


# API

def test_crud_and_team_isolation(client, db, make_user, make_project):
    owner = make_user()
    member = make_user(team=owner["team"])
    other = make_user()
    h, oh = owner["headers"], other["headers"]

    r = client.post("/tm/", headers=h, json={"source_language": "Italian", "target_language": "en-GB",
                                             "source_text": "Stato Civile", "translated_text": "Civil Status"})
    assert r.status_code == 200, r.text
    entry = r.json()
    assert (entry["source_language"], entry["target_language"], entry["origin"]) == ("it", "en-GB", "manual")
    assert client.post("/tm/", headers=h, json={"source_language": "auto", "target_language": "en",
                                                "source_text": "a", "translated_text": "b"}).status_code == 400
    project = make_project(owner)
    _add(db, owner["team"].id, "Ufficio Anagrafe", "Registry Office", "machine", project_id=project.id)

    listed = client.get("/tm/", headers=member["headers"], params={"limit": 1}).json()
    assert listed["total"] == 2 and len(listed["items"]) == 1
    assert client.get("/tm/", headers=h, params={"origin": "machine"}).json()["items"][0]["project_name"] == "source.pdf"
    assert client.get("/tm/", headers=h, params={"project_id": str(project.id)}).json()["total"] == 1
    assert client.get("/tm/", headers=h, params={"source": "Italian", "target": "English (UK)"}).json()["total"] == 1
    assert client.get("/tm/", headers=h, params={"q": "civil"}).json()["total"] == 1
    assert client.get("/tm/", headers=oh).json()["total"] == 0

    summary = client.get("/tm/summary", headers=h).json()
    assert summary["total_units"] == 2 and summary["origins"]["manual"] == 1 and summary["origins"]["machine"] == 1
    assert {(p["source"], p["target"]) for p in summary["language_pairs"]} == {("it", "en-GB"), ("it", "en")}
    assert summary["projects"][0]["id"] == str(project.id) and summary["source_words_indexed"] == 4

    eid = entry["id"]
    assert client.patch(f"/tm/{eid}", headers=oh, json={"translated_text": "Hacked"}).status_code == 404
    assert client.delete(f"/tm/{eid}", headers=oh).status_code == 404
    assert client.post("/tm/bulk-delete", headers=oh, json={"ids": [eid]}).json()["deleted"] == 0

    machine_id = client.get("/tm/", headers=h, params={"origin": "machine"}).json()["items"][0]["id"]
    r = client.patch(f"/tm/{machine_id}", headers=h, json={"translated_text": "Registry Office (checked)"})
    assert r.status_code == 200 and r.json()["origin"] == "manual"
    _add(db, owner["team"].id, "Prefettura", "Prefecture", "machine")
    assert client.patch(f"/tm/{machine_id}", headers=h, json={"source_text": " Prefettura"}).status_code == 409
    r = client.patch(f"/tm/{machine_id}", headers=h, json={"source_text": "Ufficio dell'Anagrafe"})
    assert r.status_code == 200 and r.json()["source_text"] == "Ufficio dell'Anagrafe"

    assert client.delete(f"/tm/{eid}", headers=h).json() == {"deleted": 1}
    left = [e["id"] for e in client.get("/tm/", headers=h).json()["items"]]
    assert client.post("/tm/bulk-delete", headers=h, json={"ids": left}).json()["deleted"] == 2
    assert db.query(TranslationMemory).count() == 0


def test_tm_requires_pro_plan(client, make_user):
    trial = make_user(plan="TRIAL")
    assert client.get("/tm/", headers=trial["headers"]).status_code == 403


def test_tmx_round_trip(client, db, make_user):
    owner, other = make_user(), make_user()
    _add(db, owner["team"].id, "Stato Civile & <Anagrafe>", "Civil Status & <Registry>", "approved", tgt="en-GB")
    _add(db, owner["team"].id, "Comune", "Municipalité", "manual", tgt="fr")
    r = client.get("/tm/export", headers=owner["headers"], params={"source": "it", "target": "en-GB"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert b'<tmx version="1.4">' in r.content and b"Municipalit" not in r.content

    full = client.get("/tm/export", headers=owner["headers"]).content
    r = client.post("/tm/import", headers=other["headers"], files={"file": ("tm.tmx", full, "application/xml")})
    assert r.status_code == 200, r.text
    assert r.json()["imported"] == 2 and r.json()["units"] == 2
    rows = _entries(db, team_id=other["team"].id)
    assert {(r.source_text, r.translated_text, r.target_language, r.origin) for r in rows} == {
        ("Comune", "Municipalité", "fr", "import"),
        ("Stato Civile & <Anagrafe>", "Civil Status & <Registry>", "en-GB", "import"),
    }


def test_tmx_import_reads_inline_codes_and_multiple_targets():
    data = b"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE tmx SYSTEM "tmx14.dtd">
<tmx version="1.4"><header srclang="it-IT" creationtool="x" creationtoolversion="1" segtype="sentence" o-tmf="x" adminlang="en" datatype="plaintext"/>
<body><tu><tuv xml:lang="en-US"><seg>Registry <bpt i="1">&lt;b&gt;</bpt>Office<ept i="1">&lt;/b&gt;</ept></seg></tuv>
<tuv xml:lang="it-IT"><seg>Ufficio Anagrafe</seg></tuv><tuv xml:lang="de"><seg>Meldeamt</seg></tuv></tu></body></tmx>"""
    pairs, units = tmx.parse(data)
    assert units == 1
    assert pairs == [
        {"source_language": "it-IT", "target_language": "en-US", "source_text": "Ufficio Anagrafe", "translated_text": "Registry Office"},
        {"source_language": "it-IT", "target_language": "de", "source_text": "Ufficio Anagrafe", "translated_text": "Meldeamt"},
    ]


def test_tmx_import_rejects_entities_and_oversized_files(client, make_user, monkeypatch):
    h = make_user()["headers"]
    bomb = b"""<?xml version="1.0"?>
<!DOCTYPE tmx [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>
<tmx version="1.4"><header srclang="it"/><body><tu><tuv xml:lang="it"><seg>&b;</seg></tuv><tuv xml:lang="en"><seg>x</seg></tuv></tu></body></tmx>"""
    r = client.post("/tm/import", headers=h, files={"file": ("bomb.tmx", bomb, "application/xml")})
    assert r.status_code == 400 and "entities" in r.json()["detail"]
    xxe = b"""<?xml version="1.0"?><!DOCTYPE tmx [<!ENTITY x SYSTEM "file:///etc/passwd">]><tmx version="1.4"><body/></tmx>"""
    assert client.post("/tm/import", headers=h, files={"file": ("x.tmx", xxe, "application/xml")}).status_code == 400
    assert client.post("/tm/import", headers=h, files={"file": ("x.tmx", b"not xml", "application/xml")}).status_code == 400

    big = b"<tmx>" + b" " * (tmx.MAX_BYTES + 10) + b"</tmx>"
    assert client.post("/tm/import", headers=h, files={"file": ("big.tmx", big, "application/xml")}).status_code == 413

    monkeypatch.setattr(tmx, "MAX_UNITS", 1)
    two = b"""<tmx version="1.4"><header srclang="it"/><body>
<tu><tuv xml:lang="it"><seg>a</seg></tuv><tuv xml:lang="en"><seg>b</seg></tuv></tu>
<tu><tuv xml:lang="it"><seg>c</seg></tuv><tuv xml:lang="en"><seg>d</seg></tuv></tu></body></tmx>"""
    r = client.post("/tm/import", headers=h, files={"file": ("two.tmx", two, "application/xml")})
    assert r.status_code == 400 and "translation units" in r.json()["detail"]
