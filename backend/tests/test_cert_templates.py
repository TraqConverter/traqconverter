import io
import posixpath
import uuid
import zipfile
from datetime import date

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
from PIL import Image

from app.services import cert_fields, docx_blocks, docx_cert_template, docx_certification, docx_images

W = docx_blocks.W_NS
WP = docx_blocks.WP_NS
R = docx_blocks.R_NS
A = docx_blocks.A_NS
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _png(color=(20, 90, 200), size=(90, 40)) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    buf.seek(0)
    return buf


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _fmt(run, bold=True, size=14, font="Georgia"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = font
    return run


def _complex_field(p, instr_parts, result, **fmt):
    """A MERGEFIELD the way Word writes it: begin / instrText (split over runs) / separate / result / end."""

    def run_with(tag, **attrs):
        run = _fmt(p.add_run(), **fmt)
        el = etree.SubElement(run._r, qn(tag))
        for k, v in attrs.items():
            el.set(qn(k), v)
        return el

    run_with("w:fldChar", **{"w:fldCharType": "begin"})
    for part in instr_parts:
        el = run_with("w:instrText")
        el.text = part
        el.set(XML_SPACE, "preserve")
    run_with("w:fldChar", **{"w:fldCharType": "separate"})
    _fmt(p.add_run(result), **fmt)
    run_with("w:fldChar", **{"w:fldCharType": "end"})


def _simple_field(p, instr, result, **fmt):
    fld = etree.SubElement(p._p, qn("w:fldSimple"))
    fld.set(qn("w:instr"), instr)
    run = etree.SubElement(fld, qn("w:r"))
    if fmt:
        rpr = etree.SubElement(run, qn("w:rPr"))
        etree.SubElement(rpr, qn("w:b"))
        etree.SubElement(rpr, qn("w:sz")).set(qn("w:val"), str(fmt.get("size", 12) * 2))
    etree.SubElement(run, qn("w:t")).text = result


def _template(images=False, extras=False) -> bytes:
    doc = Document()
    style = doc.styles.add_style("CertHeading", 1)
    style.font.name = "Garamond"
    style.font.size = Pt(18)
    style.font.bold = True
    if images:
        doc.sections[0].header.paragraphs[0].add_run().add_picture(_png((200, 30, 30)), width=Pt(72))
        doc.sections[0].header.add_paragraph("STUDIO ROSSI TRADUZIONI")
    doc.add_paragraph("DICHIARAZIONE DI TRADUZIONE", style="CertHeading")
    if images:
        doc.add_picture(_png(), width=Pt(80))
    p = doc.add_paragraph("Io sottoscritto ")
    _complex_field(p, [" MERGEFIELD  Tradu", "ttore \\* MERGEFORMAT "], "«Traduttore»")
    p.add_run(", in data ")
    _simple_field(p, " MERGEFIELD Data \\* MERGEFORMAT ", "«Data»")
    p.add_run(", ho tradotto dal {{Source_Language}} verso ")
    _complex_field(p, [" MERGEFIELD LinguaDestinazione "], "«LinguaDestinazione»", bold=False, size=11, font="Arial")
    p.add_run(" il documento ")
    _simple_field(p, " MERGEFIELD Documento ", "«Documento»", size=12)
    p.add_run(" di ")
    _complex_field(p, [" MERGEFIELD Pagine "], "«Pagine»", bold=False, size=11, font="Arial")
    p.add_run(" pagine. ")
    _simple_field(p, ' MERGEFIELD Client \\b "Cliente: " ', "«Client»")
    q = doc.add_paragraph("Timbro: ")
    _simple_field(q, " MERGEFIELD Timbro ", "«Timbro»")
    q.add_run(" {{codice_albo}}")
    table = doc.add_table(rows=1, cols=2)
    table.style = doc.styles["Table Grid"]
    table.cell(0, 0).text = "Source"
    _simple_field(table.cell(0, 1).paragraphs[0], " MERGEFIELD SourceLanguage ", "«SourceLanguage»")
    if extras:
        _hyperlink(doc.add_paragraph("Web: "), "https://studio-rossi.test", "studio-rossi.test")
        _text_box(doc.add_paragraph())
        doc.add_paragraph("Primo punto", style="List Number")
    doc.add_paragraph("Firma ______________________")
    return _save(doc)


def _hyperlink(p, url, text):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = etree.SubElement(p._p, qn("w:hyperlink"))
    link.set(qn("r:id"), rid)
    run = etree.SubElement(link, qn("w:r"))
    etree.SubElement(run, qn("w:t")).text = text


def _text_box(p):
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    run = etree.SubElement(p._p, qn("w:r"))
    alt = etree.SubElement(run, f"{{{mc}}}AlternateContent", nsmap={"mc": mc})
    etree.SubElement(alt, f"{{{mc}}}Fallback")


def _content(**kw):
    values = {
        "translator": "Anna Rossi",
        "date": "29 September 2026",
        "source_language": "Italian",
        "target_language": "English",
        "document": "Birth certificate",
        "pages": "2",
        "client": "Bianchi Srl",
    }
    return docx_certification.CertContent(lang="en", values=values, day=date(2026, 9, 29), pages=2, **kw)


def _translation(with_image=False) -> bytes:
    doc = Document()
    doc.add_paragraph("CERTIFICATE OF RESIDENCE")
    if with_image:
        doc.add_picture(_png((10, 10, 10)), width=Pt(50))
    doc.add_paragraph("Mr. BIANCHI LUCA, born in Bari")
    return docx_blocks.tag_blocks(_save(doc))


def _xml(data: bytes, name="word/document.xml"):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return etree.fromstring(z.read(name))


def _texts(data):
    body = _xml(data).find(f"{{{W}}}body")
    return [docx_blocks.paragraph_text(p) for p in body.iter(f"{{{W}}}p")]


def _sdt(data, name):
    for sdt in _xml(data).iter(f"{{{W}}}sdt"):
        tag = sdt.find(f"{{{W}}}sdtPr/{{{W}}}tag")
        if tag is not None and tag.get(f"{{{W}}}val") == f"cert.{name}":
            return sdt
    return None


def _assert_valid_package(data: bytes):
    Document(io.BytesIO(data))
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = set(z.namelist())
        rels = etree.fromstring(z.read("word/_rels/document.xml.rels"))
        types = z.read("[Content_Types].xml")
    by_id = {r.get("Id"): r for r in rels}
    assert len(by_id) == len(rels), "duplicate relationship ids"
    body = _xml(data)
    embeds = [el.get(f"{{{R}}}embed") for el in body.iter(f"{{{A}}}blip")]
    for rid in embeds:
        rel = by_id[rid]
        assert rel.get("Type").endswith("/image")
        assert posixpath.normpath(posixpath.join("word", rel.get("Target"))) in names
    for link in body.iter(f"{{{W}}}hyperlink"):
        rid = link.get(f"{{{R}}}id")
        if rid:
            assert by_id[rid].get("TargetMode") == "External"
    docpr = [d.get("id") for d in body.iter(f"{{{WP}}}docPr")]
    assert len(docpr) == len(set(docpr))
    if embeds:
        assert b'Extension="png"' in types
    return embeds


# --- names ------------------------------------------------------------------


def test_field_names_are_matched_loosely():
    for name, field in [
        ("Date", "date"), ("DATA", "date"), ("Translator", "translator"), ("Traduttore", "translator"),
        ("SourceLanguage", "source_language"), ("Source_Language", "source_language"), ("source language", "source_language"),
        ("LinguaOrigine", "source_language"), ("Lingua_Origine", "source_language"), ("TargetLanguage", "target_language"),
        ("LinguaDestinazione", "target_language"), ("Document", "document"), ("Documento", "document"),
        ("Pages", "pages"), ("Pagine", "pages"), ("Übersetzer", "translator"), ("translator_name", "translator"),
        ("page_count", "pages"), ("Cliente", "client"), ("certificate_number", "certificate_number"),
    ]:
        assert cert_fields.resolve(name) == field, name
    assert cert_fields.resolve("Timbro") is None
    for field, aliases in cert_fields.ALIASES.items():
        for alias in aliases:
            assert cert_fields.resolve(alias) == field, alias


# --- merge fields -----------------------------------------------------------


def test_merge_fields_and_tokens_become_content_controls_keeping_formatting():
    report = docx_cert_template.TemplateReport()
    data = docx_certification.add_certification(_translation(), _content(template_docx=_template()), report)
    texts = _texts(data)
    assert "DICHIARAZIONE DI TRADUZIONE" in texts
    assert (
        "Io sottoscritto Anna Rossi, in data 29 September 2026, ho tradotto dal Italian verso English "
        "il documento Birth certificate di 2 pagine. Cliente: Bianchi Srl"
    ) in texts
    assert "Timbro: «Timbro» {{codice_albo}}" in texts
    assert "Italian" in texts
    body = _xml(data)
    assert next(body.iter(f"{{{W}}}fldChar"), None) is None
    assert next(body.iter(f"{{{W}}}instrText"), None) is None
    assert next(body.iter(f"{{{W}}}fldSimple"), None) is None

    fields = docx_certification.read_fields(data)
    assert fields["translator"] == "Anna Rossi" and fields["date_iso"] == "2026-09-29"
    assert fields["pages"] == "2" and fields["client"] == "Bianchi Srl"

    run = _sdt(data, "translator").find(f"{{{W}}}sdtContent/{{{W}}}r")
    rpr = run.find(f"{{{W}}}rPr")
    assert rpr.find(f"{{{W}}}b") is not None
    assert rpr.find(f"{{{W}}}sz").get(f"{{{W}}}val") == "28"
    assert rpr.find(f"{{{W}}}rFonts").get(f"{{{W}}}ascii") == "Georgia"
    target = _sdt(data, "target_language").find(f"{{{W}}}sdtContent/{{{W}}}r/{{{W}}}rPr")
    assert target.find(f"{{{W}}}b").get(f"{{{W}}}val") == "0"
    assert target.find(f"{{{W}}}rFonts").get(f"{{{W}}}ascii") == "Arial"
    assert _sdt(data, "date").find(f"{{{W}}}sdtPr/{{{W}}}date") is not None

    found = {(f["name"], f["kind"]): f["field"] for f in report.fields}
    assert found[("Traduttore", "merge")] == "translator"
    assert found[("Data", "merge")] == "date"
    assert found[("Source_Language", "token")] == "source_language"
    assert found[("LinguaDestinazione", "merge")] == "target_language"
    assert found[("Documento", "merge")] == "document"
    assert found[("SourceLanguage", "merge")] == "source_language"
    assert {(u["name"], u["kind"]) for u in report.unknown} == {("Timbro", "merge"), ("codice_albo", "token")}


def test_merge_field_page_updates_in_place():
    data = docx_certification.add_certification(_translation(), _content(template_docx=_template()))
    data, changed = docx_certification.update_fields(
        data, {"translator": "Anna Bianchi", "date": "2026-10-01", "target_language": "Spanish"}, "en"
    )
    assert set(changed) == {"translator", "date", "target_language"}
    fields = docx_certification.read_fields(data)
    assert fields["translator"] == "Anna Bianchi" and fields["date"] == "1 October 2026"
    assert any("verso Spanish il documento" in t for t in _texts(data))
    rpr = _sdt(data, "translator").find(f"{{{W}}}sdtContent/{{{W}}}r/{{{W}}}rPr")
    assert rpr.find(f"{{{W}}}b") is not None and rpr.find(f"{{{W}}}sz").get(f"{{{W}}}val") == "28"
    again = docx_certification.add_certification(data, _content(template_docx=_template()))
    assert _texts(again).count("DICHIARAZIONE DI TRADUZIONE") == 1


def test_template_styles_links_and_dropped_parts():
    report = docx_cert_template.TemplateReport()
    data = docx_certification.add_certification(_translation(), _content(template_docx=_template(extras=True)), report)
    _assert_valid_package(data)
    styles = _xml(data, "word/styles.xml")
    by_id = {s.get(f"{{{W}}}styleId"): s for s in styles.findall(f"{{{W}}}style")}
    body = _xml(data)
    heading = next(p for p in body.iter(f"{{{W}}}p") if docx_blocks.paragraph_text(p) == "DICHIARAZIONE DI TRADUZIONE")
    sid = heading.find(f"{{{W}}}pPr/{{{W}}}pStyle").get(f"{{{W}}}val")
    assert sid.startswith(docx_cert_template.STYLE_PREFIX)
    style = by_id[sid]
    assert style.find(f"{{{W}}}rPr/{{{W}}}rFonts").get(f"{{{W}}}ascii") == "Garamond"
    assert style.find(f"{{{W}}}rPr/{{{W}}}sz").get(f"{{{W}}}val") == "36"
    assert style.find(f"{{{W}}}pPr/{{{W}}}spacing") is not None
    table = next(body.iter(f"{{{W}}}tbl"))
    table_style = by_id[table.find(f"{{{W}}}tblPr/{{{W}}}tblStyle").get(f"{{{W}}}val")]
    based = table_style.find(f"{{{W}}}basedOn").get(f"{{{W}}}val")
    assert based in by_id and based.startswith(docx_cert_template.STYLE_PREFIX)
    link = next(body.iter(f"{{{W}}}hyperlink"))
    assert docx_blocks.paragraph_text(link) == "studio-rossi.test"
    rels = {r.get("Id"): r for r in _xml(data, "word/_rels/document.xml.rels")}
    assert rels[link.get(f"{{{R}}}id")].get("Target") == "https://studio-rossi.test"
    assert report.dropped == {"Text boxes and shapes": 1, "List numbering and bullets": 1}
    assert "Primo punto" in _texts(data)
    translation_styles = [p.find(f"{{{W}}}pPr/{{{W}}}pStyle") for p in body.iter(f"{{{W}}}p") if docx_blocks.paragraph_text(p) == "Mr. BIANCHI LUCA, born in Bari"]
    assert translation_styles == [None]


def test_template_pictures_are_carried_over_with_valid_relationships():
    report = docx_cert_template.TemplateReport()
    logo = docx_images.prepare_image(_png((0, 200, 0)).getvalue())
    stamp = docx_images.prepare_image(_png((0, 0, 0), (120, 120)).getvalue())
    data = docx_certification.add_certification(
        _translation(with_image=True), _content(template_docx=_template(images=True), logo=logo, stamp=stamp), report
    )
    embeds = _assert_valid_package(data)
    assert len(set(embeds)) == 4  # translation picture, template header + body pictures, stamp; no logo
    texts = _texts(data)
    assert "STUDIO ROSSI TRADUZIONI" in texts
    assert texts.index("STUDIO ROSSI TRADUZIONI") < texts.index("DICHIARAZIONE DI TRADUZIONE")
    assert any("logo" in n for n in report.notes) and any("header" in n for n in report.notes)
    ids = [i["id"] for i in docx_images.list_images(data)]
    assert len(ids) == len(set(ids)) == 4
    removed = docx_certification.remove_certification(data)
    assert len(_assert_valid_package(removed)) == 1


def test_unreadable_template_pictures_are_dropped_and_reported():
    tpl = _template(images=True)
    with zipfile.ZipFile(io.BytesIO(tpl)) as z:
        files = {n: z.read(n) for n in z.namelist()}
    for name in files:
        if name.startswith("word/media/"):
            files[name] = b"\x01\x00\x00\x00 EMF picture"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, raw in files.items():
            z.writestr(name, raw)
    report = docx_cert_template.TemplateReport()
    data = docx_certification.add_certification(_translation(), _content(template_docx=buf.getvalue()), report)
    _assert_valid_package(data)
    assert sum(report.dropped.values()) == 2
    assert "DICHIARAZIONE DI TRADUZIONE" in _texts(data)


def test_ready_check_only_asks_for_fields_the_template_has():
    from app.services.review_checks import check_certification

    doc = Document()
    _complex_field(doc.add_paragraph("Translated by "), [" MERGEFIELD Translator "], "«Translator»")
    data = docx_certification.add_certification(_translation(), _content(template_docx=_save(doc)))
    assert set(docx_certification.read_fields(data)) == {"translator", "date_iso"}
    assert check_certification(None, data, True) == []


def test_legacy_substitution_fills_merge_fields_too():
    from app.services.cert_template_service import substitute_in_docx

    out = substitute_in_docx(_template(), {"translator_name": "Anna Rossi", "date_long": "29 September 2026", "document_name": "x.pdf", "source_language": "Italian", "target_language": "English", "page_count": "2"})
    doc = Document(io.BytesIO(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Io sottoscritto Anna Rossi, in data 29 September 2026, ho tradotto dal Italian verso English" in text
    assert "«Timbro»" in text


# --- API --------------------------------------------------------------------


def _upload(client, owner, data: bytes, name="certificazione.docx"):
    r = client.post(
        "/certifications/upload",
        headers=owner["headers"],
        files={"file": (name, data, "application/octet-stream")},
        data={"kind": "SWORN_DECLARATION"},
    )
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture()
def owner_project(db, storage, make_user, make_project):
    owner = make_user()
    project = make_project(owner)
    key = f"uploads/{uuid.uuid4()}_authored.docx"
    storage["objects"][key] = _translation()
    project.authored_docx_s3_key = key
    db.commit()
    return owner, project


def test_template_check_preview_and_default(client, db, storage, owner_project, make_user):
    owner, project = owner_project
    storage["objects"]["uploads/stamp.png"] = _png((0, 0, 0), (120, 120)).getvalue()
    owner["team"].stamp_s3_key = "uploads/stamp.png"
    db.commit()
    cert = _upload(client, owner, _template(images=True, extras=True))
    assert cert["is_template"] and not cert["is_default"]

    r = client.get(f"/certifications/{cert['id']}/check", headers=owner["headers"])
    assert r.status_code == 200, r.text
    check = r.json()
    by_name = {f["name"]: f for f in check["fields"]}
    assert by_name["Traduttore"]["value"] == "Test User" and by_name["Traduttore"]["label"] == "Translator"
    assert by_name["Pagine"]["value"] == "2" and by_name["Documento"]["value"] == "Birth certificate"
    assert {u["name"] for u in check["unknown"]} == {"Timbro", "codice_albo"}
    assert {d["what"] for d in check["dropped"]} == {"Text boxes and shapes", "List numbering and bullets"}
    assert any("stamp" in n for n in check["notes"])

    r = client.get(f"/certifications/{cert['id']}/preview", headers=owner["headers"])
    assert r.status_code == 200
    _assert_valid_package(r.content)
    assert "DICHIARAZIONE DI TRADUZIONE" in _texts(r.content)

    stranger = make_user()
    assert client.get(f"/certifications/{cert['id']}/check", headers=stranger["headers"]).status_code == 404
    assert client.put(f"/certifications/{cert['id']}/default", headers=stranger["headers"], json={"is_default": True}).status_code == 404

    pdf = _upload(client, owner, b"%PDF-1.4 x", name="invoice.pdf")
    assert client.put(f"/certifications/{pdf['id']}/default", headers=owner["headers"], json={"is_default": True}).status_code == 400
    assert client.get(f"/certifications/{pdf['id']}/check", headers=owner["headers"]).status_code == 400

    other = _upload(client, owner, _template(), name="other.docx")
    assert client.put(f"/certifications/{other['id']}/default", headers=owner["headers"], json={"is_default": True}).json()["is_default"]
    assert client.put(f"/certifications/{cert['id']}/default", headers=owner["headers"], json={}).json()["is_default"]
    items = {c["id"]: c for c in client.get("/certifications", headers=owner["headers"]).json()["items"]}
    assert items[cert["id"]]["is_default"] and not items[other["id"]]["is_default"]

    # The editor's Certification button sends no template id: the team default is used.
    doc = client.get(f"/projects/{project.id}/document", headers=owner["headers"])
    v = int(doc.headers["X-Document-Version"])
    r = client.post(f"/projects/{project.id}/document/certification", headers=owner["headers"], json={"version": v})
    assert r.status_code == 200, r.text
    v = r.json()["version"]
    data = client.get(f"/projects/{project.id}/document", headers=owner["headers"]).content
    assert "DICHIARAZIONE DI TRADUZIONE" in _texts(data)
    assert "STUDIO ROSSI TRADUZIONI" in _texts(data)
    _assert_valid_package(data)

    r = client.put(
        f"/projects/{project.id}/document/certification",
        headers=owner["headers"],
        json={"version": v, "fields": {"translator": "Anna Verdi", "pages": "3"}},
    )
    assert r.status_code == 200, r.text
    assert r.json()["fields"]["translator"] == "Anna Verdi" and r.json()["fields"]["pages"] == "3"

    out = client.get(f"/projects/{project.id}/export", headers=owner["headers"])
    assert out.status_code == 200, out.text
    texts = _texts(out.content)
    assert texts.count("DICHIARAZIONE DI TRADUZIONE") == 1
    assert texts.index("Mr. BIANCHI LUCA, born in Bari") < texts.index("DICHIARAZIONE DI TRADUZIONE")
    assert any("Io sottoscritto Anna Verdi" in t for t in texts)
    body = _xml(out.content)
    with zipfile.ZipFile(io.BytesIO(out.content)) as z:
        names = set(z.namelist())
        rels = {r.get("Id"): r for r in etree.fromstring(z.read("word/_rels/document.xml.rels"))}
        styles = z.read("word/styles.xml")
    for blip in body.iter(f"{{{A}}}blip"):
        assert posixpath.normpath(posixpath.join("word", rels[blip.get(f"{{{R}}}embed")].get("Target"))) in names
    for link in body.iter(f"{{{W}}}hyperlink"):
        assert rels[link.get(f"{{{R}}}id")].get("TargetMode") == "External"
    for ps in body.iter(f"{{{W}}}pStyle"):
        if ps.get(f"{{{W}}}val").startswith(docx_cert_template.STYLE_PREFIX):
            assert f'w:styleId="{ps.get(f"{{{W}}}val")}"'.encode() in styles


def test_export_appends_the_default_template_when_the_document_has_no_page(client, owner_project):
    owner, project = owner_project
    cert = _upload(client, owner, _template())
    client.put(f"/certifications/{cert['id']}/default", headers=owner["headers"], json={"is_default": True})
    client.get(f"/projects/{project.id}/document", headers=owner["headers"])
    out = client.get(f"/projects/{project.id}/export", headers=owner["headers"])
    assert out.status_code == 200, out.text
    texts = _texts(out.content)
    assert texts.count("DICHIARAZIONE DI TRADUZIONE") == 1
    assert "CERTIFIED TRANSLATION" not in texts
    assert any(t.startswith("Io sottoscritto Test User") for t in texts)


def test_download_url_is_signed_and_team_scoped(client, owner_project, make_user):
    owner, _ = owner_project
    cert = _upload(client, owner, b"%PDF-1.4 invoice", name="invoice-KVNLZE4L-0002.pdf")
    url = f"/certifications/{cert['id']}/download-url"
    assert client.get(url).status_code == 401
    r = client.get(url, headers=owner["headers"])
    assert r.status_code == 200, r.text
    assert r.json()["file_name"] == "invoice-KVNLZE4L-0002.pdf"
    assert r.json()["url"].startswith("https://storage.test/uploads/")
    assert client.get(url, headers=make_user()["headers"]).status_code == 404
    assert client.get("/certifications/not-a-uuid/download-url", headers=owner["headers"]).status_code == 404


def test_italian_template_gets_italian_values_even_for_english_target():
    import io

    from docx import Document

    from app.services import cert_locale, cert_page

    doc = Document()
    doc.add_paragraph("Io sottoscritto traduttore dichiaro che la traduzione del documento dal testo originale")
    doc.add_paragraph("è fedele e conforme all'originale, composto da pagine, per uso della persona interessata.")
    buf = io.BytesIO()
    doc.save(buf)
    assert cert_page.page_language("en-GB", buf.getvalue()) == "it"
    assert cert_page.page_language("en-GB", None) == "en"
    assert cert_locale.detect_language("I hereby certify that the attached translation of the document is true and accurate") == "en"
    assert cert_locale.detect_language("short text") is None
