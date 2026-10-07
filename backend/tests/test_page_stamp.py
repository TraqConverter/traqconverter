"""Copying pictures, the editable page stamp and the certification page's pictures."""
import io
import uuid
import zipfile

import pytest
from docx import Document
from lxml import etree

from app.services import docx_blocks, docx_certification, docx_images, docx_page_stamp
from tests.test_document_media import (  # noqa: F401  (project_with_doc is a fixture)
    W,
    WP,
    _body_xml,
    _content,
    _docx,
    _export,
    _get,
    _png,
    _tagged,
    _texts,
    _upload,
    project_with_doc,
)


def _media_files(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [n for n in z.namelist() if n.startswith("word/media/")]


# --- copy and paste -------------------------------------------------------------


def test_duplicate_reuses_the_media_part_and_gets_a_new_id():
    data, ids = _tagged("Title", "Body", "End")
    data, image_id = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[0], "inline", 3, "center")
    data, copy_id = docx_images.duplicate_image(data, image_id, ids[2])
    assert copy_id != image_id and docx_blocks.IMAGE_ID_RE.match(copy_id)
    info = {i["id"]: i for i in docx_images.list_images(data)}
    assert set(info) == {image_id, copy_id}
    assert info[copy_id]["width_cm"] == 3.0 and not info[copy_id]["floating"]
    assert len(_media_files(data)) == 1
    body = _body_xml(data)
    assert len({b.get(f"{{{docx_blocks.R_NS}}}embed") for b in body.iter(f"{{{docx_blocks.A_NS}}}blip")}) == 1
    docpr = [d.get("id") for d in body.iter(f"{{{WP}}}docPr")]
    assert len(docpr) == len(set(docpr))
    texts = _texts(data)
    assert texts[texts.index("End") + 1] == "", "the copy gets its own paragraph right after the target"
    copy_p = next(p for p in body.iter(f"{{{W}}}p") if any(d.get("name") == copy_id for d in p.iter(f"{{{WP}}}docPr")))
    assert copy_p.find(f"{{{W}}}pPr/{{{W}}}jc").get(f"{{{W}}}val") == "center"
    assert [b.get(f"{{{W}}}name") for b in body.iter(f"{{{W}}}bookmarkStart")].count("_" + copy_id) == 1
    Document(io.BytesIO(data))


def test_duplicate_keeps_floating_offsets():
    data, ids = _tagged("Title", "Body", "End")
    data, image_id = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[0], "after", 3)
    data = docx_images.float_image(data, image_id, ids[1], 1_000_000, -200_000)
    pasted, copy_id = docx_images.duplicate_image(data, image_id, ids[2])
    info = {i["id"]: i for i in docx_images.list_images(pasted)}
    assert info[copy_id]["floating"] and info[copy_id]["block_id"] == ids[2]
    assert (info[copy_id]["x_emu"], info[copy_id]["y_emu"]) == (1_000_000, -200_000)
    assert info[image_id]["block_id"] == ids[1]
    assert len(_media_files(pasted)) == 1
    # Onto its own paragraph the copy is shifted so it doesn't hide under the original.
    same, copy_id = docx_images.duplicate_image(data, image_id, ids[1])
    copy = next(i for i in docx_images.list_images(same) if i["id"] == copy_id)
    assert copy["x_emu"] > 1_000_000 and copy["y_emu"] > -200_000
    with pytest.raises(docx_blocks.DocxEditError):
        docx_images.duplicate_image(data, "img_00000000", ids[2])


def test_copy_to_pages_adds_one_per_page_and_skips_the_certification_page():
    data, ids = _tagged("Page one", "Page two", "Page three")
    data, image_id = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[0], "inline", 3)
    data = docx_certification.add_certification(data, _content())
    texts = docx_blocks.paragraph_texts(data, docx_blocks.block_ids(data))
    cert_block = next(b for b, t in texts.items() if t == "CERTIFIED TRANSLATION")
    source_block = docx_images.list_images(data)[0]["block_id"]
    targets = [(ids[1], 500_000, 100_000), (ids[2], 500_000, 100_000), (ids[2], 1, 1), (cert_block, 0, 0), (source_block, 0, 0)]
    out, made = docx_images.copy_to_blocks(data, image_id, targets)
    assert len(made) == 2
    info = {i["id"]: i for i in docx_images.list_images(out)}
    assert [info[m]["block_id"] for m in made] == [ids[1], ids[2]]
    assert all(info[m]["floating"] and (info[m]["x_emu"], info[m]["y_emu"]) == (500_000, 100_000) for m in made)
    assert _media_files(out) == _media_files(data)
    again, more = docx_images.copy_to_blocks(out, image_id, targets)
    assert more == [] and again is out, "pages already showing the picture are skipped"


def test_duplicate_and_copy_to_pages_api_are_versioned_with_one_undo(client, project_with_doc):
    owner, project = project_with_doc()
    data, v = _get(client, owner, project)
    ids = docx_blocks.block_ids(data)
    url = f"/projects/{project.id}/document"
    r = _upload(client, owner, project, v, ids[0], position="inline", width_cm=3)
    image_id, v = r.json()["image_id"], r.json()["version"]

    r = client.post(f"{url}/images/{image_id}/duplicate", headers=owner["headers"], json={"version": v, "target_block_id": ids[2]})
    assert r.status_code == 200, r.text
    copy_id, v = r.json()["image_id"], r.json()["version"]
    data, _ = _get(client, owner, project)
    assert {i["id"] for i in docx_images.list_images(data)} == {image_id, copy_id}
    stale = client.post(f"{url}/images/{image_id}/duplicate", headers=owner["headers"], json={"version": v - 1, "target_block_id": ids[2]})
    assert stale.status_code == 409

    targets = [{"block_id": b, "x_emu": 200000, "y_emu": 0} for b in ids[1:]]
    r = client.post(f"{url}/images/{image_id}/copy-to-pages", headers=owner["headers"], json={"version": v, "targets": targets})
    assert r.status_code == 200, r.text
    assert r.json()["version"] == v + 1 and len(r.json()["image_ids"]) == 2
    data, _ = _get(client, owner, project)
    assert len(docx_images.list_images(data)) == 4
    assert client.post(f"{url}/undo", headers=owner["headers"], json={}).status_code == 200
    data, _ = _get(client, owner, project)
    assert {i["id"] for i in docx_images.list_images(data)} == {image_id, copy_id}
    assert client.post(f"{url}/undo", headers=owner["headers"], json={}).status_code == 200
    data, _ = _get(client, owner, project)
    assert [i["id"] for i in docx_images.list_images(data)] == [image_id]


# --- page stamp -------------------------------------------------------------------


def _footer_stamps(data: bytes) -> list[tuple[str, str]]:
    """(footer part, paragraph alignment) of each page stamp picture."""
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in z.namelist():
            if not name.startswith("word/footer"):
                continue
            for d in etree.fromstring(z.read(name)).iter(f"{{{WP}}}docPr"):
                if d.get("name") == docx_page_stamp.STAMP_NAME:
                    p = d
                    while p.tag != f"{{{W}}}p":
                        p = p.getparent()
                    out.append((name, p.find(f"{{{W}}}pPr/{{{W}}}jc").get(f"{{{W}}}val")))
    return out


def _stamp_image():
    return docx_images.prepare_image(_png((200, 200), paper=True), remove_background=True)


def test_page_stamp_goes_in_every_section_footer():
    doc = Document()
    doc.add_paragraph("One")
    doc.add_section()
    doc.add_paragraph("Two")
    buf = io.BytesIO()
    doc.save(buf)
    data = docx_blocks.tag_blocks(docx_page_stamp.install(buf.getvalue(), _stamp_image(), "center"))
    assert docx_page_stamp.is_marked(data)
    assert _footer_stamps(data) == [("word/footer1.xml", "center")]
    refs = [s.find(f"{{{W}}}footerReference") for s in _body_xml(data).iter(f"{{{W}}}sectPr")]
    assert len(refs) == 2 and all(r is not None for r in refs)
    assert docx_page_stamp.state(data) == {"managed": True, "enabled": True, "align": "center", "width_cm": 3.0, "asset_id": None}
    assert docx_images.list_images(data) == [], "the footer stamp isn't one of the body pictures"
    assert len(_footer_stamps(docx_page_stamp.install(data, _stamp_image()))) == 1
    Document(io.BytesIO(data))

    word = Document(io.BytesIO(_docx("Body")))
    word.sections[0].footer.paragraphs[0].text = "Page footer text"
    buf = io.BytesIO()
    word.save(buf)
    out = docx_page_stamp.install(buf.getvalue(), _stamp_image(), "left")
    assert len(_footer_stamps(out)) == 1
    assert Document(io.BytesIO(out)).sections[0].footer.paragraphs[0].text == "Page footer text"

    unmarked = docx_page_stamp.install(_docx("x"), None)
    assert docx_page_stamp.is_marked(unmarked) and _footer_stamps(unmarked) == []
    assert not docx_page_stamp.is_marked(_docx("x"))


def test_body_and_footer_share_the_stamp_file_safely():
    stamp = _stamp_image()
    data = docx_blocks.tag_blocks(docx_page_stamp.install(_docx("A", "B"), stamp, "right"))
    data = docx_certification.add_certification(data, _content(stamp=stamp))
    removed = docx_certification.remove_certification(data)
    assert len(_media_files(removed)) == 1 and docx_page_stamp.state(removed)["enabled"]
    off = docx_page_stamp.update(data, enabled=False)
    assert len(_media_files(off)) == 1 and len(docx_images.list_images(off)) == 1
    assert _media_files(docx_page_stamp.update(removed, enabled=False)) == []


def _team_stamp(db, storage, owner):
    from tests.test_media_library import add_media

    return add_media(db, storage, owner, "stamp", None)


def test_new_editor_document_gets_the_team_stamp_and_the_controls_change_it(client, db, storage, project_with_doc):
    owner, project = project_with_doc()
    stamp = _team_stamp(db, storage, owner)
    data, v = _get(client, owner, project)
    assert docx_page_stamp.is_marked(data)
    assert [a for _, a in _footer_stamps(data)] == ["right"]
    url = f"/projects/{project.id}/document/page-stamp"
    state = client.get(url, headers=owner["headers"]).json()
    assert [s["id"] for s in state.pop("stamps")] == [str(stamp.id)]
    assert state == {
        "version": v, "available": True, "managed": True, "enabled": True, "align": "right", "width_cm": 3.0,
        "asset_id": str(stamp.id),
    }

    r = client.put(url, headers=owner["headers"], json={"version": v, "align": "right", "width_cm": 5})
    assert r.status_code == 200, r.text
    assert (r.json()["align"], r.json()["width_cm"], r.json()["version"]) == ("right", 5.0, v + 1)
    assert client.put(url, headers=owner["headers"], json={"version": v + 1, "width_cm": 9}).status_code == 422
    assert client.put(url, headers=owner["headers"], json={"version": v, "align": "left"}).status_code == 409

    r = client.put(url, headers=owner["headers"], json={"version": v + 1, "enabled": False})
    assert r.status_code == 200 and r.json()["enabled"] is False
    data, _ = _get(client, owner, project)
    assert _footer_stamps(data) == []
    r = client.put(url, headers=owner["headers"], json={"version": v + 2, "enabled": True, "align": "center", "width_cm": 4})
    assert r.status_code == 200, r.text
    assert (r.json()["enabled"], r.json()["align"], r.json()["width_cm"]) == (True, "center", 4.0)

    assert client.post(f"/projects/{project.id}/document/undo", headers=owner["headers"], json={}).status_code == 200
    assert client.get(url, headers=owner["headers"]).json()["enabled"] is False


def _stamped_export(client, owner, project, monkeypatch, tmp_path):
    """Footer pictures in the export, and whether the old export-time stamp was looked up."""
    from app.services import export_service

    path = tmp_path / "team_stamp.png"
    path.write_bytes(_png((90, 90), color=(10, 10, 200)))
    calls = []

    def resolve(_project, _tmp):
        calls.append(1)
        return str(path), "right"

    monkeypatch.setattr(export_service, "_resolve_team_stamp", resolve)
    out = _export(client, owner, project)
    pictures = 0
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        for name in z.namelist():
            if name.startswith("word/footer"):
                pictures += len(list(etree.fromstring(z.read(name)).iter(f"{{{W}}}drawing")))
    return pictures, calls


def test_export_adds_no_second_stamp_to_a_marked_document(client, db, storage, project_with_doc, monkeypatch, tmp_path):
    owner, project = project_with_doc()
    _team_stamp(db, storage, owner)
    _, v = _get(client, owner, project)
    pictures, calls = _stamped_export(client, owner, project, monkeypatch, tmp_path)
    assert pictures == 1 and calls == []

    r = client.put(f"/projects/{project.id}/document/page-stamp", headers=owner["headers"], json={"version": v, "enabled": False})
    assert r.status_code == 200
    pictures, calls = _stamped_export(client, owner, project, monkeypatch, tmp_path)
    assert pictures == 0 and calls == []


def test_exported_editor_stamp_stays_off_the_original_pages(storage, tmp_path):
    from types import SimpleNamespace

    from app.services.export_wrapper import build_full_export_docx
    from tests.conftest import make_pdf

    storage["objects"]["uploads/src.pdf"] = make_pdf(tmp_path / "src.pdf", pages=2).read_bytes()
    project = SimpleNamespace(file_name="src.pdf", file_path="uploads/src.pdf", team_id=None)
    data = docx_page_stamp.install(_docx("Translated text"), _stamp_image(), "center", 4)
    out = build_full_export_docx(docx_blocks.strip_blocks(data), project, None, append_certification=False).getvalue()
    sects = list(_body_xml(out).iter(f"{{{W}}}sectPr"))
    assert len(sects) == 3, "two original pages and the translation"
    assert [s.find(f"{{{W}}}footerReference") is not None for s in sects] == [False, False, True]
    word = Document(io.BytesIO(out))
    footer = word.sections[-1].footer
    assert footer.paragraphs[0].alignment == 1 and len(footer.paragraphs[0].runs[0].element.findall(f".//{{{W}}}drawing")) == 1
    assert footer._element.find(f".//{{{WP}}}extent").get("cx") == str(4 * 360000)


def test_old_editor_documents_keep_the_export_stamp(client, db, storage, project_with_doc, monkeypatch, tmp_path):
    from app.models.document_version import DocumentVersion

    owner, project = project_with_doc()
    key = "uploads/old_editor.docx"
    storage["objects"][key] = docx_blocks.tag_blocks(_docx("Old translation"))
    project.authored_docx_s3_key = key
    project.document_version = 1
    db.add(DocumentVersion(project_id=project.id, version=1, s3_key=key, note="Initial translation"))
    db.commit()
    _team_stamp(db, storage, owner)
    data, v = _get(client, owner, project)
    assert not docx_page_stamp.is_marked(data)
    url = f"/projects/{project.id}/document/page-stamp"
    assert client.get(url, headers=owner["headers"]).json()["managed"] is False
    assert client.put(url, headers=owner["headers"], json={"version": v, "align": "left"}).status_code == 409
    pictures, calls = _stamped_export(client, owner, project, monkeypatch, tmp_path)
    assert pictures == 1 and calls == [1]


def test_editable_copies_get_no_page_stamp(client, db, storage, project_with_doc):
    owner, project = project_with_doc()
    _team_stamp(db, storage, owner)
    project.mode = "dtp"
    db.commit()
    data, v = _get(client, owner, project)
    assert not docx_page_stamp.is_marked(data) and _footer_stamps(data) == []
    state = client.get(f"/projects/{project.id}/document/page-stamp", headers=owner["headers"]).json()
    assert state["available"] is False and state["managed"] is False
    r = client.put(f"/projects/{project.id}/document/page-stamp", headers=owner["headers"], json={"version": v, "enabled": True})
    assert r.status_code == 409


# --- certification page pictures ---------------------------------------------------


def test_certification_logo_and_stamp_are_editor_pictures_and_are_not_duplicated():
    data, ids = _tagged("CERTIFICATE OF RESIDENCE", "Body")
    logo = docx_images.prepare_image(_png())
    stamp = _stamp_image()
    data = docx_certification.add_certification(data, _content(logo=logo, stamp=stamp, template_ref="standard"))
    roles = {}
    for docpr in _body_xml(data).iter(f"{{{WP}}}docPr"):
        assert docx_blocks.IMAGE_ID_RE.match(docpr.get("name"))
        run = docpr.getparent().getparent().getparent()
        assert run.getprevious().getprevious().get(f"{{{W}}}name") == "_" + docpr.get("name")
        roles[docpr.get("descr")] = docpr.get("name")
    assert set(roles) == {docx_certification.LOGO_ROLE, docx_certification.STAMP_ROLE}
    logo_id = roles[docx_certification.LOGO_ROLE]

    # Resized and floated in the editor, then the template is switched: still one logo, at the user's size.
    data = docx_images.resize_image(data, logo_id, 6)
    own = next(i["block_id"] for i in docx_images.list_images(data) if i["id"] == logo_id)
    data = docx_images.float_image(data, logo_id, own, 100000, 0)
    switched = docx_certification.add_certification(data, _content(logo=logo, stamp=stamp, template_ref=str(uuid.uuid4())))
    pictures = docx_images.list_images(switched)
    assert len(pictures) == 2 and 6.0 in [p["width_cm"] for p in pictures]
    descr = sorted(d.get("descr") for d in _body_xml(switched).iter(f"{{{WP}}}docPr"))
    assert descr == sorted([docx_certification.LOGO_ROLE, docx_certification.STAMP_ROLE])

    # Dragged out onto the translation: the rebuilt page doesn't add another one.
    moved = docx_images.move_image(data, logo_id, ids[1], "after")
    rebuilt = docx_certification.add_certification(moved, _content(logo=logo, stamp=stamp, template_ref="standard"))
    descr = [d.get("descr") for d in _body_xml(rebuilt).iter(f"{{{WP}}}docPr")]
    assert descr.count(docx_certification.LOGO_ROLE) == 1 and descr.count(docx_certification.STAMP_ROLE) == 1
    assert logo_id in [i["id"] for i in docx_images.list_images(rebuilt)]


def test_certification_template_switch_via_api_keeps_one_logo(client, db, storage, project_with_doc):
    from tests.test_document_media import ITALIAN, _cert_template

    owner, project = project_with_doc()
    from tests.test_media_library import add_media

    add_media(db, storage, owner, "logo", None)
    italian = _cert_template(db, storage, owner, *ITALIAN, name="italiano.docx")
    url = f"/projects/{project.id}/document/certification"
    _, v = _get(client, owner, project)
    v = client.post(url, headers=owner["headers"], json={"version": v}).json()["version"]
    for template in (str(italian.id), "standard", str(italian.id)):
        r = client.put(url, headers=owner["headers"], json={"version": v, "template_id": template})
        assert r.status_code == 200, r.text
        v = r.json()["version"]
        data, _ = _get(client, owner, project)
        logos = [d for d in _body_xml(data).iter(f"{{{WP}}}docPr") if d.get("descr") == docx_certification.LOGO_ROLE]
        assert len(logos) == 1
