import io

import pytest
from docx import Document
from docx.shared import Pt

from app.services import docx_blocks as db


def _sample() -> bytes:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Born in ")
    b = p.add_run("Bari")
    b.bold = True
    p.add_run(" on 12/03/1987")
    t = doc.add_paragraph()
    t.add_run("Ref.")
    t.add_run().add_tab()
    t.add_run("0001")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "[Signature]"
    table.cell(0, 1).text = "[Round stamp: Registry Office]"
    doc.sections[0].header.paragraphs[0].text = "MUNICIPALITY OF ESEMPIO"
    buf = io.BytesIO()
    doc.save(buf)
    return db.tag_blocks(buf.getvalue())


def _texts(data):
    doc = Document(io.BytesIO(data))
    return [p.text for p in doc.paragraphs]


def test_every_paragraph_gets_one_unique_id_and_tagging_is_stable():
    data = _sample()
    ids = db.block_ids(data)
    assert len(ids) == len(set(ids)) >= 5
    assert db.block_ids(db.tag_blocks(data)) == ids


def test_text_edit_keeps_formatting_outside_the_change():
    data = _sample()
    first = db.block_ids(data)[0]
    out, changed = db.apply_text_edits(data, [(first, "Born in Bari on 13/03/1987")])
    assert changed == [first]
    runs = Document(io.BytesIO(out)).paragraphs[0].runs
    assert "".join(r.text for r in runs) == "Born in Bari on 13/03/1987"
    assert any(r.bold and r.text == "Bari" for r in runs)


def test_text_edit_preserves_tabs_and_handles_edges():
    data = _sample()
    second = db.block_ids(data)[1]
    out, _ = db.apply_text_edits(data, [(second, "Ref. no.\t0002")])
    assert _texts(out)[1] == "Ref. no.\t0002"
    out, _ = db.apply_text_edits(out, [(second, "X\tRef. no.\t0002")])
    assert _texts(out)[1] == "X\tRef. no.\t0002"
    out, _ = db.apply_text_edits(out, [(second, "")])
    assert _texts(out)[1] == ""
    out, _ = db.apply_text_edits(out, [(second, "restored")])
    assert _texts(out)[1] == "restored"


def test_outline_lists_body_tables_and_header():
    text = db.outline(_sample())
    assert "table 1x2" in text and "[Signature]" in text
    assert "(header" in text and "MUNICIPALITY OF ESEMPIO" in text


def test_replace_operation_restructures_and_retags():
    data = _sample()
    ids = db.block_ids(data)
    uids, xml = db.units_for_blocks(data, [ids[0]], neighbours=0)
    assert uids == [ids[0]] and "<w:p" in xml[ids[0]]
    new_xml = (
        f'<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Born in Bari</w:t></w:r></w:p>'
        f'<w:p><w:r><w:t xml:space="preserve">on 12/03/1987</w:t></w:r></w:p>'
    )
    out, changed = db.apply_operations(data, [{"op": "replace", "target": ids[0], "xml": new_xml}])
    assert _texts(out)[:2] == ["Born in Bari", "on 12/03/1987"]
    assert len(changed) == 2
    all_ids = db.block_ids(out)
    assert len(all_ids) == len(set(all_ids))


def test_copied_bookmarks_are_deduplicated():
    data = _sample()
    ids = db.block_ids(data)
    _, xml = db.units_for_blocks(data, [ids[0]], neighbours=0)
    out, _ = db.apply_operations(data, [{"op": "insert_after", "target": ids[0], "xml": xml[ids[0]]}])
    all_ids = db.block_ids(out)
    assert len(all_ids) == len(set(all_ids)) == len(ids) + 1


def test_table_cell_target_resolves_to_whole_table_and_delete_works():
    data = _sample()
    cell_block = [b for b in db.block_ids(data)][2]
    out, _ = db.apply_operations(data, [{"op": "delete", "target": cell_block}])
    assert len(Document(io.BytesIO(out)).tables) == 0


@pytest.mark.parametrize("bad", [
    '<w:p><w:r><w:drawing/></w:r></w:p>',
    '<w:p xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:hyperlink r:id="rId9"/></w:p>',
    '<w:sectPr/>',
    '<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><w:p><w:r><w:t>&e;</w:t></w:r></w:p>',
    '<w:p><w:r><w:t>unclosed</w:r></w:p>',
    '<script/>',
])
def test_unsafe_or_invalid_fragments_rejected(bad):
    data = _sample()
    target = db.block_ids(data)[0]
    with pytest.raises(db.DocxEditError):
        db.apply_operations(data, [{"op": "replace", "target": target, "xml": bad}])


def test_unknown_target_rejected():
    with pytest.raises(db.DocxEditError):
        db.apply_operations(_sample(), [{"op": "delete", "target": "_b00000000"}])


def test_strip_blocks_removes_editor_bookmarks():
    assert db.block_ids(db.strip_blocks(_sample())) == []


def test_fields_and_footnote_refs_are_left_untouched():
    from lxml import etree

    data = _sample()
    doc = db._Doc.load(data)
    p = next(doc.trees["word/document.xml"].iter(db.w("p")))
    field = etree.fromstring(
        f'<w:root xmlns:w="{db.W_NS}">'
        '<w:r><w:t xml:space="preserve"> page </w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>PAGE</w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>7</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t xml:space="preserve"> end</w:t></w:r>'
        '</w:root>'
    )
    for child in list(field):
        p.append(child)
    data = doc.dump()
    first = db.block_ids(data)[0]
    before = db.paragraph_text(next(db._Doc.load(data).trees["word/document.xml"].iter(db.w("p"))))
    assert before.endswith("1987 page  end")
    out, _ = db.apply_text_edits(data, [(first, before.replace(" end", " finish"))])
    xml = db._Doc.load(out).files["word/document.xml"].decode()
    assert ">7<" in xml and "PAGE" in xml and "finish" in xml


def test_set_text_inside_a_table_cell_keeps_the_rest_of_the_table():
    data = _sample()
    cell = db.block_ids(data)[2]
    out, changed = db.apply_operations(data, [{"op": "set_text", "target": cell, "content": "[Illegible signature]"}])
    table = Document(io.BytesIO(out)).tables[0]
    assert table.cell(0, 0).text == "[Illegible signature]"
    assert table.cell(0, 1).text == "[Round stamp: Registry Office]"
    assert changed == [cell]


def test_replace_paragraph_restyles_one_cell_paragraph():
    data = _sample()
    cell = db.block_ids(data)[3]
    xml = '<w:p><w:pPr><w:jc w:val="right"/></w:pPr><w:r><w:rPr><w:b/></w:rPr><w:t>[Round stamp: Registry Office]</w:t></w:r></w:p>'
    out, changed = db.apply_operations(data, [{"op": "replace_paragraph", "target": cell, "content": xml}])
    cell_p = Document(io.BytesIO(out)).tables[0].cell(0, 1).paragraphs[0]
    assert cell_p.runs[0].bold and cell_p.alignment == 2
    assert len(changed) == 1
    with pytest.raises(db.DocxEditError):
        db.apply_operations(data, [{"op": "replace_paragraph", "target": cell, "content": "<w:tbl/>"}])


def test_paragraph_texts_for_selected_ids():
    data = _sample()
    ids = db.block_ids(data)
    assert db.paragraph_texts(data, ids[:2]) == {ids[0]: "Born in Bari on 12/03/1987", ids[1]: "Ref.\t0001"}
