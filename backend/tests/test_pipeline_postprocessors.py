import io
import re
import zipfile

from docx import Document

from app.services import claude_authored_rebuild as car
from app.services import script_sandbox


def _docx(paragraphs):
    doc = Document()
    for para in paragraphs:
        if isinstance(para, str):
            doc.add_paragraph(para)
        else:
            p = doc.add_paragraph()
            for run_text in para:
                p.add_run(run_text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _texts(docx_bytes):
    return [p.text for p in Document(io.BytesIO(docx_bytes)).paragraphs]


REGISTRAR = "I hereby certify that this is a true extract from the civil status register of the Municipality."


def test_cert_strip_keeps_registrar_attestation_and_removes_translator_block():
    src = _docx([
        "MUNICIPALITY OF ESEMPIO",
        REGISTRAR,
        "Signature: ________",
        "Body paragraph 01/01/1990.",
        "CERTIFIED TRANSLATION",
        "I hereby certify that I am competent to translate from Italian into English "
        "and that this translation is accurate and complete.",
        "Translator: jane@example.com",
        "Signature: ________",
        "Date: 2026-01-01 10:00 UTC",
    ])
    out = _texts(car._strip_inline_cert_blocks(src))
    assert out == [
        "MUNICIPALITY OF ESEMPIO",
        REGISTRAR,
        "Signature: ________",
        "Body paragraph 01/01/1990.",
    ]


def test_cert_strip_leaves_document_without_translator_block_untouched():
    src = _docx(["CERTIFICATE OF RESIDENCE", REGISTRAR, "The Registrar"])
    assert car._strip_inline_cert_blocks(src) is src


def test_is_translator_cert_text():
    assert not car.is_translator_cert_text(REGISTRAR)
    assert car.is_translator_cert_text("CERTIFIED TRANSLATION")
    assert car.is_translator_cert_text(
        "I hereby certify that the foregoing is a true and accurate translation of the attached document."
    )


def test_html_strip_removes_real_tags_only():
    src = _docx([
        '<p style="text-align: center;">AUSTRIAN EMBASSY<br>LONDON</p>',
        "[Stamp: Ministry of the Interior]",
        "[Signature]",
        "Contact: Mario Rossi <mario@example.com>",
        "Minimum grade > 18/30",
        "Values: a < b and c > d",
        "Name: {full_name}",
        ["<b>", "Bold heading", "</b>"],
    ])
    out = _texts(car._strip_html_and_bracket_artifacts(src))
    assert out == [
        "AUSTRIAN EMBASSY LONDON",
        "[Stamp: Ministry of the Interior]",
        "[Signature]",
        "Contact: Mario Rossi <mario@example.com>",
        "Minimum grade > 18/30",
        "Values: a < b and c > d",
        "Name:",
        "Bold heading",
    ]


def test_html_strip_is_noop_on_clean_document():
    src = _docx(["Grade 28/30 > pass mark", "[Seal]"])
    assert car._strip_html_and_bracket_artifacts(src) is src


def test_force_borderless_table_styles_nils_table_grid():
    doc = Document()
    doc.add_table(rows=1, cols=1).style = "Table Grid"
    buf = io.BytesIO()
    doc.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
        styles = z.read("word/styles.xml").decode("utf-8")

    new, changed = car._force_borderless_table_styles(styles, "")
    assert changed >= 1
    grid = re.search(r'<w:style\b[^>]*w:styleId="TableGrid".*?</w:style>', new, re.S).group(0)
    assert 'w:val="single"' not in grid
    assert '<w:insideV w:val="nil"/>' in grid
    from xml.etree import ElementTree as ET
    ET.fromstring(new.encode("utf-8"))
    assert car._force_borderless_table_styles(new, "") == (new, 0)


def test_author_prompt_has_no_real_personal_data_and_valid_recipe():
    prompt = car._AUTHOR_PROMPT_TEMPLATE.format(
        source_lang="Italian", target_lang="English", output_path="/tmp/out/rebuild.docx",
        image_list="", table_list="", sandbox_rules=car.SANDBOX_RULES, notation_rules=car.NOTATION_RULES,
    )
    assert "MARIO ROSSI" in prompt
    # Example IDs must be zeroed placeholders, never real-looking numbers.
    assert not re.search(r"\b[1-9]\d{6,}\b", prompt)
    assert "Import only: docx" in prompt
    recipe = re.search(r"```python\n(.*?)```", prompt, re.S).group(1)
    script_sandbox.validate_script(recipe)
