"""The export (and so the PDF) keeps the editor document's page layout: page setup, styles, defaults, headers."""
import io
import shutil
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.services import docx_blocks, docx_page_stamp
from app.services.export_wrapper import SKIP_SOURCE_PAGES, build_full_export_docx
from tests.conftest import make_pdf


def _cm(*lengths) -> tuple:
    return tuple(round(x.cm, 2) for x in lengths)


def _authored(pages: int = 2, lines: int = 40) -> bytes:
    """An AI-style document: A4, its own margins, Times New Roman and tight spacing on Normal, a header, explicit page breaks."""
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(3), Cm(1.5), Cm(2.5), Cm(1.8)
    s.header.paragraphs[0].text = "UNIVERSITY OF EXAMPLE"
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0
    for page in range(pages):
        for i in range(lines):
            doc.add_paragraph(f"Page {page + 1} line {i + 1}: the undersigned certifies the records below.")
        if page < pages - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _project():
    return SimpleNamespace(file_name="src.pdf", file_path="uploads/src.pdf", team_id=None, id="p")


def _export(data: bytes, skip_source: bool = True) -> bytes:
    token = SKIP_SOURCE_PAGES.set(skip_source)
    try:
        return build_full_export_docx(docx_blocks.strip_blocks(data), _project(), None, append_certification=False).getvalue()
    finally:
        SKIP_SOURCE_PAGES.reset(token)


def test_export_keeps_page_setup_header_and_normal_style(storage):
    out = Document(io.BytesIO(_export(docx_page_stamp.install(docx_blocks.tag_blocks(_authored()), None))))
    s = out.sections[-1]
    assert _cm(s.page_width, s.page_height) == (21, 29.7), "A4, not python-docx's US Letter"
    assert _cm(s.top_margin, s.bottom_margin, s.left_margin, s.right_margin) == (3, 1.5, 2.5, 1.8)
    assert s.header.paragraphs[0].text == "UNIVERSITY OF EXAMPLE"
    normal = out.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size == Pt(11)
    assert normal.paragraph_format.space_after == Pt(2) and normal.paragraph_format.line_spacing == 1.0


def test_export_keeps_every_section_of_a_mixed_orientation_document(storage):
    doc = Document(io.BytesIO(_authored(pages=1, lines=3)))
    land = doc.add_section(WD_SECTION.NEW_PAGE)
    land.orientation = WD_ORIENT.LANDSCAPE
    land.page_width, land.page_height = Cm(29.7), Cm(21)
    doc.add_paragraph("Wide table page")
    back = doc.add_section(WD_SECTION.NEW_PAGE)
    back.orientation = WD_ORIENT.PORTRAIT
    back.page_width, back.page_height = Cm(21), Cm(29.7)
    doc.add_paragraph("Back to portrait")
    buf = io.BytesIO()
    doc.save(buf)
    out = Document(io.BytesIO(_export(docx_blocks.tag_blocks(buf.getvalue()))))
    assert [_cm(s.page_width, s.page_height) for s in out.sections] == [(21, 29.7), (29.7, 21), (21, 29.7)]


def test_original_pages_come_first_without_the_translation_header(storage, tmp_path):
    storage["objects"]["uploads/src.pdf"] = make_pdf(tmp_path / "src.pdf", pages=2).read_bytes()
    authored = _authored(pages=1, lines=3)
    doc = Document(io.BytesIO(authored))
    # A continuous first section would otherwise run the translation onto the last original page.
    sect = doc.sections[0]._sectPr
    kind = sect.makeelement(qn("w:type"), {qn("w:val"): "continuous"})
    sect.insert(0, kind)
    buf = io.BytesIO()
    doc.save(buf)
    out = Document(io.BytesIO(_export(buf.getvalue(), skip_source=False)))
    sects = out.sections
    assert len(sects) == 3, "two original pages and the translation"
    for original in sects[:2]:
        assert original._sectPr.find(qn("w:headerReference")) is None and original._sectPr.find(qn("w:footerReference")) is None
        assert _cm(original.page_width, original.top_margin) == (21, 1)
    assert sects[2]._sectPr.find(qn("w:type")) is None
    assert sects[2].header.paragraphs[0].text == "UNIVERSITY OF EXAMPLE"
    assert _cm(sects[2].page_width, sects[2].page_height, sects[2].top_margin) == (21, 29.7, 3)
    body = out.element.body
    pictures = [p for p in body.iter(qn("w:p")) if p.find(".//" + qn("w:drawing")) is not None]
    assert len(pictures) == 2 and all(p.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None for p in pictures)
    assert out.paragraphs[2].text.startswith("Page 1 line 1"), "the translation follows the original's pages"


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice isn't installed")
def test_pdf_has_the_pages_the_document_has(storage):
    import fitz

    from app.services.export_service import _convert_docx_to_pdf

    # Two explicit pages that each fit on A4 with the document's own spacing; python-docx's defaults would not.
    data = docx_blocks.tag_blocks(_authored(pages=2, lines=40))
    direct = fitz.open(stream=_convert_docx_to_pdf(docx_blocks.strip_blocks(data)), filetype="pdf")
    exported = fitz.open(stream=_convert_docx_to_pdf(_export(data)), filetype="pdf")
    assert len(direct) == 2
    assert len(exported) == len(direct)
    for a, b in zip(direct, exported):
        assert a.get_text().split("\n")[:3] == b.get_text().split("\n")[:3]
        assert (round(b.rect.width), round(b.rect.height)) == (595, 842)
