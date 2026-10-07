import io
import uuid
import zipfile
from datetime import date

import pytest
from docx import Document
from lxml import etree
from PIL import Image, ImageDraw

from app.services import cert_locale, docx_blocks, docx_certification, docx_images

W = docx_blocks.W_NS
WP = docx_blocks.WP_NS


def _png(size=(120, 60), color=(200, 30, 30), paper=False) -> bytes:
    img = Image.new("RGB", size, "white" if paper else color)
    if paper:
        ImageDraw.Draw(img).ellipse((5, 5, size[0] - 5, size[1] - 5), outline=color, width=6)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _docx(*paras: str) -> bytes:
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _body_xml(data: bytes) -> etree._Element:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return etree.fromstring(z.read("word/document.xml"))


def _texts(data):
    body = _body_xml(data).find(f"{{{W}}}body")
    return [docx_blocks.paragraph_text(p) for p in body.findall(f"{{{W}}}p")]


def _tagged(*paras):
    data = docx_blocks.tag_blocks(_docx(*paras))
    return data, docx_blocks.block_ids(data)


# --- document model ---------------------------------------------------------


def test_prepare_image_reencodes_and_rejects_non_images():
    jpeg = io.BytesIO()
    img = Image.new("RGB", (50, 40), (10, 20, 30))
    exif = Image.Exif()
    exif[0x010F] = "SecretCam"
    img.save(jpeg, "JPEG", exif=exif.tobytes())
    prepared = docx_images.prepare_image(jpeg.getvalue())
    assert prepared.ext == "jpeg" and (prepared.width_px, prepared.height_px) == (50, 40)
    assert b"SecretCam" not in prepared.data
    with pytest.raises(docx_images.ImageError):
        docx_images.prepare_image(b"<svg xmlns='http://www.w3.org/2000/svg'/>")
    with pytest.raises(docx_images.ImageError):
        docx_images.prepare_image(b"x" * (docx_images.MAX_UPLOAD_BYTES + 1))


def test_stamp_on_white_paper_gets_a_transparent_background():
    prepared = docx_images.prepare_image(_png((200, 200), paper=True), remove_background=True)
    img = Image.open(io.BytesIO(prepared.data))
    assert prepared.ext == "png" and img.mode == "RGBA"
    assert img.getpixel((2, 2))[3] == 0
    assert img.getpixel((100, 7))[3] > 200


def test_insert_move_float_resize_delete_round_trip():
    data, ids = _tagged("Title", "Body text", "Signature: ____")
    prepared = docx_images.prepare_image(_png())
    data, image_id = docx_images.insert_image(data, prepared, ids[0], "after", 3, "center")
    info = docx_images.list_images(data)
    assert [i["id"] for i in info] == [image_id] and info[0]["width_cm"] == 3.0 and not info[0]["floating"]
    assert len(docx_blocks.block_ids(data)) == 4

    data = docx_images.float_image(data, image_id, ids[2], 2 * docx_images.EMU_PER_CM, -300000)
    info = docx_images.list_images(data)[0]
    assert info["floating"] and info["block_id"] == ids[2] and info["y_emu"] == -300000
    assert len(docx_blocks.block_ids(data)) == 3, "the emptied image paragraph is removed"
    assert _texts(data)[2] == "Signature: ____"

    data = docx_images.resize_image(data, image_id, 6)
    assert docx_images.list_images(data)[0]["width_cm"] == 6.0
    assert docx_images.list_images(data)[0]["height_cm"] == 3.0

    data = docx_images.move_image(data, image_id, ids[1], "inline", "right")
    info = docx_images.list_images(data)[0]
    assert not info["floating"] and info["block_id"] == ids[1]
    assert _texts(data)[1] == "Body text"

    data = docx_images.delete_image(data, image_id)
    assert docx_images.list_images(data) == []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert not [n for n in z.namelist() if n.startswith("word/media/tq_")]
    Document(io.BytesIO(data))


def test_nudging_a_picture_within_its_own_paragraph_keeps_it():
    data, ids = _tagged("Title", "Body text")
    data, image_id = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[0], "after", 3)
    own = docx_images.list_images(data)[0]["block_id"]
    data = docx_images.float_image(data, image_id, own, 100000, 20000)
    data = docx_images.float_image(data, image_id, own, 200000, 40000)
    info = docx_images.list_images(data)
    assert [i["id"] for i in info] == [image_id] and info[0]["block_id"] == own and info[0]["x_emu"] == 200000
    data = docx_images.move_image(data, image_id, own, "inline")
    assert [i["id"] for i in docx_images.list_images(data)] == [image_id]


def test_moving_a_picture_keeps_a_symbol_or_field_sharing_its_paragraph():
    data, ids = _tagged("Title", "Body text")
    data, image_id = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[0], "inline", 3)
    doc = docx_blocks._Doc.load(data)
    title = doc.find_block(ids[0])
    for p in list(doc.trees["word/document.xml"].iter(f"{{{W}}}p")):
        if p is not title and next(p.iter(f"{{{W}}}drawing"), None) is not None:
            raise AssertionError("expected the picture inline in the title")
    # A checkbox glyph (w:sym) and a page field: no plain text, so paragraph_text sees nothing.
    title.append(etree.fromstring(f'<w:r xmlns:w="{W}"><w:sym w:font="Wingdings" w:char="F0A8"/></w:r>'))
    title.append(etree.fromstring(f'<w:fldSimple xmlns:w="{W}" w:instr="PAGE"/>'))
    for r in [r for r in title.findall(f"{{{W}}}r") if r.find(f"{{{W}}}t") is not None]:
        title.remove(r)
    data = doc.dump()

    moved = docx_images.move_image(data, image_id, ids[1], "after")
    body = _body_xml(moved).find(f"{{{W}}}body")
    assert body.find(f".//{{{W}}}sym") is not None and body.find(f".//{{{W}}}fldSimple") is not None
    assert ids[0] in docx_blocks.block_ids(moved)

    deleted = docx_images.delete_image(data, image_id)
    assert _body_xml(deleted).find(f".//{{{W}}}sym") is not None


def test_every_picture_has_an_editor_marker_right_before_it():
    data, ids = _tagged("One", "Two")
    data, image_id = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[1], "inline")
    body = _body_xml(data)
    run = next(body.iter(f"{{{W}}}drawing")).getparent()
    start = run.getprevious().getprevious()
    assert start.get(f"{{{W}}}name") == "_" + image_id
    again = docx_blocks.tag_blocks(data)
    assert len([b for b in _body_xml(again).iter(f"{{{W}}}bookmarkStart") if (b.get(f"{{{W}}}name") or "").startswith("_img_")]) == 1


def test_identical_images_share_one_media_part():
    data, ids = _tagged("One", "Two")
    prepared = docx_images.prepare_image(_png())
    data, _ = docx_images.insert_image(data, prepared, ids[0])
    data, _ = docx_images.insert_image(data, prepared, ids[1])
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert len([n for n in z.namelist() if n.startswith("word/media/")]) == 1


def test_set_text_keeps_an_image_in_the_paragraph():
    data, ids = _tagged("Signed by the Registrar")
    data, image_id = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[0], "inline")
    out, changed = docx_blocks.apply_operations(data, [{"op": "set_text", "target": ids[0], "content": "Signed by the Mayor"}])
    assert _texts(out)[0] == "Signed by the Mayor"
    assert [i["id"] for i in docx_images.list_images(out)] == [image_id]
    out, _ = docx_blocks.apply_text_edits(data, [(ids[0], "")])
    assert [i["id"] for i in docx_images.list_images(out)] == [image_id]


def _drawing_xml(data: bytes) -> str:
    body = _body_xml(data)
    return etree.tostring(next(body.iter(f"{{{W}}}drawing")), encoding="unicode")


def test_ai_fragments_may_reuse_existing_images_only():
    data, ids = _tagged("Title", "Body")
    data, image_id = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[0], "inline")
    drawing = _drawing_xml(data)
    moved, _ = docx_blocks.apply_operations(data, [
        {"op": "replace_paragraph", "target": ids[0], "content": "<w:p><w:r><w:t>Title</w:t></w:r></w:p>"},
        {"op": "insert_after", "target": ids[1], "content": f"<w:p><w:r>{drawing}</w:r></w:p>"},
    ])
    info = docx_images.list_images(moved)
    assert [i["id"] for i in info] == [image_id]
    assert info[0]["block_id"] == docx_blocks.block_ids(moved)[-1]

    for bad in (
        drawing.replace('r:embed="rId', 'r:embed="rIdX'),
        drawing.replace('r:embed=', 'r:link='),
        "<w:drawing/>",
    ):
        with pytest.raises(docx_blocks.DocxEditError):
            docx_blocks.apply_operations(data, [{"op": "insert_after", "target": ids[1], "content": f"<w:p><w:r>{bad}</w:r></w:p>"}])
    with pytest.raises(docx_blocks.DocxEditError):
        docx_blocks.apply_operations(data, [{"op": "insert_after", "target": ids[1], "content": '<w:p><w:hyperlink r:id="rId1"><w:r><w:t>x</w:t></w:r></w:hyperlink></w:p>'}])


def test_duplicated_drawings_get_unique_ids():
    data, ids = _tagged("Title", "Body")
    data, image_id = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[0], "inline")
    drawing = _drawing_xml(data)
    out, _ = docx_blocks.apply_operations(data, [{"op": "insert_after", "target": ids[1], "content": f"<w:p><w:r>{drawing}</w:r></w:p>"}])
    found = [i["id"] for i in docx_images.list_images(out)]
    assert len(found) == 2 and len(set(found)) == 2 and image_id in found
    docpr_ids = [d.get("id") for d in _body_xml(out).iter(f"{{{WP}}}docPr")]
    assert len(set(docpr_ids)) == 2


def test_strip_blocks_keeps_images_and_drops_editor_markers():
    data, ids = _tagged("Title", "Body")
    data, _ = docx_images.insert_image(data, docx_images.prepare_image(_png()), ids[0])
    clean = docx_blocks.strip_blocks(data)
    body = _body_xml(clean)
    assert next(body.iter(f"{{{W}}}drawing"), None) is not None
    assert [b.get(f"{{{W}}}name") for b in body.iter(f"{{{W}}}bookmarkStart")] == []


def _content(**kw):
    values = {
        "translator": "Anna Rossi",
        "date": "29 September 2026",
        "source_language": "Italian",
        "target_language": "English",
        "document": "Certificate of residence",
    }
    return docx_certification.CertContent(lang="en", values=values, day=date(2026, 9, 29), pages=2, **kw)


def test_certification_section_fields_update_in_place_and_keep_manual_edits():
    data, _ = _tagged("CERTIFICATE OF RESIDENCE", "Body")
    stamp = docx_images.prepare_image(_png((200, 200), paper=True), remove_background=True)
    data = docx_certification.add_certification(data, _content(stamp=stamp))
    assert docx_certification.has_certification(data)
    fields = docx_certification.read_fields(data)
    assert fields["translator"] == "Anna Rossi" and fields["date_iso"] == "2026-09-29"
    texts = _texts(data)
    assert "CERTIFIED TRANSLATION" in texts and "Source language: Italian" in texts
    signature = [i for i in docx_images.list_images(data) if i["floating"]]
    assert len(signature) == 1

    ids = docx_blocks.block_ids(data)
    statement = next(b for b, t in docx_blocks.paragraph_texts(data, ids).items() if t.startswith("I, Anna Rossi"))
    data, _ = docx_blocks.apply_text_edits(data, [(statement, "I, Anna Rossi, sworn translator, certify this.")])
    data, changed = docx_certification.update_fields(data, {"date": "2026-10-01", "translator": "Anna Bianchi"}, "en")
    assert set(changed) == {"date", "translator"}
    fields = docx_certification.read_fields(data)
    assert fields["date"] == "1 October 2026" and fields["date_iso"] == "2026-10-01"
    texts = _texts(data)
    assert "I, Anna Bianchi, sworn translator, certify this." in texts
    assert texts[-1] == "Anna Bianchi"

    again = docx_certification.add_certification(data, _content())
    assert _texts(again).count("CERTIFIED TRANSLATION") == 1
    removed = docx_certification.remove_certification(again)
    assert not docx_certification.has_certification(removed)
    assert _texts(removed) == ["CERTIFICATE OF RESIDENCE", "Body"]


def test_certification_opens_a_new_page_in_word_and_in_the_preview():
    data, _ = _tagged("Body")
    data = docx_certification.add_certification(data, _content())
    body = _body_xml(data)
    first = next(p for p in body.iter(f"{{{W}}}p") if any(b.get(f"{{{W}}}name") == "_cert_start" for b in p.iter(f"{{{W}}}bookmarkStart")))
    assert first.find(f"{{{W}}}pPr/{{{W}}}pageBreakBefore") is not None
    assert first.find(f"{{{W}}}pPr/{{{W}}}pStyle").get(f"{{{W}}}val") == docx_certification.PAGE_STYLE_ID
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert docx_certification.PAGE_STYLE_ID.encode() in z.read("word/styles.xml")


def test_typed_and_ai_edits_work_inside_content_controls():
    data, _ = _tagged("Body")
    data = docx_certification.add_certification(data, _content())
    texts = docx_blocks.paragraph_texts(data, docx_blocks.block_ids(data))
    date_block = next(b for b, t in texts.items() if t.startswith("Date: "))
    out, _ = docx_blocks.apply_operations(data, [{"op": "set_text", "target": date_block, "content": "Date: 2 October 2026"}])
    assert docx_certification.read_fields(out)["date"] == "2 October 2026"


def test_template_tokens_become_content_controls():
    tpl = Document()
    tpl.add_paragraph("AFFIDAVIT OF ACCURACY")
    tpl.add_paragraph("I, {{translator_name}}, translated {{document_name}} from {{source_language}} on {{date_long}}. Cert {{certificate_number}}.")
    buf = io.BytesIO()
    tpl.save(buf)
    data, _ = _tagged("Body")
    data = docx_certification.add_certification(
        data, _content(template_docx=buf.getvalue(), extra_tokens={"certificate_number": "CERT-2026-AB"})
    )
    texts = _texts(data)
    assert "AFFIDAVIT OF ACCURACY" in texts
    assert "I, Anna Rossi, translated Certificate of residence from Italian on 29 September 2026. Cert CERT-2026-AB." in texts
    data, _ = docx_certification.update_fields(data, {"source_language": "Spanish"}, "en")
    assert any("from Spanish on" in t for t in _texts(data))


def test_locale_helpers():
    assert cert_locale.language_code("it-IT") == "it"
    assert cert_locale.language_code("Italian") == "it"
    assert cert_locale.language_code("English (UK)") == "en"
    assert cert_locale.language_code("auto") == ""
    assert cert_locale.cert_language("ja-JP") == "en"
    assert cert_locale.language_name("en-GB", "it") == "Inglese"
    assert cert_locale.format_date(date(2026, 9, 29), "es") == "29 de septiembre de 2026"
    assert cert_locale.format_date(date(2026, 3, 1), "de") == "1. März 2026"


def test_document_title_guess_prefers_the_heading():
    data, _ = _tagged("CERTIFICATE OF RESIDENCE", "Mr. Bianchi 12/03/1987")
    assert docx_certification.guess_document_title(data, "scan 001") == "CERTIFICATE OF RESIDENCE"
    data, _ = _tagged("12/03/1987 - 00123", "x")
    assert docx_certification.guess_document_title(data, "scan 001") == "scan 001"


# --- API --------------------------------------------------------------------


@pytest.fixture()
def project_with_doc(db, storage, make_user, make_project):
    def _make(plan="PRO", owner=None):
        owner = owner or make_user(plan=plan)
        project = make_project(owner)
        key = f"uploads/{uuid.uuid4()}_authored.docx"
        storage["objects"][key] = _docx("CERTIFICATE OF RESIDENCE", "Mr. BIANCHI LUCA, born in Bari", "Signature: ________")
        project.authored_docx_s3_key = key
        db.commit()
        return owner, project

    return _make


def _get(client, owner, project):
    r = client.get(f"/projects/{project.id}/document", headers=owner["headers"])
    assert r.status_code == 200, r.text
    return r.content, int(r.headers["X-Document-Version"])


def _upload(client, owner, project, version, block_id, content=None, name="stamp.png", **form):
    return client.post(
        f"/projects/{project.id}/document/images",
        headers=owner["headers"],
        data={"version": str(version), "block_id": block_id, **{k: str(v) for k, v in form.items()}},
        files={"file": (name, content if content is not None else _png(), "image/png")},
    )


def test_image_api_flow_with_undo(client, project_with_doc):
    owner, project = project_with_doc()
    data, v = _get(client, owner, project)
    ids = docx_blocks.block_ids(data)
    url = f"/projects/{project.id}/document"

    r = _upload(client, owner, project, v, ids[2], position="inline", width_cm=3)
    assert r.status_code == 200, r.text
    image_id, v = r.json()["image_id"], r.json()["version"]
    assert v == 2

    listed = client.get(f"{url}/images", headers=owner["headers"]).json()
    assert listed["images"][0]["id"] == image_id

    r = client.post(f"{url}/images/{image_id}/position", headers=owner["headers"], json={"version": v, "target_block_id": ids[2], "x_emu": 1800000, "y_emu": -200000})
    assert r.status_code == 200, r.text
    v = r.json()["version"]
    r = client.post(f"{url}/images/{image_id}/resize", headers=owner["headers"], json={"version": v, "width_cm": 5})
    v = r.json()["version"]
    r = client.post(f"{url}/images/{image_id}/move", headers=owner["headers"], json={"version": v, "target_block_id": ids[0], "position": "after", "align": "center"})
    assert r.status_code == 200, r.text
    v = r.json()["version"]
    r = client.post(f"{url}/images/{image_id}/align", headers=owner["headers"], json={"version": v, "align": "right"})
    v = r.json()["version"]
    data, _ = _get(client, owner, project)
    info = docx_images.list_images(data)[0]
    assert info["width_cm"] == 5.0 and not info["floating"]

    stale = client.post(f"{url}/images/{image_id}/resize", headers=owner["headers"], json={"version": 1, "width_cm": 2})
    assert stale.status_code == 409

    r = client.delete(f"{url}/images/{image_id}?version={v}", headers=owner["headers"])
    assert r.status_code == 200
    data, v = _get(client, owner, project)
    assert docx_images.list_images(data) == []
    assert client.post(f"{url}/undo", headers=owner["headers"], json={}).status_code == 200
    data, _ = _get(client, owner, project)
    assert [i["id"] for i in docx_images.list_images(data)] == [image_id]


def test_image_upload_validation_and_tenancy(client, project_with_doc, make_user):
    owner, project = project_with_doc()
    data, v = _get(client, owner, project)
    bid = docx_blocks.block_ids(data)[0]
    assert _upload(client, owner, project, v, bid, content=b"not an image").status_code == 400
    assert _upload(client, owner, project, v, bid, name="evil.svg").status_code == 400
    assert _upload(client, owner, project, v, "_b00000000").status_code == 409
    assert _upload(client, owner, project, v + 5, bid).status_code == 409
    other = make_user()
    assert _upload(client, other, project, v, bid).status_code == 404
    assert client.get(f"/projects/{project.id}/document/page-stamp", headers=other["headers"]).status_code == 404


def test_certification_api(client, db, storage, project_with_doc):
    owner, project = project_with_doc()
    from tests.test_media_library import add_media

    add_media(db, storage, owner, "stamp", None)
    url = f"/projects/{project.id}/document/certification"
    _, v = _get(client, owner, project)
    assert client.get(url, headers=owner["headers"]).json()["present"] is False

    r = client.post(url, headers=owner["headers"], json={"version": v})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["fields"]["translator"] == "Test User"
    assert body["fields"]["source_language"] == "Italian" and body["fields"]["target_language"] == "English"
    assert body["fields"]["document"] == "CERTIFICATE OF RESIDENCE"
    assert body["fields"]["date"] == cert_locale.format_date(date.today(), "en")

    r = client.put(url, headers=owner["headers"], json={"version": body["version"], "fields": {"date": "2026-01-05", "document": "Residence certificate"}})
    assert r.status_code == 200, r.text
    assert r.json()["fields"]["date"] == "5 January 2026"
    status = client.get(f"/projects/{project.id}/document/status", headers=owner["headers"]).json()
    assert status["certification"] is True and status["images"] == 1

    r = client.delete(f"{url}?version={r.json()['version']}", headers=owner["headers"])
    assert r.status_code == 200
    assert client.get(url, headers=owner["headers"]).json()["present"] is False


def test_certification_is_a_pro_feature(client, project_with_doc):
    owner, project = project_with_doc(plan="BASIC")
    _, v = _get(client, owner, project)
    url = f"/projects/{project.id}/document/certification"
    assert client.post(url, headers=owner["headers"], json={"version": v}).status_code == 403
    assert client.get(url, headers=owner["headers"]).status_code == 403


def test_certification_template_must_belong_to_the_team(client, db, project_with_doc, make_user):
    from app.models.certification import Certification

    owner, project = project_with_doc()
    stranger = make_user()
    cert = Certification(team_id=stranger["team"].id, file_name="t.docx", file_path="uploads/t.docx", file_hash="0" * 64, size_bytes=1)
    db.add(cert)
    db.commit()
    _, v = _get(client, owner, project)
    r = client.post(f"/projects/{project.id}/document/certification", headers=owner["headers"], json={"version": v, "template_id": str(cert.id)})
    assert r.status_code == 404


def _export(client, owner, project) -> bytes:
    r = client.get(f"/projects/{project.id}/export", headers=owner["headers"])
    assert r.status_code == 200, r.text
    return r.content


def test_export_contains_exactly_one_certification_and_the_images(client, db, storage, project_with_doc):
    owner, project = project_with_doc()
    from tests.test_media_library import add_media

    add_media(db, storage, owner, "logo", None)
    data, v = _get(client, owner, project)
    r = _upload(client, owner, project, v, docx_blocks.block_ids(data)[2], position="inline")
    v = r.json()["version"]
    v = client.post(f"/projects/{project.id}/document/certification", headers=owner["headers"], json={"version": v}).json()["version"]

    out = _export(client, owner, project)
    Document(io.BytesIO(out))
    texts = _texts(out)
    assert texts.count("CERTIFIED TRANSLATION") == 1
    assert "Mr. BIANCHI LUCA, born in Bari" in texts
    assert texts.index("Mr. BIANCHI LUCA, born in Bari") < texts.index("CERTIFIED TRANSLATION")
    body = _body_xml(out)
    assert len(list(body.iter(f"{{{W}}}drawing"))) == 2
    docpr = [d.get("id") for d in body.iter(f"{{{WP}}}docPr")]
    assert len(docpr) == len(set(docpr))
    assert not [b for b in body.iter(f"{{{W}}}bookmarkStart") if (b.get(f"{{{W}}}name") or "").startswith(("_b", "_img_", "_cert_"))]
    tags = {t.get(f"{{{W}}}val") for t in body.iter(f"{{{W}}}tag")}
    assert "cert.date" in tags
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        assert docx_certification.PAGE_STYLE_ID.encode() in z.read("word/styles.xml")


def test_export_without_certification_page_keeps_the_appended_one(client, project_with_doc):
    owner, project = project_with_doc()
    _get(client, owner, project)
    texts = _texts(_export(client, owner, project))
    assert texts.count("CERTIFIED TRANSLATION") == 1


def test_certification_title_prefers_document_heading_over_issuer_lines():
    from docx import Document
    from docx.shared import Pt

    from app.services import docx_blocks, docx_certification

    doc = Document()
    for text in ("ITALIAN REPUBLIC", "MUNICIPALITY OF ESEMPIO", "CERTIFICATE OF RESIDENCE", "Mr Luca Bianchi resides in Esempio."):
        run = doc.add_paragraph().add_run(text)
        run.bold = text.isupper()
        run.font.size = Pt(14)
    buf = __import__("io").BytesIO()
    doc.save(buf)
    data = docx_blocks.tag_blocks(buf.getvalue())
    assert docx_certification.guess_document_title(data, "x") == "CERTIFICATE OF RESIDENCE"
    assert docx_certification.guess_document_title(data, "x", "residence certificate") == "CERTIFICATE OF RESIDENCE"


# --- picking the certification template ---------------------------------------


def _cert_template(db, storage, owner, title, statement, name="template.docx", default=False):
    from app.models.certification import Certification

    key = f"uploads/{uuid.uuid4()}_{name}"
    storage["objects"][key] = _docx(title, statement, "Firma / Signature ________")
    cert = Certification(
        team_id=owner["team"].id, file_name=name, file_path=key, file_hash="0" * 64, size_bytes=1, is_default=default,
    )
    db.add(cert)
    db.commit()
    return cert


ENGLISH = (
    "STATEMENT OF ACCURACY",
    "I, {{translator}}, certify that this is a true translation of the {{document}} from {{source_language}} "
    "to {{target_language}} and that it is complete. Date: {{date}}",
)
ITALIAN = (
    "DICHIARAZIONE DI TRADUZIONE",
    "Io sottoscritto {{translator}} dichiaro che la traduzione del documento {{document}} dal {{source_language}} "
    "al {{target_language}} è conforme al testo originale e che la presente è completa. Data: {{date}}",
)


def test_certification_page_can_be_added_from_a_chosen_template_or_the_standard_one(client, db, storage, project_with_doc):
    from app.services import cert_page

    owner, project = project_with_doc()
    english = _cert_template(db, storage, owner, *ENGLISH, name="english.docx")
    italian = _cert_template(db, storage, owner, *ITALIAN, name="italiano.docx", default=True)
    _cert_template(db, storage, owner, "Invoice", "x", name="invoice.pdf")
    url = f"/projects/{project.id}/document/certification"
    _, v = _get(client, owner, project)

    got = client.get(url, headers=owner["headers"]).json()
    assert got["present"] is False and got["template_id"] == str(italian.id)
    assert [t["id"] for t in got["templates"]] == [str(italian.id), str(english.id)]
    assert got["templates"][0]["is_default"] and got["templates"][1]["name"] == "english.docx"

    r = client.post(url, headers=owner["headers"], json={"version": v, "template_id": str(english.id)})
    assert r.status_code == 200, r.text
    assert r.json()["template_id"] == str(english.id)
    data, v = _get(client, owner, project)
    texts = _texts(data)
    assert "STATEMENT OF ACCURACY" in texts and "DICHIARAZIONE DI TRADUZIONE" not in texts
    db.refresh(project)
    assert project.certification_template_id == english.id and not project.certification_standard
    assert client.get(url, headers=owner["headers"]).json()["template_id"] == str(english.id)

    r = client.post(url, headers=owner["headers"], json={"version": v, "template_id": "standard"})
    assert r.status_code == 200, r.text
    assert r.json()["template_id"] == "standard"
    data, v = _get(client, owner, project)
    texts = _texts(data)
    assert texts.count("CERTIFIED TRANSLATION") == 1 and "STATEMENT OF ACCURACY" not in texts
    db.refresh(project)
    assert project.certification_template_id is None and project.certification_standard
    assert cert_page.template_for(db, project) is None
    assert client.get(url, headers=owner["headers"]).json()["template_id"] == "standard"

    # Standard sticks for the export even though the team has a default template.
    v = client.delete(f"{url}?version={v}", headers=owner["headers"]).json()["version"]
    texts = _texts(_export(client, owner, project))
    assert texts.count("CERTIFIED TRANSLATION") == 1 and "DICHIARAZIONE DI TRADUZIONE" not in texts

    assert client.post(url, headers=owner["headers"], json={"version": v, "template_id": "nope"}).status_code == 422


def test_switching_template_keeps_the_values_and_replaces_the_page(client, db, storage, project_with_doc):
    owner, project = project_with_doc()
    italian = _cert_template(db, storage, owner, *ITALIAN, name="italiano.docx")
    url = f"/projects/{project.id}/document/certification"
    _, v = _get(client, owner, project)
    v = client.post(url, headers=owner["headers"], json={"version": v}).json()["version"]
    r = client.put(url, headers=owner["headers"], json={"version": v, "fields": {"date": "2026-01-05", "translator": "Anna Verdi"}})
    v = r.json()["version"]

    # The panel sends its current draft with the switch, including a document title not saved yet.
    r = client.put(
        url,
        headers=owner["headers"],
        json={
            "version": v,
            "template_id": str(italian.id),
            "fields": {"date": "2026-01-05", "translator": "Anna Verdi", "document": "Residence certificate",
                       "source_language": "Italian", "target_language": "English"},
        },
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["template_id"] == str(italian.id)
    fields = body["fields"]
    assert fields["translator"] == "Anna Verdi" and fields["document"] == "Residence certificate"
    assert fields["date_iso"] == "2026-01-05" and fields["date"] == cert_locale.format_date(date(2026, 1, 5), "it")
    assert fields["source_language"] == "Italiano" and fields["target_language"] == "Inglese"

    data, v = _get(client, owner, project)
    texts = _texts(data)
    assert texts.count("DICHIARAZIONE DI TRADUZIONE") == 1 and "CERTIFIED TRANSLATION" not in texts
    assert texts.index("Mr. BIANCHI LUCA, born in Bari") < texts.index("DICHIARAZIONE DI TRADUZIONE")
    starts = [b for b in _body_xml(data).iter(f"{{{W}}}bookmarkStart") if b.get(f"{{{W}}}name") == "_cert_start"]
    assert len(starts) == 1
    db.refresh(project)
    assert project.certification_template_id == italian.id and not project.certification_standard

    # Same template again: fields update in place, nothing is rebuilt.
    r = client.put(url, headers=owner["headers"], json={"version": v, "template_id": str(italian.id), "fields": {"translator": "Anna B."}})
    assert r.status_code == 200 and r.json()["fields"]["translator"] == "Anna B."

    # A normal version: undo brings back the standard page and the panel reads it.
    assert client.post(f"/projects/{project.id}/document/undo", headers=owner["headers"], json={}).status_code == 200
    assert client.post(f"/projects/{project.id}/document/undo", headers=owner["headers"], json={}).status_code == 200
    data, v = _get(client, owner, project)
    assert "CERTIFIED TRANSLATION" in _texts(data) and "DICHIARAZIONE DI TRADUZIONE" not in _texts(data)
    got = client.get(url, headers=owner["headers"]).json()
    assert got["template_id"] == "standard" and got["fields"]["translator"] == "Anna Verdi"

    # The export and later re-adds use the remembered pick.
    client.delete(f"{url}?version={v}", headers=owner["headers"])
    texts = _texts(_export(client, owner, project))
    assert texts.count("DICHIARAZIONE DI TRADUZIONE") == 1 and "CERTIFIED TRANSLATION" not in texts


def test_switch_needs_a_page_and_a_word_template_of_the_team(client, db, storage, project_with_doc, make_user):
    owner, project = project_with_doc()
    stranger = make_user()
    theirs = _cert_template(db, storage, stranger, *ENGLISH, name="theirs.docx")
    pdf = _cert_template(db, storage, owner, "Certificate", "x", name="iso17100.pdf")
    mine = _cert_template(db, storage, owner, *ENGLISH, name="mine.docx")
    url = f"/projects/{project.id}/document/certification"
    _, v = _get(client, owner, project)

    assert client.post(url, headers=owner["headers"], json={"version": v, "template_id": str(theirs.id)}).status_code == 404
    assert client.post(url, headers=owner["headers"], json={"version": v, "template_id": str(pdf.id)}).status_code == 422
    assert client.put(url, headers=owner["headers"], json={"version": v, "template_id": str(mine.id)}).status_code == 409

    v = client.post(url, headers=owner["headers"], json={"version": v}).json()["version"]
    assert client.put(url, headers=owner["headers"], json={"version": v, "template_id": str(theirs.id)}).status_code == 404
    assert client.put(url, headers=owner["headers"], json={"version": v, "template_id": str(pdf.id)}).status_code == 422
    db.refresh(project)
    assert project.certification_template_id is None and not project.certification_standard


def test_a_saved_pdf_pick_falls_back_to_a_word_page(client, db, storage, project_with_doc):
    owner, project = project_with_doc()
    pdf = _cert_template(db, storage, owner, "Certificate", "x", name="iso17100.pdf")
    project.certification_template_id = pdf.id
    db.commit()
    url = f"/projects/{project.id}/document/certification"
    _, v = _get(client, owner, project)
    assert client.get(url, headers=owner["headers"]).json()["template_id"] == "standard"
    r = client.post(url, headers=owner["headers"], json={"version": v})
    assert r.status_code == 200, r.text


def test_rebuilt_page_stays_where_it_was():
    data, _ = _tagged("Heading", "Body text")
    data = docx_certification.add_certification(data, _content(template_ref="standard"))
    doc = docx_blocks._Doc.load(data)
    body = doc.trees["word/document.xml"].find(f"{{{W}}}body")
    tail = etree.SubElement(body, f"{{{W}}}p")
    etree.SubElement(etree.SubElement(tail, f"{{{W}}}r"), f"{{{W}}}t").text = "Annex"
    body.append(body.find(f"{{{W}}}sectPr"))
    data = doc.dump()
    assert docx_certification.page_template(data) == "standard"

    ref = str(uuid.uuid4())
    out = docx_certification.add_certification(data, _content(template_ref=ref, statement_override="Other wording"))
    texts = _texts(out)
    assert texts[-1] == "Annex" and texts.count("CERTIFIED TRANSLATION") == 1 and "Other wording" in texts
    assert docx_certification.page_template(out) == ref
    names = {b.get(f"{{{W}}}name") for b in _body_xml(docx_blocks.strip_blocks(out)).iter(f"{{{W}}}bookmarkStart")}
    assert not any(n.startswith("_cert_") for n in names)
