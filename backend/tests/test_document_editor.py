import io
import json
import uuid
from types import SimpleNamespace

import pytest
from docx import Document

from app.services import claude_params, docx_blocks


def _docx(*paras: str) -> bytes:
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture()
def project_with_doc(db, storage, make_user, make_project):
    def _make(plan="PRO"):
        owner = make_user(plan=plan)
        project = make_project(owner)
        key = f"uploads/{uuid.uuid4()}_authored.docx"
        storage["objects"][key] = _docx("CERTIFICATE OF RESIDENCE", "Mr. BIANCHI LUCA, born in Bari", "[Signature]")
        project.authored_docx_s3_key = key
        db.commit()
        return owner, project

    return _make


def _get(client, owner, project):
    r = client.get(f"/projects/{project.id}/document", headers=owner["headers"])
    assert r.status_code == 200, r.text
    return r.content, int(r.headers["X-Document-Version"])


def _texts(data):
    return [p.text for p in Document(io.BytesIO(data)).paragraphs]


def test_first_load_snapshots_version_one_with_block_ids(client, project_with_doc):
    owner, project = project_with_doc()
    data, version = _get(client, owner, project)
    assert version == 1
    assert len(docx_blocks.block_ids(data)) == 3
    again, version2 = _get(client, owner, project)
    assert version2 == 1 and docx_blocks.block_ids(again) == docx_blocks.block_ids(data)


def test_typed_edit_undo_and_stale_version(client, project_with_doc):
    owner, project = project_with_doc()
    data, v = _get(client, owner, project)
    bid = docx_blocks.block_ids(data)[1]
    url = f"/projects/{project.id}/document"

    r = client.post(f"{url}/edits", headers=owner["headers"], json={"version": v, "edits": [{"block_id": bid, "text": "Mr. BIANCHI LUCA, born in Bari (BA)"}]})
    assert r.status_code == 200 and r.json()["version"] == 2
    data, v = _get(client, owner, project)
    assert _texts(data)[1] == "Mr. BIANCHI LUCA, born in Bari (BA)"

    stale = client.post(f"{url}/edits", headers=owner["headers"], json={"version": 1, "edits": [{"block_id": bid, "text": "x"}]})
    assert stale.status_code == 409

    assert client.post(f"{url}/undo", headers=owner["headers"], json={}).json()["version"] == 1
    data, _ = _get(client, owner, project)
    assert _texts(data)[1] == "Mr. BIANCHI LUCA, born in Bari"
    assert client.post(f"{url}/undo", headers=owner["headers"], json={}).status_code == 404


def test_other_tenant_cannot_touch_document(client, project_with_doc, make_user):
    owner, project = project_with_doc()
    other = make_user()
    url = f"/projects/{project.id}/document"
    assert client.get(url, headers=other["headers"]).status_code == 404
    assert client.post(f"{url}/edits", headers=other["headers"], json={"version": 1, "edits": []}).status_code == 404
    assert client.post(f"{url}/undo", headers=other["headers"], json={}).status_code == 404


def test_trial_sees_watermark_but_stored_document_stays_clean(client, db, storage, project_with_doc):
    owner, project = project_with_doc(plan="TRIAL")
    data, _ = _get(client, owner, project)
    header = lambda d: " ".join(p.text for s in Document(io.BytesIO(d)).sections for p in s.header.paragraphs)
    assert "PREVIEW ONLY" in header(data)
    db.refresh(project)
    assert "PREVIEW ONLY" not in header(storage["objects"][project.authored_docx_s3_key])


class _FakeClaude:
    def __init__(self, answers):
        self.answers = list(answers)
        self.requests = []

    def __call__(self, client, **kwargs):
        self.requests.append(kwargs)
        body = self.answers.pop(0)
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=json.dumps(body))],
            stop_reason="end_turn",
            usage=None,
            model=kwargs["model"],
        )


@pytest.fixture()
def fake_claude(monkeypatch):
    def install(*answers):
        fake = _FakeClaude(answers)
        monkeypatch.setattr(claude_params, "create_message", fake)
        monkeypatch.setattr(claude_params, "api_key", lambda: "test-key")
        return fake

    return install


def test_chat_edit_applies_operations_and_reports_changed_blocks(client, project_with_doc, fake_claude):
    owner, project = project_with_doc()
    data, v = _get(client, owner, project)
    ids = docx_blocks.block_ids(data)
    fake = fake_claude({
        "reply": "Split the name onto its own line.",
        "operations": [{
            "op": "replace",
            "target": ids[1],
            "content": "<w:p><w:r><w:t>Mr. BIANCHI LUCA</w:t></w:r></w:p><w:p><w:r><w:t>born in Bari</w:t></w:r></w:p>",
        }],
    })
    r = client.post(
        f"/projects/{project.id}/document/chat",
        headers=owner["headers"],
        json={"version": v, "block_ids": [ids[1]], "selected_text": "BIANCHI LUCA", "message": "Put the name on its own line"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["version"] == 2 and len(body["changed_block_ids"]) == 2
    assert body["reply"].startswith("Split")
    new_data, _ = _get(client, owner, project)
    assert _texts(new_data)[1:3] == ["Mr. BIANCHI LUCA", "born in Bari"]

    request_text = fake.requests[0]["messages"][0]["content"][-1]["text"]
    assert "BIANCHI LUCA" in request_text and ids[1] in request_text
    assert fake.requests[0]["output_config"]["format"]["type"] == "json_schema"
    assert fake.requests[0]["messages"][0]["content"][0]["type"] == "document"


def test_chat_retries_once_when_operations_are_invalid(client, project_with_doc, fake_claude):
    owner, project = project_with_doc()
    data, v = _get(client, owner, project)
    target = docx_blocks.block_ids(data)[2]
    fake = fake_claude(
        {"reply": "x", "operations": [{"op": "replace", "target": target, "content": "<w:p><w:r><w:drawing/></w:r></w:p>"}]},
        {"reply": "Done.", "operations": [{"op": "replace", "target": target, "content": "<w:p><w:r><w:t>[Signature]</w:t></w:r></w:p>"}]},
    )
    r = client.post(f"/projects/{project.id}/document/chat", headers=owner["headers"], json={"version": v, "message": "tidy"})
    assert r.status_code == 200, r.text
    assert len(fake.requests) == 2
    assert "could not be applied" in fake.requests[1]["messages"][-1]["content"]


def test_chat_question_without_operations_keeps_version(client, project_with_doc, fake_claude):
    owner, project = project_with_doc()
    _, v = _get(client, owner, project)
    fake_claude({"reply": "The stamp reads 'Registry Office'.", "operations": []})
    r = client.post(f"/projects/{project.id}/document/chat", headers=owner["headers"], json={"version": v, "message": "What does the stamp say?"})
    assert r.status_code == 200 and r.json()["version"] == v and r.json()["changed_block_ids"] == []


def test_chat_gives_up_after_two_bad_answers(client, project_with_doc, fake_claude):
    owner, project = project_with_doc()
    _, v = _get(client, owner, project)
    bad = {"reply": "x", "operations": [{"op": "delete", "target": "_b00000000", "content": ""}]}
    fake_claude(bad, bad)
    r = client.post(f"/projects/{project.id}/document/chat", headers=owner["headers"], json={"version": v, "message": "x"})
    assert r.status_code == 502


def test_staff_see_provider_error_detail(client, db, project_with_doc, monkeypatch):
    import anthropic
    import httpx

    owner, project = project_with_doc()
    _, v = _get(client, owner, project)

    def boom(client_, **kwargs):
        req = httpx.Request("POST", "https://api.anthropic.com/v1/messages")
        raise anthropic.AuthenticationError("invalid x-api-key", response=httpx.Response(401, request=req), body=None)

    monkeypatch.setattr(claude_params, "create_message", boom)
    monkeypatch.setattr(claude_params, "api_key", lambda: "k")
    url = f"/projects/{project.id}/document/chat"
    r = client.post(url, headers=owner["headers"], json={"version": v, "message": "x"})
    assert r.status_code == 502 and "invalid x-api-key" not in r.json()["detail"]

    owner["user"].role = "SUPERUSER"
    db.commit()
    r = client.post(url, headers=owner["headers"], json={"version": v, "message": "x"})
    assert "AuthenticationError 401" in r.json()["detail"]
