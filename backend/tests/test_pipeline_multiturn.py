import io
import re
from types import SimpleNamespace

import anthropic
import httpx2
import pytest
from docx import Document

from app.services import claude_authored_rebuild as car
from app.services import claude_multiturn_rebuild as mt

SOURCE_TEXT = (
    "UNIVERSITA DEGLI STUDI DI ESEMPIO\n"
    "Certificato n. 0000123456 Matricola 7654321\n"
    "Mario Rossi, nato il 01/02/1990, voto 105/110, data 15/07/2020, crediti 120"
)


def _docx(paragraphs, header=None):
    doc = Document()
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


FULL_BODY = [
    "UNIVERSITY OF ESEMPIO",
    "Mario Rossi, born on 01/02/1990, passed the final examination on 15/07/2020 "
    "with a grade of 105/110, having earned 120 credits. " + "This paragraph pads the body text. " * 5,
]
GOOD_DOCX = _docx(FULL_BODY, header="Certificate No. 0000123456 Student ID 7654321")


def test_extract_number_tokens_normalises():
    assert mt.extract_number_tokens("No. 0012345, 02/06/1990, 29/30, page 1 of 2") == {
        "12345", "2", "6", "1990", "29", "30",
    }
    assert mt.extract_number_tokens("١٢٣") == {"123"}


def test_missing_source_numbers_and_threshold():
    missing, total = mt.missing_source_numbers(SOURCE_TEXT, "0000123456 7654321 01/02/1990 105/110 15/07/2020 120")
    assert (missing, total) == ([], 11)
    missing, total = mt.missing_source_numbers(SOURCE_TEXT, "Certificate 123456, born 01/02/1990")
    assert missing == ["7", "15", "105", "110", "120", "2020", "7654321"]
    assert mt.coverage_failed(missing, total)
    assert not mt.coverage_failed(["7"], total)
    assert not mt.coverage_failed(["7", "15"], 40)


def test_inspection_reads_header_numbers_and_flags_dropped_content():
    report = mt._inspect_docx(GOOD_DOCX, source_text=SOURCE_TEXT)
    assert report["source_numbers_missing"] == 0
    assert "warnings" not in report

    dropped = _docx([FULL_BODY[0], "Mario Rossi passed the final examination. " * 6])
    report = mt._inspect_docx(dropped, source_text=SOURCE_TEXT)
    assert report["source_numbers_missing"] == 11
    assert any(w.startswith("CRITICAL") and "7654321" in w for w in report["warnings"])


def test_inspection_does_not_flag_registrar_attestation():
    doc = _docx([
        "MUNICIPALITY OF ESEMPIO",
        "I hereby certify that this is a true extract from the civil status register. " * 4,
    ])
    report = mt._inspect_docx(doc)
    assert report["forbidden_string_hits"] == []


def _response(blocks, stop_reason="tool_use"):
    usage = SimpleNamespace(input_tokens=10, output_tokens=5, cache_read_input_tokens=0, cache_creation_input_tokens=0)
    return SimpleNamespace(content=blocks, stop_reason=stop_reason, usage=usage, model="claude-opus-4-8")


def _tool(code, tool_id):
    return SimpleNamespace(type="tool_use", id=tool_id, input={"code": code})


class _FakeClient:
    def __init__(self, responses, requests):
        self._responses = list(responses)
        self._requests = requests
        client = self

        class _Stream:
            def __init__(self, resp):
                self._resp = resp

            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

            def get_final_message(self):
                return self._resp

        class _Messages:
            def create(self, **kwargs):
                client._requests.append(("create", kwargs))
                return client._next(kwargs)

            def stream(self, **kwargs):
                client._requests.append(("stream", kwargs))
                return _Stream(client._next(kwargs))

        self.messages = _Messages()

    def _next(self, kwargs):
        resp = self._responses.pop(0)
        return resp(kwargs) if callable(resp) else resp


@pytest.fixture
def loop_env(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    monkeypatch.setattr(car, "_extract_tables_via_vision", lambda *a, **k: [])
    monkeypatch.setattr(mt, "_pdf_text_layer", lambda pdf: SOURCE_TEXT)
    requests = []

    def install(responses):
        monkeypatch.setattr(anthropic, "Anthropic", lambda **kw: _FakeClient(responses, requests))
        return requests

    return install


def _run(**kwargs):
    return mt._author_rebuild_docx_multiturn_core(
        b"%PDF-fake", "it", "en", _force_doc_type="CERTIFICATE", **kwargs
    )


def test_loop_keeps_best_output_not_last(loop_env, monkeypatch):
    unclean_but_complete = _docx(["Mario Rossi heading in lower case"] + FULL_BODY[1:],
                                 header="Certificate No. 0000123456 Student ID 7654321")
    missing_numbers = _docx([FULL_BODY[0], "Mario Rossi passed the final examination. " * 6])
    runs = iter([(True, "", unclean_but_complete), (True, "", missing_numbers)])
    monkeypatch.setattr(mt, "_run_in_sandbox", lambda *a, **k: next(runs))
    requests = loop_env([
        _response([_tool("a", "t1")]),
        _response([_tool("b", "t2")]),
        _response([SimpleNamespace(type="text", text="done")], stop_reason="end_turn"),
    ])

    out = _run()

    assert len(requests) == 3
    texts = [p.text for p in Document(io.BytesIO(out)).paragraphs]
    assert texts[0] == "Mario Rossi heading in lower case"


def test_loop_stops_after_first_clean_inspection(loop_env, monkeypatch):
    monkeypatch.setattr(mt, "_run_in_sandbox", lambda *a, **k: (True, "", GOOD_DOCX))
    requests = loop_env([_response([_tool("a", "t1")]), _response([_tool("b", "t2")])])

    _run()

    assert len(requests) == 1
    kind, kwargs = requests[0]
    assert kind == "stream"
    assert kwargs["thinking"] == {"type": "adaptive"}
    assert "temperature" not in kwargs and "extra_body" not in kwargs
    assert kwargs["cache_control"] == {"type": "ephemeral", "ttl": "1h"}
    first_user = kwargs["messages"][0]["content"]
    assert first_user[0]["type"] == "document"
    assert first_user[1]["cache_control"] == {"type": "ephemeral", "ttl": "1h"}


def test_validation_error_is_fed_back_as_error_tool_result(loop_env):
    def bad_script(kwargs):
        prompt = kwargs["messages"][0]["content"][1]["text"]
        path = re.search(r'Save the final DOCX to exactly: r"(.*?)"', prompt).group(1)
        code = f'import subprocess\nfrom docx import Document\ndoc = Document()\ndoc.save(r"{path}")\n'
        return _response([_tool(code, "t1")])

    requests = loop_env([bad_script, _response([SimpleNamespace(type="text", text="giving up")], "end_turn")])

    with pytest.raises(RuntimeError, match="no successful DOCX"):
        _run()

    result = requests[1][1]["messages"][2]["content"][0]
    assert result["type"] == "tool_result"
    assert result["is_error"] is True
    assert "REJECTED" in result["content"] and "subprocess" in result["content"]
    assert "Import only: docx" in result["content"]


def test_page_by_page_failure_raises_instead_of_skipping(monkeypatch):
    monkeypatch.setattr(mt, "_split_pdf_per_page", lambda pdf: [b"p1", b"p2", b"p3"])
    monkeypatch.setattr(mt, "_pdf_page_count", lambda pdf: 3)

    def core(page_pdf, *a, **k):
        if page_pdf == b"p2":
            raise RuntimeError("no successful DOCX")
        return GOOD_DOCX

    monkeypatch.setattr(mt, "_author_rebuild_docx_multiturn_core", core)
    with pytest.raises(RuntimeError, match="page 2 of 3"):
        mt._author_rebuild_form_page_by_page(b"pdf", "it", "en")


def test_page_split_shortfall_raises(monkeypatch):
    monkeypatch.setattr(mt, "_split_pdf_per_page", lambda pdf: [b"p1", b"p2"])
    monkeypatch.setattr(mt, "_pdf_page_count", lambda pdf: 3)
    with pytest.raises(RuntimeError, match="2 of 3 pages"):
        mt._author_rebuild_form_page_by_page(b"pdf", "it", "en")


def test_multipage_form_routes_before_running_prepasses(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")

    def forbidden(*a, **k):
        raise AssertionError("whole-document pre-pass should not run")

    monkeypatch.setattr(car, "_extract_tables_via_vision", forbidden)
    monkeypatch.setattr(mt, "_classify_document", lambda pdf: "FORM")
    monkeypatch.setattr(mt, "_pdf_page_count", lambda pdf: 2)
    monkeypatch.setattr(mt, "_author_rebuild_form_page_by_page", lambda *a, **k: b"merged")

    assert mt._author_rebuild_docx_multiturn_core(b"pdf", "it", "en") == b"merged"


def test_default_max_turns_is_six():
    import inspect
    assert inspect.signature(mt.author_rebuild_docx_multiturn).parameters["max_turns"].default == 6


def _status_error(cls, status):
    request = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
    return cls("boom", response=httpx2.Response(status, request=request), body=None)


def test_single_shot_falls_back_down_model_chain(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    requests = []

    def overloaded(kwargs):
        raise _status_error(anthropic.OverloadedError, 529)

    ok = _response([SimpleNamespace(type="text", text="```python\nprint(1)\n```")], "end_turn")
    monkeypatch.setattr(anthropic, "Anthropic", lambda **kw: _FakeClient([overloaded, ok], requests))

    raw = car._call_claude_to_author(b"%PDF", "it", "en", "/tmp/out.docx", [], [], model="claude-opus-4-8")

    assert "print(1)" in raw
    assert [kw["model"] for _, kw in requests] == ["claude-opus-4-8", "claude-sonnet-4-6"]
    for kind, kw in requests:
        assert kind == "stream"
        assert kw["thinking"] == {"type": "adaptive"}
        assert "temperature" not in kw and "extra_body" not in kw


def test_single_shot_does_not_fall_back_on_auth_error(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    requests = []

    def denied(kwargs):
        raise _status_error(anthropic.AuthenticationError, 401)

    monkeypatch.setattr(anthropic, "Anthropic", lambda **kw: _FakeClient([denied], requests))
    with pytest.raises(anthropic.AuthenticationError):
        car._call_claude_to_author(b"%PDF", "it", "en", "/tmp/out.docx", [], [])
    assert len(requests) == 1


def test_author_rebuild_uses_rebuild_default_model(monkeypatch):
    from app.services import claude_params

    seen = {}
    monkeypatch.setattr(claude_params, "REBUILD_MODEL", "claude-opus-4-6")
    monkeypatch.setattr(
        mt, "author_rebuild_docx_multiturn",
        lambda pdf, s, t, model=None: seen.setdefault("model", model) and b"docx",
    )
    car.author_rebuild_docx(b"pdf", "it", "en")
    assert seen["model"] == "claude-opus-4-6"


def test_table_prepass_sends_no_temperature_to_opus_48(monkeypatch):
    import fitz

    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    pdf = fitz.open()
    pdf.new_page()
    pdf_bytes = pdf.tobytes()
    requests = []
    reply = _response([SimpleNamespace(type="text", text='{"tables": []}')], "end_turn")
    monkeypatch.setattr(anthropic, "Anthropic", lambda **kw: _FakeClient([reply], requests))

    assert car._extract_tables_via_vision(pdf_bytes, model="claude-opus-4-8") == []
    kind, kwargs = requests[0]
    assert kind == "create"
    assert "temperature" not in kwargs and "extra_body" not in kwargs and "thinking" not in kwargs
