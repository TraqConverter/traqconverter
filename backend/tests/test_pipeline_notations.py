import io
import re

from docx import Document

from app.services import claude_authored_rebuild as car
from app.services import claude_multiturn_rebuild as mt
from app.services import script_sandbox

NOTES = [
    "[Signature]",
    "[Round stamp: Municipality of Rome – Registry Office]",
    "[Stamp, illegible]",
    "Born on [illegible]/05/1990",
    "[Revenue stamp: €16.00]",
    "[Coat of arms of the Italian Republic]",
]


def _author_prompt():
    return car._AUTHOR_PROMPT_TEMPLATE.format(
        source_lang="Italian", target_lang="English", output_path="/tmp/out/rebuild.docx",
        image_list=car._format_image_list([]), table_list="",
        sandbox_rules=car.SANDBOX_RULES, notation_rules=car.NOTATION_RULES,
    )


def test_both_prompts_carry_notation_rules():
    assert "[illegible]" in car.NOTATION_RULES and "[Signature]" in car.NOTATION_RULES
    assert car.NOTATION_RULES in _author_prompt()
    assert "{notation_rules}" in mt._UNIVERSAL_PROMPT


def test_prompts_never_ask_for_pasted_images():
    for prompt in (_author_prompt(), mt._UNIVERSAL_PROMPT):
        assert not re.search(r"add_picture\(\s*[\"']images/", prompt)
    recipe = re.search(r"```python\n(.*?)```", _author_prompt(), re.S).group(1)
    assert "add_picture" not in recipe
    assert "[Signature]" in recipe
    script_sandbox.validate_script(recipe)


def test_notations_survive_post_processing():
    doc = Document()
    doc.add_paragraph("CERTIFICATE OF RESIDENCE")
    for note in NOTES:
        doc.add_paragraph(note)
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    for step in (
        car._strip_inline_cert_blocks,
        car._strip_html_and_bracket_artifacts,
        car._merge_adjacent_compatible_tables,
        car._collapse_pre_section_whitespace,
        car._strip_broken_image_drawings,
    ):
        data = step(data)
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    for note in NOTES:
        assert note in text, note
