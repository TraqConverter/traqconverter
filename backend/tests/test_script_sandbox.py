import tempfile
import textwrap
from pathlib import Path

import pytest

from app.services import script_sandbox

LEGIT = textwrap.dedent('''
    import os
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement, parse_xml

    OUTPUT_PATH = r"{out}"

    def _no_borders(tbl):
        tblPr = tbl._tbl.find(qn('w:tblPr'))
        if tblPr is None:
            tblPr = OxmlElement('w:tblPr')
            tbl._tbl.insert(0, tblPr)
        borders = OxmlElement('w:tblBorders')
        for edge in ("top", "left", "bottom", "right"):
            e = OxmlElement(f'w:{{edge}}')
            e.set(qn('w:val'), 'nil')
            borders.append(e)
        tblPr.append(borders)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    for p in list(section.header.paragraphs):
        p._element.getparent().remove(p._element)
    t = doc.add_table(rows=2, cols=2)
    _no_borders(t)
    t.cell(0, 0).merge(t.cell(0, 1))
    r = doc.add_paragraph().add_run("Hello")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0, 0, 0)
    r.add_break(WD_BREAK.PAGE)
    if os.path.exists("images/logo.png"):
        doc.add_paragraph().add_run().add_picture("images/logo.png", width=Cm(3))
    doc.save(OUTPUT_PATH)
''')


def _run(script, tmp):
    work = Path(tempfile.mkdtemp(dir=tmp))
    out = Path(tempfile.mkdtemp(dir=tmp))
    (work / "images").mkdir()
    from PIL import Image
    Image.new("RGB", (20, 20), "red").save(work / "images" / "logo.png")
    out_path = out / "rebuild.docx"
    proc = script_sandbox.run_script(script.format(out=out_path), work, [out], timeout_seconds=60)
    return proc, out_path


def test_legit_script_produces_docx(tmp_path):
    proc, out_path = _run(LEGIT, tmp_path)
    assert proc.returncode == 0, proc.stderr.decode()
    assert out_path.stat().st_size > 1000


@pytest.mark.parametrize("snippet", [
    "import socket",
    "import subprocess",
    "import ctypes",
    "import sys",
    "import importlib",
    "from lxml import etree",
    "import os\nos.system('id')",
    "x = getattr(object, 'mro')",
    "x = ().__class__.__bases__",
    "from docx.oxml import OxmlElement\nOxmlElement('w:p').getroottree().write('/tmp/pwn')",
    "import docx\ndocx.oxml.parser.etree",
    "eval('1')",
    "import os\nprint(os.environ)",
    "g = (i for i in [1])\ng.gi_frame",
])
def test_rejected_at_validation(snippet):
    with pytest.raises(ValueError):
        script_sandbox.validate_script(snippet)


@pytest.mark.parametrize("snippet", [
    "print(open('/etc/hosts').read())",
    "f = open('{out_dir}/../escape.txt', 'w')\nf.close()",
    "import os\nprint(os.listdir('/'))",
    "from docx import Document\nDocument('/etc/hosts')",
    "import pathlib\nprint(pathlib.Path('/etc/hosts').read_text())",
    "import os\nos.remove('/etc/hosts')",
])
def test_blocked_at_runtime(snippet, tmp_path):
    proc, _ = _run(snippet.replace("{out_dir}", "{{out_dir}}").replace("{{out_dir}}", str(tmp_path)), tmp_path)
    assert proc.returncode != 0
    assert b"sandbox" in proc.stderr


def test_external_entities_not_resolved(tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text("TOPSECRET")
    script = (
        "from docx.oxml import parse_xml\n"
        f"el = parse_xml('<!DOCTYPE r [<!ENTITY e SYSTEM \"file://{secret}\">]><r>&e;</r>')\n"
        "assert 'TOPSECRET' not in (el.text or '')\n"
    )
    proc, _ = _run(script, tmp_path)
    assert b"TOPSECRET" not in proc.stdout + proc.stderr
