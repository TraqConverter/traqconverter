"""Addressable editing of a DOCX: every paragraph carries a hidden bookmark id the editor can target."""
from __future__ import annotations

import io
import re
import secrets
import zipfile
from dataclasses import dataclass

from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
WP14_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
NSMAP = {"w": W_NS, "w14": W14_NS}
DRAWING_NAMESPACES = (WP_NS, WP14_NS, A_NS, PIC_NS)


def w(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


BLOCK_PREFIX = "_b"
BLOCK_ID_RE = re.compile(r"^_b[0-9a-f]{8}$")
IMAGE_ID_RE = re.compile(r"^img_[0-9a-f]{8}$")
IMAGE_MARKER_PREFIX = "_"
CERT_MARKERS = ("_cert_start", "_cert_end")
UNIT_TAGS = {w("p"), w("tbl"), w("sdt")}
CONTAINER_TAGS = {w("body"), w("hdr"), w("ftr")}
PART_RE = re.compile(r"^word/(document|header\d*|footer\d*)\.xml$")
FORBIDDEN_TAGS = {w("altChunk"), w("subDoc"), w("object"), w("pict"), w("control"), w("sectPr")}
_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False, remove_blank_text=False)


class DocxEditError(ValueError):
    pass


@dataclass
class _Doc:
    files: dict[str, bytes]
    trees: dict[str, etree._Element]

    @classmethod
    def load(cls, data: bytes) -> "_Doc":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            files = {n: z.read(n) for n in z.namelist()}
        trees = {n: etree.fromstring(b, _PARSER) for n, b in files.items() if PART_RE.match(n)}
        if "word/document.xml" not in trees:
            raise DocxEditError("Not a Word document")
        return cls(files, trees)

    def dump(self) -> bytes:
        for name, tree in self.trees.items():
            self.files[name] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            # [Content_Types].xml first, as Word expects.
            for name in sorted(self.files, key=lambda n: n != "[Content_Types].xml"):
                z.writestr(name, self.files[name])
        return buf.getvalue()

    def paragraphs(self):
        for tree in self.trees.values():
            yield from tree.iter(w("p"))

    def find_block(self, block_id: str):
        for tree in self.trees.values():
            for bm in tree.iter(w("bookmarkStart")):
                if bm.get(w("name")) == block_id:
                    return bm.getparent()
        return None

    def part_of(self, el) -> str:
        root = el.getroottree().getroot()
        for name, tree in self.trees.items():
            if tree is root:
                return name
        return "word/document.xml"

    def rels(self, part: str) -> etree._Element:
        folder, _, base = part.rpartition("/")
        name = f"{folder}/_rels/{base}.rels"
        if name not in self.trees:
            raw = self.files.get(name)
            self.trees[name] = (
                etree.fromstring(raw, _PARSER) if raw else etree.Element(f"{{{PKG_REL_NS}}}Relationships", nsmap={None: PKG_REL_NS})
            )
        return self.trees[name]

    def image_rel_ids(self, part: str) -> set[str]:
        return {
            r.get("Id")
            for r in self.rels(part)
            if r.get("Type") == IMAGE_REL and (r.get("TargetMode") or "Internal") != "External"
        }

    def next_bookmark_id(self) -> int:
        top = 0
        for tree in self.trees.values():
            for bm in tree.iter(w("bookmarkStart")):
                try:
                    top = max(top, int(bm.get(w("id"), "0")))
                except ValueError:
                    pass
        return top + 1


def _block_id_of(p) -> str | None:
    for bm in p.findall(w("bookmarkStart")):
        name = bm.get(w("name")) or ""
        if BLOCK_ID_RE.match(name):
            return name
    return None


def _strip_rsids(tree) -> None:
    for el in tree.iter():
        for attr in [a for a in el.attrib if a.startswith(f"{{{W_NS}}}rsid")]:
            del el.attrib[attr]


def _tag(doc: _Doc) -> None:
    """Give every paragraph exactly one unique block bookmark; drop duplicates copied by edits."""
    max_id = 0
    for tree in doc.trees.values():
        for bm in tree.iter(w("bookmarkStart")):
            try:
                max_id = max(max_id, int(bm.get(w("id"), "0")))
            except ValueError:
                pass
    seen: set[str] = set()
    for p in doc.paragraphs():
        own = None
        for bm in p.findall(w("bookmarkStart")):
            name = bm.get(w("name")) or ""
            if not BLOCK_ID_RE.match(name):
                continue
            if own is None and name not in seen:
                own = name
                continue
            bid = bm.get(w("id"))
            p.remove(bm)
            for end in p.findall(w("bookmarkEnd")):
                if end.get(w("id")) == bid:
                    p.remove(end)
        if own is None:
            max_id += 1
            own = BLOCK_PREFIX + secrets.token_hex(4)
            while own in seen:
                own = BLOCK_PREFIX + secrets.token_hex(4)
            start = etree.Element(w("bookmarkStart"))
            start.set(w("id"), str(max_id))
            start.set(w("name"), own)
            end = etree.Element(w("bookmarkEnd"))
            end.set(w("id"), str(max_id))
            ppr = p.find(w("pPr"))
            index = 1 if ppr is not None else 0
            p.insert(index, start)
            p.insert(index + 1, end)
        seen.add(own)
    _sync_images(doc)


def _marker_pair(bid: int, name: str):
    start = etree.Element(w("bookmarkStart"))
    start.set(w("id"), str(bid))
    start.set(w("name"), name)
    end = etree.Element(w("bookmarkEnd"))
    end.set(w("id"), str(bid))
    return start, end


def _remove_bookmark(tree, start) -> None:
    bid = start.get(w("id"))
    start.getparent().remove(start)
    for end in list(tree.iter(w("bookmarkEnd"))):
        if end.get(w("id")) == bid:
            end.getparent().remove(end)
            break


def drawing_run(drawing):
    el = drawing.getparent()
    while el is not None and el.tag != w("r"):
        el = el.getparent()
    return el


def _sync_images(doc: _Doc) -> None:
    """Every body picture gets a stable id (wp:docPr name) and a hidden bookmark right before its run so the editor can find it."""
    tree = doc.trees["word/document.xml"]
    docpr_tag = f"{{{WP_NS}}}docPr"
    used_ids: set[str] = set()
    used_nums: set[int] = set()
    wanted: dict[str, object] = {}
    for drawing in tree.iter(w("drawing")):
        docpr = next(drawing.iter(docpr_tag), None)
        run = drawing_run(drawing)
        if docpr is None or run is None:
            continue
        name = docpr.get("name") or ""
        if not IMAGE_ID_RE.match(name) or name in used_ids:
            name = "img_" + secrets.token_hex(4)
            while name in used_ids:
                name = "img_" + secrets.token_hex(4)
            docpr.set("name", name)
        used_ids.add(name)
        try:
            num = int(docpr.get("id", "0"))
        except ValueError:
            num = 0
        if num <= 0 or num in used_nums:
            num = max(used_nums | {0}) + 1
            docpr.set("id", str(num))
        used_nums.add(num)
        wanted[IMAGE_MARKER_PREFIX + name] = run
    for bm in list(tree.iter(w("bookmarkStart"))):
        name = bm.get(w("name")) or ""
        if not name.startswith(IMAGE_MARKER_PREFIX + "img_"):
            continue
        run = wanted.get(name)
        prev = run.getprevious() if run is not None else None
        prev2 = prev.getprevious() if prev is not None else None
        if run is not None and prev2 is bm and prev.tag == w("bookmarkEnd") and prev.get(w("id")) == bm.get(w("id")):
            wanted.pop(name)
            continue
        _remove_bookmark(tree, bm)
    next_id = doc.next_bookmark_id()
    for name, run in wanted.items():
        start, end = _marker_pair(next_id, name)
        next_id += 1
        run.addprevious(start)
        run.addprevious(end)


def tag_blocks(data: bytes) -> bytes:
    doc = _Doc.load(data)
    for tree in doc.trees.values():
        _strip_rsids(tree)
    _tag(doc)
    return doc.dump()


def strip_blocks(data: bytes) -> bytes:
    """Remove editor bookmarks for export."""
    doc = _Doc.load(data)
    for tree in doc.trees.values():
        ids = set()
        for bm in list(tree.iter(w("bookmarkStart"))):
            if _is_editor_marker(bm.get(w("name")) or ""):
                ids.add(bm.get(w("id")))
                bm.getparent().remove(bm)
        for end in list(tree.iter(w("bookmarkEnd"))):
            if end.get(w("id")) in ids:
                end.getparent().remove(end)
    return doc.dump()


def _is_editor_marker(name: str) -> bool:
    return bool(BLOCK_ID_RE.match(name)) or name in CERT_MARKERS or (
        name.startswith(IMAGE_MARKER_PREFIX) and bool(IMAGE_ID_RE.match(name[len(IMAGE_MARKER_PREFIX):]))
    )


def block_ids(data: bytes) -> list[str]:
    doc = _Doc.load(data)
    return [b for b in (_block_id_of(p) for p in doc.paragraphs()) if b]


def _pieces(p):
    """Text-bearing elements of a paragraph in order, excluding nested text-box paragraphs."""
    out = []
    # Field codes and their results (page numbers, dates) aren't visible to the editor; leave them untouched.
    in_field = [0]

    def walk(el):
        for child in el:
            tag = child.tag
            if tag in (w("txbxContent"), w("p"), w("fldSimple"), w("footnoteReference"), w("endnoteReference")):
                continue
            if tag == w("fldChar"):
                kind = child.get(w("fldCharType"))
                if kind == "begin":
                    in_field[0] += 1
                elif kind == "end" and in_field[0]:
                    in_field[0] -= 1
                continue
            if tag == w("t"):
                if not in_field[0]:
                    out.append((child, child.text or ""))
            elif tag == w("tab") and el.tag == w("r"):
                if not in_field[0]:
                    out.append((child, "\t"))
            elif tag in (w("br"), w("cr")) and el.tag == w("r"):
                if not in_field[0]:
                    out.append((child, "\n"))
            else:
                walk(child)

    walk(p)
    return out


def paragraph_text(p) -> str:
    return "".join(text for _, text in _pieces(p))


def _make_t(text: str):
    t = etree.Element(w("t"))
    t.text = text
    t.set(f"{{{XML_NS}}}space", "preserve")
    return t


def _elements_for(text: str) -> list:
    els = []
    for part in re.split(r"(\t|\n)", text):
        if part == "\t":
            els.append(etree.Element(w("tab")))
        elif part == "\n":
            els.append(etree.Element(w("br")))
        elif part:
            els.append(_make_t(part))
    return els


def set_paragraph_text(p, new_text: str) -> bool:
    """Replace a paragraph's text, keeping run formatting outside the changed region. Returns True if changed."""
    pieces = _pieces(p)
    old = "".join(text for _, text in pieces)
    if old == new_text:
        return False
    pre = 0
    limit = min(len(old), len(new_text))
    while pre < limit and old[pre] == new_text[pre]:
        pre += 1
    suf = 0
    while suf < limit - pre and old[len(old) - 1 - suf] == new_text[len(new_text) - 1 - suf]:
        suf += 1
    del_start, del_end = pre, len(old) - suf
    insert = new_text[pre:len(new_text) - suf]

    if not pieces:
        run = etree.SubElement(p, w("r"))
        mark_rpr = p.find(f"{w('pPr')}/{w('rPr')}")
        if mark_rpr is not None:
            run.append(_copy_rpr(mark_rpr))
        for el in _elements_for(new_text):
            run.append(el)
        return True

    # Anchor the insertion in the text element where the change starts (or the one it follows).
    anchor_el, anchor_offset = None, 0
    pos = 0
    for el, text in pieces:
        end = pos + len(text)
        if el.tag == w("t") and pos <= del_start <= end:
            anchor_el, anchor_offset = el, del_start - pos
            break
        pos = end
    if anchor_el is None:
        # The change starts at a tab/break boundary: put a fresh text node right before that piece.
        anchor_el, anchor_offset = _make_t(""), 0
        pos, following = 0, None
        for el, text in pieces:
            if pos >= del_start:
                following = el
                break
            pos += len(text)
        if following is not None:
            following.addprevious(anchor_el)
        else:
            pieces[-1][0].addnext(anchor_el)

    # Delete the old region piece by piece.
    pos = 0
    for el, text in pieces:
        start, end = pos, pos + len(text)
        pos = end
        lo, hi = max(start, del_start), min(end, del_end)
        if lo >= hi:
            continue
        if el.tag == w("t"):
            s = el.text or ""
            el.text = s[: lo - start] + s[hi - start:]
            el.set(f"{{{XML_NS}}}space", "preserve")
            if el is anchor_el and lo - start < anchor_offset:
                anchor_offset = lo - start
        else:
            el.getparent().remove(el)

    if insert:
        s = anchor_el.text or ""
        before, after = s[:anchor_offset], s[anchor_offset:]
        anchor_el.text = before
        anchor_el.set(f"{{{XML_NS}}}space", "preserve")
        new_els = _elements_for(insert)
        if after:
            new_els.append(_make_t(after))
        ref = anchor_el
        for el in new_els:
            ref.addnext(el)
            ref = el
    return True


def _copy_rpr(mark_rpr):
    rpr = etree.Element(w("rPr"))
    for child in mark_rpr:
        if child.tag not in (w("ins"), w("del"), w("moveFrom"), w("moveTo")):
            rpr.append(etree.fromstring(etree.tostring(child)))
    return rpr


def apply_text_edits(data: bytes, edits: list[tuple[str, str]]) -> tuple[bytes, list[str]]:
    doc = _Doc.load(data)
    changed = []
    for block_id, text in edits:
        bm_parent = doc.find_block(block_id)
        if bm_parent is None or bm_parent.tag != w("p"):
            raise DocxEditError(f"Unknown block {block_id}")
        if set_paragraph_text(bm_parent, text.replace("\r\n", "\n").replace("\r", "\n")):
            changed.append(block_id)
    return doc.dump(), changed


def _unit_of(el):
    while el is not None and el.getparent() is not None and el.getparent().tag not in CONTAINER_TAGS:
        el = el.getparent()
    return el if el is not None and el.getparent() is not None else None


def _unit_id(unit) -> str | None:
    for bm in unit.iter(w("bookmarkStart")):
        name = bm.get(w("name")) or ""
        if BLOCK_ID_RE.match(name):
            return name
    return None


def _units(doc: _Doc):
    for name in sorted(doc.trees, key=lambda n: (n != "word/document.xml", n)):
        tree = doc.trees[name]
        container = tree.find(w("body")) if tree.tag == w("document") else tree
        if container is None:
            continue
        for child in container:
            if child.tag in UNIT_TAGS:
                yield name, child


def _describe(unit) -> str:
    if unit.tag == w("tbl"):
        rows = unit.findall(w("tr"))
        cols = max((len(r.findall(w("tc"))) for r in rows), default=0)
        cells = []
        for r in rows:
            cells.append(" | ".join(paragraph_text_all(tc) for tc in r.findall(w("tc"))))
        text = " // ".join(cells)
        return f"table {len(rows)}x{cols}: {text}"
    return "paragraph: " + paragraph_text_all(unit)


def paragraph_text_all(el) -> str:
    return " ".join(t for t in (paragraph_text(p) for p in el.iter(w("p"))) if t.strip())


def paragraph_texts(data: bytes, ids: list[str]) -> dict[str, str]:
    doc = _Doc.load(data)
    out = {}
    for bid in ids:
        p = doc.find_block(bid)
        if p is not None and p.tag == w("p"):
            out[bid] = paragraph_text(p)
    return out


def outline(data: bytes, limit: int = 160) -> str:
    doc = _Doc.load(data)
    lines = []
    for part, unit in _units(doc):
        uid = _unit_id(unit)
        if not uid:
            continue
        where = "body" if part == "word/document.xml" else part.split("/")[1].replace(".xml", "")
        desc = _describe(unit)
        if len(desc) > limit:
            desc = desc[: limit - 1] + "…"
        lines.append(f"[{uid}] ({where}) {desc}")
    return "\n".join(lines)


def units_for_blocks(data: bytes, ids: list[str], neighbours: int = 1) -> tuple[list[str], dict[str, str]]:
    """Unit ids containing the given blocks (plus neighbours) and their XML."""
    doc = _Doc.load(data)
    ordered = [(part, unit, _unit_id(unit)) for part, unit in _units(doc)]
    wanted: set[int] = set()
    for bid in ids:
        el = doc.find_block(bid)
        unit = _unit_of(el) if el is not None else None
        for i, (_, u, _) in enumerate(ordered):
            if u is unit:
                for j in range(i - neighbours, i + neighbours + 1):
                    if 0 <= j < len(ordered) and ordered[j][0] == ordered[i][0]:
                        wanted.add(j)
    uids, xml = [], {}
    for i in sorted(wanted):
        uid = ordered[i][2]
        if uid:
            uids.append(uid)
            xml[uid] = etree.tostring(ordered[i][1], encoding="unicode")
    return uids, xml


def _parse_fragment(xml: str, image_rels: set[str] | frozenset = frozenset()) -> list:
    """Parse model-written XML; pictures are allowed only when they reuse an image already in the part."""
    wrapped = (
        f'<root xmlns:w="{W_NS}" xmlns:w14="{W14_NS}" xmlns:r="{R_NS}" xmlns:wp="{WP_NS}" '
        f'xmlns:a="{A_NS}" xmlns:pic="{PIC_NS}">{xml}</root>'
    )
    try:
        root = etree.fromstring(wrapped.encode("utf-8"), _PARSER)
    except etree.XMLSyntaxError as e:
        raise DocxEditError(f"Invalid XML: {e}") from e
    elements = [el for el in root if isinstance(el.tag, str)]
    embed_attr = f"{{{R_NS}}}embed"
    for el in elements:
        if el.tag not in UNIT_TAGS - {w("sdt")}:
            raise DocxEditError(f"Top-level elements must be <w:p> or <w:tbl>, got {el.tag}")
        for node in el.iter():
            if not isinstance(node.tag, str):
                continue
            ns = etree.QName(node).namespace
            in_drawing = ns in DRAWING_NAMESPACES
            if ns not in (W_NS, W14_NS) and not in_drawing:
                raise DocxEditError(f"Element {node.tag} is not allowed")
            if in_drawing and next((a for a in node.iterancestors(w("drawing"))), None) is None:
                raise DocxEditError(f"Element {etree.QName(node).localname} is only allowed inside <w:drawing>")
            if node.tag in FORBIDDEN_TAGS:
                raise DocxEditError(f"Element {etree.QName(node).localname} is not allowed")
            for attr in node.attrib:
                attr_ns = etree.QName(attr).namespace
                if attr == embed_attr and in_drawing:
                    if node.get(attr) not in image_rels:
                        raise DocxEditError("A drawing may only reuse an image already in the document")
                    continue
                if attr_ns not in (W_NS, W14_NS, XML_NS, None):
                    raise DocxEditError(f"Attribute {attr} is not allowed")
            if node.tag == w("drawing") and not any(b.get(embed_attr) for b in node.iter(f"{{{A_NS}}}blip")):
                raise DocxEditError("A drawing must reuse an existing image (r:embed)")
    return elements


def apply_operations(data: bytes, operations: list[dict]) -> tuple[bytes, list[str]]:
    """Apply replace / insert_before / insert_after / delete operations on top-level units."""
    doc = _Doc.load(data)
    touched = []
    for op in operations:
        kind, target = op.get("op"), op.get("target", "")
        content = op.get("content", op.get("xml", ""))
        if kind in ("set_text", "replace_paragraph"):
            p = doc.find_block(target)
            if p is None or p.tag != w("p"):
                raise DocxEditError(f"Unknown paragraph {target}")
            if kind == "set_text":
                set_paragraph_text(p, content.replace("\r\n", "\n"))
                touched.append(p)
            else:
                new = _parse_fragment(content, doc.image_rel_ids(doc.part_of(p)))
                if not new or any(el_new.tag != w("p") for el_new in new):
                    raise DocxEditError("replace_paragraph content must be one or more <w:p>")
                ref = p
                for el_new in new:
                    ref.addnext(el_new)
                    ref = el_new
                p.getparent().remove(p)
                touched.extend(new)
            continue
        el = doc.find_block(target)
        unit = _unit_of(el) if el is not None else None
        if unit is None:
            raise DocxEditError(f"Unknown target {target}")
        container = unit.getparent()
        if kind == "delete":
            siblings = [c for c in container if c.tag in UNIT_TAGS]
            container.remove(unit)
            if len(siblings) == 1:
                container.insert(0, etree.Element(w("p")))
            continue
        new = _parse_fragment(content, doc.image_rel_ids(doc.part_of(unit)))
        if not new:
            raise DocxEditError(f"Operation {kind} on {target} has no content")
        if kind == "replace":
            ref = unit
            for el_new in new:
                ref.addnext(el_new)
                ref = el_new
            container.remove(unit)
        elif kind == "insert_after":
            ref = unit
            for el_new in new:
                ref.addnext(el_new)
                ref = el_new
        elif kind == "insert_before":
            for el_new in new:
                unit.addprevious(el_new)
        else:
            raise DocxEditError(f"Unknown operation {kind}")
        touched.extend(new)
    _tag(doc)
    changed = []
    for unit in touched:
        for p in unit.iter(w("p")):
            bid = _block_id_of(p)
            if bid and bid not in changed:
                changed.append(bid)
    out = doc.dump()
    _check_opens(out)
    return out, changed


def _check_opens(data: bytes) -> None:
    from docx import Document

    try:
        Document(io.BytesIO(data))
    except Exception as e:
        raise DocxEditError(f"Edited document no longer opens: {e}") from e
