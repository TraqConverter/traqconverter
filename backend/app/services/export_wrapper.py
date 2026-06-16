"""Structured export wrapper.

Wraps the translated DOCX (whether it came from the Claude-authored
path or the user's WYSIWYG edits) in the canonical certified-
translation layout:

  [SOURCE PAGES — original PDF rendered as images, 1 page = 1 DOCX page]
  [TRANSLATED CONTENT — merged in from the authored / edited DOCX]
  [CERTIFICATION PAGE — translator + date stamp, optionally via the
   project's selected certification template with {{token}} expansion]

If the team has uploaded a company stamp in Settings, it lives in the
section footer so Word repeats it on every page.

This is invoked from `generate_docx` and `generate_pdf` in
`export_service.py` whenever the project has an authored DOCX or
edited HTML — so the visual edits go through the same wrapper as the
Claude-authored output.
"""
from __future__ import annotations

import logging
import shutil
import tempfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

logger = logging.getLogger(__name__)






def _render_source_pages_as_images(pdf_path: Path, work_dir: Path) -> list:
    """Render each page of `pdf_path` to a PNG at ~2x resolution. Returns
    a list of paths in page order. Failures yield an empty list."""
    try:
        import fitz
    except Exception:
        logger.warning("PyMuPDF not installed — source pages not embedded")
        return []

    pages_dir = work_dir / "source_pages"
    pages_dir.mkdir(parents=True, exist_ok=True)
    out = []
    try:
        doc = fitz.open(str(pdf_path))
    except Exception as e:
        logger.warning("Couldn't open source PDF for embedding: %s", e)
        return []
    try:


        zoom = fitz.Matrix(2, 2)
        for page_num, page in enumerate(doc, start=1):
            try:
                pix = page.get_pixmap(matrix=zoom, alpha=False)
                img_path = pages_dir / f"source_p{page_num}.png"
                pix.save(str(img_path))
                out.append(img_path)
            except Exception as e:
                logger.warning("Failed to render source page %d: %s", page_num, e)
    finally:
        try:
            doc.close()
        except Exception:
            pass
    return out






def _install_stamp_in_footer(section, stamp_path: Path, alignment: str = "right"):
    """Insert the team stamp image into the section's footer. Word
    auto-replicates section footers on every page of that section, so
    the stamp shows up on every page without per-page code."""
    try:
        footer = section.footer

        if footer.paragraphs:
            p = footer.paragraphs[0]
        else:
            p = footer.add_paragraph()


        for run in list(p.runs):
            run.text = ""

        align_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }
        p.alignment = align_map.get(alignment, WD_ALIGN_PARAGRAPH.RIGHT)

        run = p.add_run()


        run.add_picture(str(stamp_path), width=Cm(3))
    except Exception:
        logger.exception("Failed to install team stamp in footer")







def _detect_source_page_orientation(source_path) -> str:
    """Return 'portrait' or 'landscape' based on the source PDF's
    first-page aspect ratio. Defaults to portrait on any error so
    we never break existing behavior.
    """
    try:
        import fitz  # type: ignore
        try:
            d = fitz.open(str(source_path))
            try:
                if len(d) == 0:
                    return "portrait"
                r = d[0].rect
                return "landscape" if r.width > r.height else "portrait"
            finally:
                d.close()
        except Exception:
            pass
    except Exception:
        pass
    try:
        import pypdf
        r = pypdf.PdfReader(str(source_path))
        if not r.pages:
            return "portrait"
        box = r.pages[0].mediabox
        w = float(box.width)
        h = float(box.height)
        return "landscape" if w > h else "portrait"
    except Exception:
        return "portrait"


def _append_body_from(src_doc: Document, dst_doc: Document):
    """Copy paragraphs and tables from `src_doc.body` into `dst_doc.body`,
    preserving formatting at the XML level.

    Trims any trailing empty paragraphs (so the merged body doesn't
    add a blank page before whatever comes next), and drops the
    trailing sectPr (page-size info — we keep dst's own).
    """
    src_body = src_doc.element.body



    W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    T_TAG = "{%s}t" % W_NS
    BR_TAG = "{%s}br" % W_NS
    PICT_TAG = "{%s}pict" % W_NS
    DRAWING_TAG = "{%s}drawing" % W_NS

    def _is_empty_para(el):
        if el.tag.split("}")[-1] != "p":
            return False

        for t in el.iter(T_TAG):
            if (t.text or "").strip():
                return False

        for _ in el.iter(DRAWING_TAG):
            return False
        for _ in el.iter(PICT_TAG):
            return False

        for br in el.iter(BR_TAG):
            if br.get("{%s}type" % W_NS) == "page":
                return False
        return True

    children = list(src_body.iterchildren())

    while children and children[-1].tag.split("}")[-1] == "sectPr":
        children.pop()

    while children and _is_empty_para(children[-1]):
        children.pop()

    dst_body = dst_doc.element.body













    final_sectpr = None
    for ch in list(dst_body.iterchildren()):
        if ch.tag.split("}")[-1] == "sectPr":
            final_sectpr = ch

    R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    EMBED_ATTR = "{%s}embed" % R_NS
    try:
        from docx.opc.constants import RELATIONSHIP_TYPE as _RT
    except Exception:
        _RT = None

    src_part = src_doc.part
    dst_part = dst_doc.part
    rid_map = {}
    images_copied = 0
    for child in children:
        for drawing in child.iter(DRAWING_TAG):
            for el in drawing.iter():
                rid = el.get(EMBED_ATTR)
                if not rid or rid in rid_map:
                    continue
                try:
                    related = src_part.related_parts.get(rid)
                except Exception:
                    related = None
                if related is None:
                    continue
                try:
                    if _RT is not None:
                        new_rid = dst_part.relate_to(related, _RT.IMAGE)
                    else:
                        new_rid = dst_part.relate_to(
                            related,
                            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                        )
                    rid_map[rid] = new_rid
                    images_copied += 1
                except Exception:
                    logger.warning(
                        "Failed to re-register image rel %s during merge", rid
                    )

    if images_copied:
        logger.info(
            "Carried over %d image(s) from authored DOCX into wrapper",
            images_copied,
        )

    for child in children:
        try:
            copied = deepcopy(child)
            if rid_map:
                for el in copied.iter():
                    rid = el.get(EMBED_ATTR)
                    if rid and rid in rid_map:
                        el.set(EMBED_ATTR, rid_map[rid])
            if final_sectpr is not None:
                final_sectpr.addprevious(copied)
            else:
                dst_body.append(copied)
        except Exception:
            logger.warning(
                "Skipped a body element during merge (%s)",
                child.tag.split("}")[-1],
            )







def _append_certification(dst_doc: Document, project, user, work_dir: Path):
    """Append a certification block to the document. If the project has
    a `certification_template_id` set, that DOCX is downloaded,
    {{token}} substitutions are applied, and its body is appended.
    Otherwise we write a simple hardcoded affidavit."""
    user_email = getattr(user, "email", "") or ""




    dst_doc.add_page_break()


    try:
        from app.services.cert_template_service import (
            build_substitution_values,
            substitute_in_docx,
        )
        from app.services.s3_service import generate_presigned_download_url
        from app.models.certification import Certification
        from app.database import SessionLocal
        import requests as _req

        template_id = getattr(project, "certification_template_id", None)
        if template_id:
            db = SessionLocal()
            try:
                cert = (
                    db.query(Certification)
                    .filter(Certification.id == template_id)
                    .first()
                )






                cert_key = getattr(cert, "file_path", None) if cert else None
                cert_name = (getattr(cert, "file_name", "") or "").lower() if cert else ""
                if cert and cert_key and cert_name.endswith(".docx"):
                    url = generate_presigned_download_url(cert_key)
                    r = _req.get(url, timeout=15)
                    if r.ok:
                        team = None
                        try:
                            from app.models.team import Team
                            team = (
                                db.query(Team)
                                .filter(Team.id == project.team_id)
                                .first()
                            )
                        except Exception:
                            pass
                        values = build_substitution_values(
                            user=user, project=project, team=team, extra=None
                        )
                        templated_bytes = substitute_in_docx(r.content, values)
                        templated_doc = Document(BytesIO(templated_bytes))
                        _append_body_from(templated_doc, dst_doc)
                        return
            finally:
                db.close()
    except Exception:
        logger.exception(
            "Cert-template substitution failed — falling back to hardcoded cert"
        )


    from datetime import datetime

    p = dst_doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CERTIFIED TRANSLATION")
    r.bold = True
    r.font.size = Pt(14)

    dst_doc.add_paragraph("")
    dst_doc.add_paragraph(
        "I hereby certify that this translation is accurate and "
        "complete to the best of my knowledge and ability."
    )
    dst_doc.add_paragraph("")
    dst_doc.add_paragraph(f"Translator: {user_email}")
    dst_doc.add_paragraph(
        f"Date: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    )
    dst_doc.add_paragraph("")
    dst_doc.add_paragraph("Signature: ____________________________")






def build_full_export_docx(
    translated_docx_bytes: bytes,
    project,
    user,
) -> BytesIO:
    """Compose the final export DOCX from a translated DOCX:

        [source pages as images]
        [translated content]
        [certification page]

    Team stamp (if configured) lives in the footer and repeats on
    every page.

    Returns a BytesIO containing the finished DOCX, positioned at 0.
    """
    work_dir = Path(tempfile.mkdtemp(prefix="export_wrap_"))
    try:

        out = Document()
        section = out.sections[0]
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)


        try:
            from app.services.export_service import _resolve_team_stamp
            stamp_path, alignment = _resolve_team_stamp(project, work_dir)
            if stamp_path:
                _install_stamp_in_footer(section, Path(stamp_path), alignment)
        except Exception:
            logger.exception(
                "Team-stamp resolution failed — continuing without footer stamp"
            )






        source_added = False
        try:
            from app.services.s3_service import download_file_from_s3
            source_path = work_dir / (project.file_name or "source")
            try:
                download_file_from_s3(project.file_path, source_path)
            except Exception:
                source_path = None

            if (
                source_path
                and source_path.exists()
                and str(source_path).lower().endswith(".pdf")
            ):


                src_orient = _detect_source_page_orientation(source_path)
                if src_orient == "landscape":

                    pg_w, pg_h = Cm(29.7), Cm(21)
                    src_img_w = Cm(27.7)
                else:

                    pg_w, pg_h = Cm(21), Cm(29.7)
                    src_img_w = Cm(19)
                page_imgs = _render_source_pages_as_images(source_path, work_dir)
                if page_imgs:
                    section.page_width = pg_w
                    section.page_height = pg_h
                    section.top_margin = Cm(1)
                    section.bottom_margin = Cm(1)
                    section.left_margin = Cm(1)
                    section.right_margin = Cm(1)
                for i, img in enumerate(page_imgs):
                    if i > 0:
                        new_sect = out.add_section(WD_SECTION.NEW_PAGE)
                        new_sect.page_width = pg_w
                        new_sect.page_height = pg_h
                        new_sect.top_margin = Cm(1)
                        new_sect.bottom_margin = Cm(1)
                        new_sect.left_margin = Cm(1)
                        new_sect.right_margin = Cm(1)
                    p = out.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    run = p.add_run()

                    run.add_picture(str(img), width=src_img_w)
                if page_imgs:
                    source_added = True


                    body_sect = out.add_section(WD_SECTION.NEW_PAGE)
                    body_sect.page_width = pg_w
                    body_sect.page_height = pg_h
                    body_sect.top_margin = Cm(2)
                    body_sect.bottom_margin = Cm(2.2)
                    body_sect.left_margin = Cm(2)
                    body_sect.right_margin = Cm(2)
        except Exception:
            logger.exception(
                "Source-page embedding failed — continuing without source pages"
            )


        try:
            translated_doc = Document(BytesIO(translated_docx_bytes))
            _append_body_from(translated_doc, out)
        except Exception:
            logger.exception("Translated-body merge failed")


            out.add_paragraph(
                "Translation content could not be merged. Please contact support."
            )


        try:
            _append_certification(out, project, user, work_dir)
        except Exception:
            logger.exception("Certification append failed")


        buf = BytesIO()
        out.save(buf)
        buf.seek(0)
        logger.info(
            "Full export built: source_added=%s, size=%d bytes",
            source_added, len(buf.getvalue()),
        )
        return buf
    finally:
        try:
            shutil.rmtree(work_dir, ignore_errors=True)
        except Exception:
            pass
